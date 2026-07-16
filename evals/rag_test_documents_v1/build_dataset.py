#!/usr/bin/env python3
"""Build a hash-pinned multi-domain RAG benchmark from the local test corpus."""

from __future__ import annotations

import argparse
from datetime import date, datetime
import hashlib
import json
import mimetypes
from pathlib import Path
import re
from typing import Any


PACKAGE_DIR = Path(__file__).resolve().parent
REQUIRED_FAITHFULNESS_SCORE = 0.8
DEFAULT_SOURCE_ROOT = Path(
    "/Users/tabbasi/Desktop/_Organized Desktop/AI and RAG Projects/RAG test documents"
)

CANONICAL_SOURCES = {
    "EvaluationRAG/9789240029200-eng.pdf",
    "EvaluationRAG/Amazon-2025-Annual-Report.pdf",
    "EvaluationRAG/Apollo 11 Mission Report.pdf",
    "EvaluationRAG/Global Economic Prospects, January 2026.pdf",
    "EvaluationRAG/IPCC_AR6_SYR_SPM.pdf",
    "EvaluationRAG/Letter scanned.pdf",
    "EvaluationRAG/NIST.AI.100-1.pdf",
    "EvaluationRAG/artemis_plan-20200921.pdf",
    "EvaluationRAG/handwritten table.pdf",
    "EvaluationRAG/sd19.pdf",
    "EvaluationRAG/zero_trust_maturity_model_v2_508.pdf",
    "ISO_IEC 27001_2022.pdf",
    "Manuals/Dell PowerEdge R630 Technical Manual.pdf",
    "Manuals/GS24_Series_UG_101325_En_OS15_REV_1.1_AC_10282025.pdf",
    "Manuals/samsung repair manual.pdf",
    "Updated FIRs/FIR_02_kidnapping.pdf",
    "Updated FIRs/FIR_03_robbery.pdf",
    "Updated FIRs/FIR_04_cybercrime.pdf",
    "Updated FIRs/FIR_05_narcotics.pdf",
    "Updated FIRs/FIR_fictitious.pdf",
    "Urdu fir.jpg",
    "WHO Test Doc.pdf",
    "demo_insurance_policy.pdf",
    "rag_sample_files/Judicial RAG Sample files/c.p._769_k_2025.pdf",
    "rag_sample_files/Judicial RAG Sample files/crl.a._43_2020.pdf",
    "rag_sample_files/Judicial RAG Sample files/crl.r.p._56_2022.pdf",
    "rag_sample_files/Judicial RAG Sample files/wp-25-06-judicial-and-legal-reforms-in-pakistan-1947-2024-a-brief-review.pdf",
    "rag_sample_files/images.png",
    "rag_sample_files/sample.json",
    "rag_sample_files/sample.pdf",
    "rag_sample_files/scs_logo.jpg",
}

EXCLUSIONS: dict[str, set[str]] = {
    "evaluation_question_only": {
        "demo_rag_questions (1).pdf",
        "Updated FIRs/rag_test_questions (1).html",
    },
    "unsupported_benign": {
        "Assignments.xlsx",
        "Employees.csv",
        "Employees.xlsx",
        "Incidents.xlsx",
        "PosMis_agency_24-04-2026 03-53-38_glx4f3tr.dvx.csv",
        "PosMis_agency_24-04-2026 03-53-38_glx4f3tr.dvx.xlsx",
        "Projects.xlsx",
        "Teams.xlsx",
        "Technologies.xlsx",
        "Updated FIRs/the-silent-sentinel.zip",
        "demo_underwriting_notes.txt",
        "rag_sample_files/file_example_XLSX_1000.csv",
        "rag_sample_files/file_example_XLSX_1000.xlsx",
        "rag_sample_files/sample.csv",
        "rag_sample_files/sample.md",
        "rag_sample_files/small.txt",
        "rag_sample_files/timeline.xlsx",
        "rag_sample_files/unicode.txt",
        "sample doc.csv",
        "sample doc.xlsx",
    },
    "system_or_temporary": {
        ".DS_Store",
        "rag_sample_files/.DS_Store",
        "rag_sample_files/Judicial RAG Sample files/.DS_Store",
        "rag_sample_files/Judicial RAG Sample files/~$se_law_summary_1 (1).docx",
    },
    "exact_duplicate": {
        "UPDATED FIRs 2/FIR_02_kidnapping.pdf",
        "UPDATED FIRs 2/FIR_03_robbery.pdf",
        "UPDATED FIRs 2/FIR_04_cybercrime.pdf",
        "UPDATED FIRs 2/FIR_05_narcotics.pdf",
        "UPDATED FIRs 2/FIR_fictitious.pdf",
        "EvaluationRAG/scansmpl.pdf",
    },
    "reject_empty": {"rag_sample_files/empty.txt"},
    "reject_negative_type": {
        "rag_sample_files/Png file sample",
        "rag_sample_files/malicious.exe",
    },
    "reject_oversize": {
        "rag_sample_files/100-mb-example-jpg.jpg",
        "rag_sample_files/100mb.docx",
        "rag_sample_files/200MB-TESTFILE.ORG.pdf",
    },
    "reject_stress_unsupported": {"rag_sample_files/large.txt"},
    "pipeline_incompatible_original": {
        "Askari_Takaful_Plan.docx",
        "rag_sample_files/sample.docx",
        "sample with images.docx",
        "sample-10pages.docx",
    },
    "semantic_format_variant": {"EvaluationRAG/handwritten table.jpeg"},
}

DERIVATIVE_SHEETS = {
    "Assignments.normalized.json": "Assignments.xlsx",
    "Employees.normalized.json": "Employees.xlsx",
    "Incidents.normalized.json": "Incidents.xlsx",
    "Projects.normalized.json": "Projects.xlsx",
    "Teams.normalized.json": "Teams.xlsx",
    "Technologies.normalized.json": "Technologies.xlsx",
}
DERIVATIVE_NOTES = {
    "demo_underwriting_notes.normalized.json": "demo_underwriting_notes.txt"
}
DERIVATIVE_DOCX = {
    "Askari_Takaful_Plan.normalized.json": "Askari_Takaful_Plan.docx",
    "sample.normalized.json": "rag_sample_files/sample.docx",
    "sample-10pages.normalized.json": "sample-10pages.docx",
    "sample-with-images.normalized.json": "sample with images.docx",
}

CORE_SOURCES = {
    "Manuals/samsung repair manual.pdf",
    "Updated FIRs/FIR_02_kidnapping.pdf",
    "Updated FIRs/FIR_03_robbery.pdf",
    "Updated FIRs/FIR_04_cybercrime.pdf",
    "Updated FIRs/FIR_05_narcotics.pdf",
    "Updated FIRs/FIR_fictitious.pdf",
    "demo_insurance_policy.pdf",
    "rag_sample_files/sample.json",
}

OCR_STRESS_SOURCES = {
    "EvaluationRAG/Letter scanned.pdf",
    "EvaluationRAG/handwritten table.pdf",
    "Urdu fir.jpg",
    "rag_sample_files/images.png",
    "rag_sample_files/scs_logo.jpg",
}

FIR_TYPES = (
    ["form_field"] * 14
    + ["single_document_fact"] * 11
    + ["cross_document"] * 6
    + ["reasoning"] * 8
    + ["hallucination_trap"] * 8
    + ["temporal"] * 3
    + ["legal_reasoning"] * 4
)

FIR_AMBIGUITIES = {
    "fir-002": "Reference number is interpreted as the FIR number in the header.",
    "fir-016": "No exact abduction time is stated; 7:30 PM is departure time and 7:15–7:40 PM is the witness window.",
    "fir-021": "Six people received messages, while five transferred money.",
    "fir-027": "The source question is singular, but three FIRs contain IMEI values.",
    "fir-028": "Narcotics evidence bag identifiers are serial numbers, not exhibit numbers.",
    "fir-033": "Robbery has a spent cartridge and CCTV, but murder is the actual weapon-recovery case.",
    "fir-036": "The form and narrative render the same address differently.",
    "fir-041": "The source labels this a trap, but the FIR explicitly states current recovery status.",
    "fir-054": "Non-PPC statutes occur in three FIRs; only narcotics is exclusively outside the PPC.",
}

SAMSUNG_DOCUMENT = "samsung repair manual.pdf"
SAMSUNG_EVIDENCE_SPANS = (
    (
        51,
        (
            "Safety Goggles",
            "Safety Gloves",
            "Safety Mask",
            "Anti-static Wrist Strap",
        ),
    ),
    (
        52,
        (
            "ESD Safe Mat",
            "Ejection Pin",
            "Cross-head Screwdriver",
            "Opening Pick",
            "Opening Tool",
            "Suction Cup",
        ),
    ),
    (
        53,
        (
            "ESD Safe Tweezers",
            "Round Tip Metal Tweezers",
            "Heating Bag",
            "Acrylic Protective Cover for Broken Glass",
        ),
    ),
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def json_default(value: object) -> str:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    raise TypeError(f"Cannot serialize {type(value).__name__}")


def write_json(path: Path, payload: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, default=json_default) + "\n",
        encoding="utf-8",
    )


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.write_text(
        "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows),
        encoding="utf-8",
    )


def build_derivatives(source_root: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit("openpyxl is required to normalize XLSX evidence files") from exc

    generated_dir = PACKAGE_DIR / "generated-corpus"
    generated_dir.mkdir(parents=True, exist_ok=True)
    entries: list[dict[str, Any]] = []
    for output_name, source_name in DERIVATIVE_SHEETS.items():
        source = source_root / source_name
        workbook = load_workbook(source, read_only=True, data_only=True)
        sheets = []
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [str(value).strip() if value is not None else f"column_{index + 1}" for index, value in enumerate(rows[0])]
            records = [
                {headers[index]: value for index, value in enumerate(row) if index < len(headers)}
                for row in rows[1:]
                if any(value is not None for value in row)
            ]
            sheets.append({"name": sheet.title, "records": records})
        payload = {
            "schema_version": 1,
            "title": output_name.removesuffix(".normalized.json"),
            "source_document": source_name,
            "source_sha256": sha256(source),
            "sheets": sheets,
        }
        target = generated_dir / output_name
        write_json(target, payload)
        entries.append(generated_entry(target, source_name))

    for output_name, source_name in DERIVATIVE_NOTES.items():
        source = source_root / source_name
        payload = {
            "schema_version": 1,
            "title": "Insurance Underwriting Notes (Demo)",
            "source_document": source_name,
            "source_sha256": sha256(source),
            "content": source.read_text(encoding="utf-8"),
        }
        target = generated_dir / output_name
        write_json(target, payload)
        entries.append(generated_entry(target, source_name))

    try:
        from docx import Document
    except ImportError as exc:
        raise SystemExit("python-docx is required to normalize DOCX evidence files") from exc
    for output_name, source_name in DERIVATIVE_DOCX.items():
        source = source_root / source_name
        document = Document(source)
        payload = {
            "schema_version": 1,
            "title": output_name.removesuffix(".normalized.json"),
            "source_document": source_name,
            "source_sha256": sha256(source),
            "paragraphs": [
                {"style": paragraph.style.name if paragraph.style else None, "text": paragraph.text.strip()}
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ],
            "tables": [
                [
                    [cell.text.strip() for cell in row.cells]
                    for row in table.rows
                ]
                for table in document.tables
            ],
        }
        target = generated_dir / output_name
        write_json(target, payload)
        entries.append(generated_entry(target, source_name))
    return entries


def generated_entry(path: Path, source_name: str) -> dict[str, Any]:
    return {
        "kind": "generated_derivative",
        "relative_path": str(path.relative_to(PACKAGE_DIR)),
        "upload_title": path.name,
        "source_relative_path": source_name,
        "size_bytes": path.stat().st_size,
        "sha256": sha256(path),
        "content_type": "application/json",
        "tier": "core",
        "quality_preset": "balanced",
    }


def source_entry(source_root: Path, relative_path: str) -> dict[str, Any]:
    path = source_root / relative_path
    tier = "core" if relative_path in CORE_SOURCES else "ocr_stress" if relative_path in OCR_STRESS_SOURCES else "distractor"
    return {
        "kind": "source",
        "relative_path": relative_path,
        "upload_title": path.name,
        "size_bytes": path.stat().st_size,
        "sha256": sha256(path),
        "content_type": mimetypes.guess_type(path.name)[0] or "application/octet-stream",
        "tier": tier,
        "quality_preset": "high_accuracy" if tier == "ocr_stress" else "balanced",
    }


def build_manifests(source_root: Path, generated: list[dict[str, Any]]) -> None:
    discovered = {
        str(path.relative_to(source_root))
        for path in source_root.rglob("*")
        if path.is_file()
    }
    classified = set(CANONICAL_SOURCES)
    for paths in EXCLUSIONS.values():
        overlap = classified.intersection(paths)
        if overlap:
            raise SystemExit(f"Paths classified twice: {sorted(overlap)}")
        classified.update(paths)
    if discovered != classified:
        raise SystemExit(
            f"Unclassified={sorted(discovered - classified)}; missing={sorted(classified - discovered)}"
        )

    canonical = [source_entry(source_root, item) for item in sorted(CANONICAL_SOURCES)]
    files = canonical + generated
    tier_order = {"core": 0, "distractor": 1, "ocr_stress": 2}
    files.sort(
        key=lambda item: (
            tier_order[item["tier"]],
            0 if item["kind"] == "generated_derivative" else 1,
            item["relative_path"],
        )
    )
    manifest = {
        "schema_version": 1,
        "dataset_id": "rag-test-documents-v1",
        "source_root": str(source_root),
        "target_group_path": "/test",
        "source_file_count": len(discovered),
        "canonical_source_count": len(canonical),
        "generated_derivative_count": len(generated),
        "ingest_file_count": len(canonical) + len(generated),
        "files": files,
    }
    write_json(PACKAGE_DIR / "corpus-manifest.json", manifest)

    excluded = []
    duplicate_of = {
        "UPDATED FIRs 2/FIR_02_kidnapping.pdf": "Updated FIRs/FIR_02_kidnapping.pdf",
        "UPDATED FIRs 2/FIR_03_robbery.pdf": "Updated FIRs/FIR_03_robbery.pdf",
        "UPDATED FIRs 2/FIR_04_cybercrime.pdf": "Updated FIRs/FIR_04_cybercrime.pdf",
        "UPDATED FIRs 2/FIR_05_narcotics.pdf": "Updated FIRs/FIR_05_narcotics.pdf",
        "UPDATED FIRs 2/FIR_fictitious.pdf": "Updated FIRs/FIR_fictitious.pdf",
        "EvaluationRAG/scansmpl.pdf": "EvaluationRAG/Letter scanned.pdf",
    }
    for duplicate, canonical_path in duplicate_of.items():
        duplicate_hash = sha256(source_root / duplicate)
        canonical_hash = sha256(source_root / canonical_path)
        if duplicate_hash != canonical_hash:
            raise SystemExit(
                f"Declared duplicate drifted: {duplicate} ({duplicate_hash}) != "
                f"{canonical_path} ({canonical_hash})"
            )
    canonical_by_hash: dict[str, list[str]] = {}
    for relative_path in CANONICAL_SOURCES:
        canonical_by_hash.setdefault(sha256(source_root / relative_path), []).append(relative_path)
    unexpected_duplicates = [
        paths for paths in canonical_by_hash.values() if len(paths) > 1
    ]
    if unexpected_duplicates:
        raise SystemExit(f"Canonical corpus contains unexpected exact duplicates: {unexpected_duplicates}")
    for category, paths in EXCLUSIONS.items():
        for relative_path in sorted(paths):
            path = source_root / relative_path
            entry = {
                "relative_path": relative_path,
                "category": category,
                "size_bytes": path.stat().st_size,
                "sha256": sha256(path),
            }
            if relative_path in duplicate_of:
                entry["duplicate_of"] = duplicate_of[relative_path]
            if relative_path in set(DERIVATIVE_SHEETS.values()) | set(DERIVATIVE_NOTES.values()):
                entry["normalized_derivative"] = next(
                    name
                    for name, source in {**DERIVATIVE_SHEETS, **DERIVATIVE_NOTES}.items()
                    if source == relative_path
                )
            excluded.append(entry)
    write_json(
        PACKAGE_DIR / "exclusion-manifest.json",
        {
            "schema_version": 1,
            "dataset_id": "rag-test-documents-v1",
            "excluded_file_count": len(excluded),
            "files": excluded,
        },
    )


def common_case(**row: Any) -> dict[str, Any]:
    row.setdefault("must_include", [])
    row.setdefault("must_not_include", [])
    row.setdefault("expected_source_docs", [])
    row.setdefault("acceptable_source_pages", [])
    row.setdefault("min_sources", len(row["expected_source_docs"]))
    row.setdefault("must_cite_source", True)
    row.setdefault("min_faithfulness_score", REQUIRED_FAITHFULNESS_SCORE)
    row.setdefault("allow_degraded", False)
    row.setdefault("latency_threshold_ms", 20000)
    return row


def attach_document_bound_expectations(rows: list[dict[str, Any]]) -> None:
    manifest = json.loads(
        (PACKAGE_DIR / "corpus-manifest.json").read_text(encoding="utf-8")
    )
    hashes_by_title = {
        str(item["upload_title"]): str(item["sha256"])
        for item in manifest.get("files", [])
        if isinstance(item, dict) and item.get("upload_title") and item.get("sha256")
    }
    for row in rows:
        pages = list(row["acceptable_source_pages"])
        expectations: list[dict[str, Any]] = []
        for title in row["expected_source_docs"]:
            content_hash = hashes_by_title.get(title)
            if content_hash is None:
                raise SystemExit(
                    f"Case {row['id']} references a document missing from the corpus manifest: {title}"
                )
            if row["id"].startswith("samsung-equipment-"):
                if title != SAMSUNG_DOCUMENT:
                    raise SystemExit(
                        f"Case {row['id']} must reference {SAMSUNG_DOCUMENT}"
                    )
                expectations.extend(
                    {
                        "content_hash": content_hash,
                        "pages": [page],
                        "page_match": "all",
                        "anchors": list(anchors),
                        "min_anchor_recall": 1.0,
                    }
                    for page, anchors in SAMSUNG_EVIDENCE_SPANS
                )
                continue
            expectations.append(
                {
                    "content_hash": content_hash,
                    "pages": pages,
                    "page_match": "all" if pages else "any",
                }
            )
        row["evidence_expectations"] = expectations


def load_fir_cases(source_root: Path) -> tuple[list[dict[str, Any]], dict[str, str]]:
    gold_path = PACKAGE_DIR / "source-gold" / "fir_gold.jsonl"
    rows = [json.loads(line) for line in gold_path.read_text(encoding="utf-8").splitlines() if line.strip()]
    html = (source_root / "Updated FIRs/rag_test_questions (1).html").read_text(encoding="utf-8")
    source_questions = re.findall(r'\{\s*cat:"[^"]+",\s*q:"([^"]+)",\s*src:"[^"]+"\s*\}', html)
    if len(rows) != 54 or len(source_questions) != 54:
        raise SystemExit(f"Expected 54 FIR cases, got gold={len(rows)} html={len(source_questions)}")
    for index, (row, question) in enumerate(zip(rows, source_questions, strict=True)):
        if row["question"] != question:
            raise SystemExit(f"FIR question drift at row {index + 1}: {row['question']!r} != {question!r}")
        row["question_type"] = FIR_TYPES[index]
        row["difficulty"] = "hard" if row["question_type"] in {"cross_document", "reasoning", "hallucination_trap", "temporal", "legal_reasoning"} else "medium"
        row["acceptable_source_pages"] = [1]
        rows[index] = common_case(**row)
    return rows, FIR_AMBIGUITIES


def insurance_cases() -> list[dict[str, Any]]:
    policy = "demo_insurance_policy.pdf"
    notes = "demo_underwriting_notes.normalized.json"
    raw = [
        ("insurance-001", "What benefits does the policy provide?", "The policy provides life protection, accidental death benefits, and optional critical illness coverage under Shariah-compliant principles.", ["life protection", "accidental death", "critical illness"], [policy], "fact"),
        ("insurance-002", "What is excluded from coverage?", "Claims arising from fraud, illegal activities, or pre-existing undisclosed conditions are excluded.", ["fraud", "illegal activities", "pre-existing undisclosed conditions"], [policy], "fact"),
        ("insurance-003", "How long does claims processing take?", "Claims processing takes 7–14 working days.", ["7", "14", "working days"], [policy], "fact"),
        ("insurance-004", "Why do premiums increase with age?", "The underwriting notes state that higher age increases the premium because age is used as a risk factor in the risk multiplier.", ["higher age", "increases", "premium"], [notes], "reasoning"),
        ("insurance-005", "How does occupation affect risk?", "High-risk occupations increase the assessed risk and therefore increase premiums.", ["high-risk jobs", "increase premiums"], [notes], "reasoning"),
        ("insurance-006", "What triggers fraud detection?", "Fraud detection is triggered by inconsistent medical history, frequent claims within a short period, and missing documentation.", ["inconsistent medical history", "frequent claims", "missing documentation"], [notes], "reasoning"),
        ("insurance-007", "How does underwriting affect claim approval?", "Underwriting evaluates age, health, occupation, and lifestyle to rate risk and price premiums; a pre-existing condition that was not disclosed can also make a claim excluded.", ["risk", "premium", "pre-existing", "undisclosed"], [notes, policy], "cross_document"),
        ("insurance-008", "What is the role of reinsurance?", "Large policies are partially transferred through reinsurance to reduce the insurer's exposure.", ["large policies", "transferred", "reduce insurer exposure"], [notes], "fact"),
        ("insurance-009", "Can undisclosed illness affect claims?", "Yes. Claims involving a pre-existing condition that was not disclosed are excluded from coverage.", ["pre-existing", "undisclosed", "excluded"], [policy], "edge_case"),
        ("insurance-010", "Under the insurance policy and underwriting notes, what happens when claim documentation is missing?", "The claim cannot be fully processed until the required supporting documents are supplied, and missing documentation is also a fraud-detection trigger; the sources do not say it causes automatic denial.", ["supporting documents", "missing documentation", "fraud"], [policy, notes], "cross_document"),
    ]
    return [common_case(id=i, question=q, expected_answer=a, must_include=m, expected_source_docs=d, acceptable_source_pages=[1] if d == [policy] else [], question_type=t, difficulty="medium" if t == "fact" else "hard") for i, q, a, m, d, t in raw]


def samsung_cases() -> list[dict[str, Any]]:
    tools = [tool for _, anchors in SAMSUNG_EVIDENCE_SPANS for tool in anchors]
    answer = "The manual lists: " + "; ".join(tools) + "."
    questions = [
        "Give me a list of all the equipment required to repair a Samsung device. Do not skip or miss any.",
        "Give me a list of all the equipments required to repair a samsung. Dont skip or miss any.",
        "What tools and protective equipment are listed for Samsung disassembly and assembly? Include every item.",
        "List every repair tool from the Samsung manual's Tools for Disassembly and Assembly section.",
    ]
    return [
        common_case(
            id=f"samsung-equipment-{index:03d}", question=question,
            question_type="exhaustive_list", difficulty="hard", expected_answer=answer,
            must_include=tools, expected_source_docs=[SAMSUNG_DOCUMENT],
            acceptable_source_pages=[51, 52, 53], min_sources=1,
        )
        for index, question in enumerate(questions, start=1)
    ]


def structured_cases() -> list[dict[str, Any]]:
    e, t, p, a, tech, inc = (
        "Employees.normalized.json", "Teams.normalized.json", "Projects.normalized.json",
        "Assignments.normalized.json", "Technologies.normalized.json", "Incidents.normalized.json",
    )
    raw = [
        ("structured-001", "Who manages Sara Khan?", "Sara Khan reports to Priya Mehta.", ["Sara Khan", "Priya Mehta"], [e]),
        ("structured-002", "Which team has the larger budget, and what is its focus area?", "Platform AI has the larger budget of 3,500,000 and focuses on Model Infrastructure.", ["Platform AI", "3500000", "Model Infrastructure"], [t]),
        ("structured-003", "Which active project had a high-severity incident, and what caused it?", "VectorSearchX is active and had a high-severity incident caused by a Memory Leak.", ["VectorSearchX", "High", "Memory Leak"], [p, inc]),
        ("structured-004", "Who is assigned to VectorSearchX and what are their project roles?", "Lina Zhou is assigned as API Developer and David Kim as UI Integration.", ["Lina Zhou", "API Developer", "David Kim", "UI Integration"], [e, p, a]),
        ("structured-005", "Which technologies are used by SmartRecommend?", "SmartRecommend uses PyTorch and Neo4j.", ["SmartRecommend", "PyTorch", "Neo4j"], [p, tech]),
        ("structured-006", "Which archived project is Sara Khan assigned to, and in what role?", "Sara Khan is assigned to the archived AutoTagger project in the Research role.", ["Sara Khan", "AutoTagger", "Archived", "Research"], [e, p, a]),
        ("structured-007", "Who leads the two teams and who do both leaders report to?", "Priya Mehta leads Applied ML and Carlos Ruiz leads Platform AI; both report to Helen Brooks.", ["Priya Mehta", "Applied ML", "Carlos Ruiz", "Platform AI", "Helen Brooks"], [e, t]),
        ("structured-008", "What low-severity incident affected SmartRecommend, and on what date?", "The low-severity incident was a Feature Store Delay on 2024-08-03.", ["SmartRecommend", "Feature Store Delay", "2024-08-03"], [p, inc]),
        ("structured-009", "What is the total budget of the teams that own active projects?", "Both teams own active projects; their budgets total 5,500,000.", ["5500000"], [p, t]),
        ("structured-010", "Which project uses TensorFlow, who is assigned to it, and what is its status?", "AutoTagger uses TensorFlow; Sara Khan is assigned to it, and the project is Archived.", ["AutoTagger", "TensorFlow", "Sara Khan", "Archived"], [p, tech, a, e]),
    ]
    return [common_case(id=i, question=q, expected_answer=ans, must_include=m, expected_source_docs=docs, question_type="cross_document" if len(docs) > 1 else "structured_fact", difficulty="hard" if len(docs) > 1 else "medium") for i, q, ans, m, docs in raw]


def modality_cases() -> list[dict[str, Any]]:
    raw = [
        ("docx-001", "Which Askari Takaful annual plan is described as ideal for professionals and families?", "The PKR 50,000 annual plan.", ["PKR 50,000", "professionals", "families"], ["Askari_Takaful_Plan.normalized.json"]),
        ("docx-002", "What three principles is Takaful based on in the Askari plan?", "Mutual cooperation, risk sharing, and interest-free investments.", ["Mutual cooperation", "Risk sharing", "Interest-free"], ["Askari_Takaful_Plan.normalized.json"]),
        ("docx-003", "Under what age does the customized Askari plan mention medical and hospitalization coverage?", "It mentions coverage for people under age 55.", ["under age 55"], ["Askari_Takaful_Plan.normalized.json"]),
        ("docx-004", "In the file-format document, how do JPEG and PNG compression differ?", "JPEG uses lossy compression, while PNG uses lossless compression and supports transparency.", ["JPEG", "lossy", "PNG", "lossless", "transparency"], ["sample-10pages.normalized.json"]),
        ("docx-005", "How many pens, pencils, and highlighter colors are needed in the DOCX support demonstration table?", "The table lists 3 pens, 2 pencils, and 2 highlighter colors.", ["3", "pens", "2", "pencils", "2 colors"], ["sample-with-images.normalized.json"]),
        ("json-001", "In sample.json, what user is named, what items are listed, and what is nested.a?", "The user is test, the items are 1, 2, and 3, and nested.a is 1.", ["test", "1", "2", "3", "nested", "a"], ["sample.json"]),
    ]
    return [common_case(id=i, question=q, expected_answer=ans, must_include=m, expected_source_docs=docs, question_type="format_coverage", difficulty="medium") for i, q, ans, m, docs in raw]


def build_dataset(source_root: Path) -> list[dict[str, Any]]:
    fir, ambiguities = load_fir_cases(source_root)
    rows = fir + insurance_cases() + samsung_cases() + structured_cases() + modality_cases()
    attach_document_bound_expectations(rows)
    ids = [row["id"] for row in rows]
    if len(ids) != len(set(ids)):
        raise SystemExit("Dataset contains duplicate case ids")
    write_jsonl(PACKAGE_DIR / "dataset.jsonl", rows)
    ranking_cases = [ranking_case(row) for row in rows]
    write_json(
        PACKAGE_DIR / "ranking-gold.json",
        {
            "schema_version": 2,
            "dataset_id": "rag-test-documents-v1",
            "dataset_sha256": sha256(PACKAGE_DIR / "dataset.jsonl"),
            "corpus_manifest_sha256": sha256(
                PACKAGE_DIR / "corpus-manifest.json"
            ),
            "case_count": len(rows),
            "notes": {
                "current_evaluator_limitation": "The API stores evidence anchors; this companion file adds stage-specific ranking thresholds.",
                "fir_ambiguities": ambiguities,
            },
            "cases": ranking_cases,
        },
    )
    return rows


def ranking_case(row: dict[str, Any]) -> dict[str, Any]:
    required_pages = list(row["acceptable_source_pages"])
    evidence_spans: list[dict[str, Any]] = []
    if row["id"].startswith("samsung-equipment-"):
        evidence_spans = [
            {
                "document": SAMSUNG_DOCUMENT,
                "page": page,
                "must_contain": list(anchors),
            }
            for page, anchors in SAMSUNG_EVIDENCE_SPANS
        ]
    return {
        "id": row["id"],
        "query": row["question"],
        "required_documents": row["expected_source_docs"],
        "required_pages": required_pages,
        "required_evidence_spans": evidence_spans,
        "content_anchors": row["must_include"],
        "retrieval_gold": {
            "candidate_k": 40,
            "min_document_recall_at_k": 1.0,
            "min_required_page_recall_at_k": 1.0 if required_pages else None,
            "min_evidence_span_recall_at_k": 1.0 if evidence_spans else None,
        },
        "reranker_gold": {
            "max_candidates": 40,
            "required_status": "scored",
            "forbidden_statuses": ["fallback", "error"],
            "min_document_recall_at_k": 1.0,
            "min_required_page_recall_at_k": 1.0 if required_pages else None,
            "min_evidence_span_recall_at_k": 1.0 if evidence_spans else None,
        },
        "final_context_gold": {
            "min_sources": row["min_sources"],
            "min_document_recall": 1.0,
            "min_required_page_recall": 1.0 if required_pages else None,
            "min_evidence_span_recall": 1.0 if evidence_spans else None,
        },
    }


def validate_with_application_importer(rows: list[dict[str, Any]]) -> None:
    try:
        from rag.evaluations.datasets import normalize_dataset_content
    except ImportError as exc:
        raise SystemExit(
            "Application dataset validation is mandatory; run with PYTHONPATH=backend/src "
            "and a Python environment containing the workspace document dependencies."
        ) from exc
    content = (PACKAGE_DIR / "dataset.jsonl").read_text(encoding="utf-8")
    normalized = normalize_dataset_content(content, name="RAG test documents v1", source_format="jsonl")
    if len(normalized.cases) != len(rows):
        raise SystemExit("Application importer changed the dataset case count")
    expected_expectation_counts = {
        str(row["id"]): len(row.get("evidence_expectations") or []) for row in rows
    }
    for case in normalized.cases:
        if not case.must_include or not case.expected_source_docs:
            raise SystemExit(f"Case {case.id} lost required scoring expectations")
        expected_count = expected_expectation_counts.get(case.id)
        if expected_count is None:
            raise SystemExit(f"Application importer changed case id {case.id}")
        if len(case.evidence_expectations) != expected_count:
            raise SystemExit(
                f"Case {case.id} lost document-bound evidence expectations"
            )
        if case.must_cite_source and case.min_faithfulness_score < REQUIRED_FAITHFULNESS_SCORE:
            raise SystemExit(
                f"Case {case.id} lost the required citation-support threshold"
            )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", type=Path, default=DEFAULT_SOURCE_ROOT)
    args = parser.parse_args()
    source_root = args.source_root.expanduser().resolve()
    if not source_root.is_dir():
        raise SystemExit(f"Source root does not exist: {source_root}")
    generated = build_derivatives(source_root)
    build_manifests(source_root, generated)
    rows = build_dataset(source_root)
    validate_with_application_importer(rows)
    print(json.dumps({"dataset_cases": len(rows), "ingest_files": len(CANONICAL_SOURCES) + len(generated), "excluded_files": sum(len(paths) for paths in EXCLUSIONS.values())}, indent=2))


if __name__ == "__main__":
    main()
