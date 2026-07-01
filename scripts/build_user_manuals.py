from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    LongTable,
    PageBreak,
    PageTemplate,
    Paragraph,
    Preformatted,
    Spacer,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "docs" / "manuals"
OUTPUT_DIR = SOURCE_DIR / "generated"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(28, 37, 48)
MUTED = RGBColor(91, 103, 112)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


@dataclass(frozen=True)
class Manual:
    source: str
    docx: str
    pdf: str
    running_label: str


MANUALS = [
    Manual(
        source="end-user-manual.md",
        docx="Prudentia-AI-End-User-Manual.docx",
        pdf="Prudentia-AI-End-User-Manual.pdf",
        running_label="Prudentia AI End User Manual",
    ),
    Manual(
        source="administrator-manual.md",
        docx="Prudentia-AI-Administrator-Manual.docx",
        pdf="Prudentia-AI-Administrator-Manual.pdf",
        running_label="Prudentia AI Administrator Manual",
    ),
    Manual(
        source="operator-manual.md",
        docx="Prudentia-AI-Operator-Manual.docx",
        pdf="Prudentia-AI-Operator-Manual.pdf",
        running_label="Prudentia AI Operator Manual",
    ),
]


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for manual in MANUALS:
        blocks = parse_markdown(SOURCE_DIR / manual.source)
        build_docx(blocks, OUTPUT_DIR / manual.docx, manual.running_label)
        build_pdf(blocks, OUTPUT_DIR / manual.pdf, manual.running_label)
        print(f"Wrote {OUTPUT_DIR / manual.docx}")
        print(f"Wrote {OUTPUT_DIR / manual.pdf}")


def parse_markdown(path: Path) -> list[dict]:
    lines = path.read_text(encoding="utf-8").splitlines()
    blocks: list[dict] = []
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue

        if stripped.startswith("```"):
            code: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            blocks.append({"type": "code", "text": "\n".join(code)})
            index += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            blocks.append({"type": "heading", "level": min(level, 3), "text": stripped[level:].strip()})
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and is_table_divider(lines[index + 1]):
            table_lines = [stripped]
            index += 1
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            blocks.append({"type": "table", "rows": parse_table(table_lines)})
            continue

        bullet_match = re.match(r"^-\s+(.*)", stripped)
        if bullet_match:
            text, index = collect_list_item(lines, index, bullet_match.group(1))
            blocks.append({"type": "bullet", "text": text})
            continue

        number_match = re.match(r"^(\d+)\.\s+(.*)", stripped)
        if number_match:
            text, index = collect_list_item(lines, index, number_match.group(2))
            blocks.append({"type": "number", "number": number_match.group(1), "text": text})
            continue

        parts = [stripped]
        index += 1
        while index < len(lines) and is_paragraph_continuation(lines[index]):
            parts.append(lines[index].strip())
            index += 1
        blocks.append({"type": "paragraph", "text": " ".join(parts)})

    return blocks


def is_paragraph_continuation(line: str) -> bool:
    stripped = line.strip()
    if not stripped:
        return False
    return not (
        stripped.startswith("#")
        or stripped.startswith("```")
        or stripped.startswith("|")
        or stripped.startswith("- ")
        or re.match(r"^\d+\.\s+", stripped)
    )


def collect_list_item(lines: list[str], index: int, first_text: str) -> tuple[str, int]:
    parts = [first_text.strip()]
    index += 1
    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if not stripped:
            break
        if raw.startswith("  ") and is_paragraph_continuation(raw):
            parts.append(stripped)
            index += 1
            continue
        break
    return " ".join(parts), index


def is_table_divider(line: str) -> bool:
    stripped = line.strip()
    return bool(re.match(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?$", stripped))


def parse_table(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        if is_table_divider(line):
            continue
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


def build_docx(blocks: list[dict], path: Path, running_label: str) -> None:
    doc = Document()
    configure_docx(doc, running_label)
    title_seen = False

    for block in blocks:
        kind = block["type"]
        if kind == "heading":
            level = block["level"]
            if level == 1 and not title_seen:
                title_seen = True
                p = doc.add_paragraph(style="Manual Title")
                add_inline_runs(p, block["text"], size=24, color=DARK_BLUE, bold=True)
            else:
                p = doc.add_paragraph(style=f"Heading {level}")
                add_inline_runs(p, block["text"])
        elif kind == "paragraph":
            p = doc.add_paragraph()
            add_inline_runs(p, block["text"])
        elif kind == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, block["text"])
        elif kind == "number":
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, block["text"])
        elif kind == "code":
            for code_line in block["text"].splitlines() or [""]:
                p = doc.add_paragraph(style="Manual Code")
                p.add_run(code_line)
        elif kind == "table":
            add_docx_table(doc, block["rows"])

    doc.save(path)


def configure_docx(doc: Document, running_label: str) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title_style = styles.add_style("Manual Title", 1)
    title_style.font.name = "Calibri"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title_style.font.size = Pt(24)
    title_style.font.bold = True
    title_style.font.color.rgb = DARK_BLUE
    title_style.paragraph_format.space_after = Pt(12)
    title_style.paragraph_format.line_spacing = 1.1

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    code_style = styles.add_style("Manual Code", 1)
    code_style.font.name = "Consolas"
    code_style._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code_style.font.size = Pt(9)
    code_style.font.color.rgb = RGBColor(30, 45, 60)
    code_style.paragraph_format.left_indent = Inches(0.25)
    code_style.paragraph_format.space_after = Pt(0)
    code_style.paragraph_format.line_spacing = 1.05

    for list_style_name in ("List Bullet", "List Number"):
        style = styles[list_style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.text = running_label
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.text = "Prudentia AI manuals"
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = MUTED


def add_docx_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    col_count = len(rows[0])
    widths = docx_widths_for_columns(col_count)
    table = doc.add_table(rows=0, cols=col_count)
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_table_width(table, sum(widths), TABLE_INDENT_DXA)

    for row_index, row_values in enumerate(rows):
        row = table.add_row()
        for cell_index, value in enumerate(row_values):
            cell = row.cells[cell_index]
            cell.width = Inches(widths[cell_index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=80, bottom=80, start=120, end=120)
            shading = LIGHT_BLUE if row_index == 0 else None
            if shading:
                set_cell_shading(cell, shading)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline_runs(p, value, bold=row_index == 0)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)


def docx_widths_for_columns(col_count: int) -> list[int]:
    if col_count == 2:
        return [2700, 6660]
    if col_count == 3:
        return [2100, 3260, 4000]
    if col_count == 4:
        return [1800, 2400, 2400, 2760]
    base = TABLE_WIDTH_DXA // col_count
    widths = [base] * col_count
    widths[-1] += TABLE_WIDTH_DXA - sum(widths)
    return widths


def set_table_width(table, width_dxa: int, indent_dxa: int) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top: int, bottom: int, start: int, end: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        element = tc_mar.find(qn(f"w:{tag}"))
        if element is None:
            element = OxmlElement(f"w:{tag}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_inline_runs(paragraph, text: str, size: int | None = None, color: RGBColor | None = None, bold: bool = False) -> None:
    for token, token_style in inline_tokens(text):
        run = paragraph.add_run(token)
        if token_style == "bold" or bold:
            run.bold = True
        if token_style == "code":
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
            run.font.size = Pt(9.5 if size is None else size)
        else:
            run.font.name = "Calibri"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
            if size is not None:
                run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color


def inline_tokens(text: str) -> list[tuple[str, str]]:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    tokens: list[tuple[str, str]] = []
    position = 0
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")
    for match in pattern.finditer(text):
        if match.start() > position:
            tokens.append((text[position:match.start()], "normal"))
        value = match.group(0)
        if value.startswith("`"):
            tokens.append((value[1:-1], "code"))
        else:
            tokens.append((value[2:-2], "bold"))
        position = match.end()
    if position < len(text):
        tokens.append((text[position:], "normal"))
    return tokens


def build_pdf(blocks: list[dict], path: Path, running_label: str) -> None:
    styles = pdf_styles()
    story = []
    title_seen = False
    for block in blocks:
        kind = block["type"]
        if kind == "heading":
            level = block["level"]
            if level == 1 and not title_seen:
                title_seen = True
                story.append(Paragraph(pdf_inline(block["text"]), styles["title"]))
            else:
                story.append(Paragraph(pdf_inline(block["text"]), styles[f"h{level}"]))
        elif kind == "paragraph":
            story.append(Paragraph(pdf_inline(block["text"]), styles["body"]))
        elif kind == "bullet":
            story.append(Paragraph(f"- {pdf_inline(block['text'])}", styles["bullet"]))
        elif kind == "number":
            story.append(Paragraph(f"{block['number']}. {pdf_inline(block['text'])}", styles["number"]))
        elif kind == "code":
            story.append(Preformatted(block["text"], styles["code"]))
            story.append(Spacer(1, 6))
        elif kind == "table":
            story.append(pdf_table(block["rows"], styles))
            story.append(Spacer(1, 8))

    frame = Frame(inch, inch, 6.5 * inch, 9 * inch, id="normal")
    doc = BaseDocTemplate(str(path), pagesize=letter, leftMargin=inch, rightMargin=inch, topMargin=inch, bottomMargin=inch)
    doc.addPageTemplates([PageTemplate(id="manual", frames=[frame], onPage=lambda canvas, document: pdf_page(canvas, document, running_label))])
    doc.build(story)


def pdf_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("ManualTitle", parent=base["Title"], fontName="Helvetica-Bold", fontSize=23, leading=27, textColor=colors.HexColor("#1F4D78"), spaceAfter=14, alignment=TA_LEFT),
        "h1": ParagraphStyle("ManualH1", parent=base["Heading1"], fontName="Helvetica-Bold", fontSize=15, leading=18, textColor=colors.HexColor("#2E74B5"), spaceBefore=14, spaceAfter=8),
        "h2": ParagraphStyle("ManualH2", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=12.5, leading=15, textColor=colors.HexColor("#2E74B5"), spaceBefore=11, spaceAfter=6),
        "h3": ParagraphStyle("ManualH3", parent=base["Heading3"], fontName="Helvetica-Bold", fontSize=11.5, leading=14, textColor=colors.HexColor("#1F4D78"), spaceBefore=8, spaceAfter=4),
        "body": ParagraphStyle("ManualBody", parent=base["BodyText"], fontName="Helvetica", fontSize=10.2, leading=13.2, textColor=colors.HexColor("#1C2530"), spaceAfter=6),
        "bullet": ParagraphStyle("ManualBullet", parent=base["BodyText"], fontName="Helvetica", fontSize=10.2, leading=13.2, leftIndent=16, firstLineIndent=-9, textColor=colors.HexColor("#1C2530"), spaceAfter=3),
        "number": ParagraphStyle("ManualNumber", parent=base["BodyText"], fontName="Helvetica", fontSize=10.2, leading=13.2, leftIndent=18, firstLineIndent=-14, textColor=colors.HexColor("#1C2530"), spaceAfter=3),
        "code": ParagraphStyle("ManualCode", parent=base["Code"], fontName="Courier", fontSize=8.3, leading=10.2, leftIndent=14, rightIndent=10, backColor=colors.HexColor("#F2F4F7"), textColor=colors.HexColor("#1E2D3C"), spaceBefore=3, spaceAfter=3),
        "table": ParagraphStyle("ManualTable", parent=base["BodyText"], fontName="Helvetica", fontSize=8, leading=10, textColor=colors.HexColor("#1C2530")),
        "table_header": ParagraphStyle("ManualTableHeader", parent=base["BodyText"], fontName="Helvetica-Bold", fontSize=8, leading=10, textColor=colors.HexColor("#1C2530")),
    }


def pdf_table(rows: list[list[str]], styles: dict[str, ParagraphStyle]) -> LongTable:
    col_count = len(rows[0])
    widths = pdf_widths_for_columns(col_count)
    table_data = []
    for row_index, row in enumerate(rows):
        style = styles["table_header"] if row_index == 0 else styles["table"]
        table_data.append([Paragraph(pdf_inline(cell), style) for cell in row])
    table = LongTable(table_data, colWidths=widths, repeatRows=1, splitByRow=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B8C2CC")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def pdf_widths_for_columns(col_count: int) -> list[float]:
    if col_count == 2:
        return [1.85 * inch, 4.65 * inch]
    if col_count == 3:
        return [1.45 * inch, 2.25 * inch, 2.8 * inch]
    if col_count == 4:
        return [1.25 * inch, 1.7 * inch, 1.7 * inch, 1.85 * inch]
    return [6.5 * inch / col_count] * col_count


def pdf_inline(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = escape(text)
    text = re.sub(r"`([^`]+)`", r'<font name="Courier">\1</font>', text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", text)
    return text


def pdf_page(canvas, doc, running_label: str) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#5B6770"))
    canvas.drawString(inch, 10.35 * inch, running_label)
    canvas.drawRightString(7.5 * inch, 0.55 * inch, f"Page {doc.page}")
    canvas.restoreState()


if __name__ == "__main__":
    main()
