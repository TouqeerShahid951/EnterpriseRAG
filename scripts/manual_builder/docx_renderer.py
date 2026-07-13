"""DOCX rendering for parsed manual blocks."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(28, 37, 48)
MUTED = RGBColor(91, 103, 112)
LIGHT_BLUE = "E8EEF5"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def build_docx(blocks: list[dict], path: Path, running_label: str) -> None:
    doc = Document()
    configure_docx(doc, running_label)
    title_seen = False

    for block in blocks:
        kind = block["type"]
        if kind == "heading":
            level = block["level"]
            if level == 1 and not title_seen:
                title_seen = True
                p = doc.add_paragraph(style="Manual Title")
                add_inline_runs(p, block["text"], size=24, color=DARK_BLUE, bold=True)
            else:
                p = doc.add_paragraph(style=f"Heading {level}")
                add_inline_runs(p, block["text"])
        elif kind == "paragraph":
            p = doc.add_paragraph()
            add_inline_runs(p, block["text"])
        elif kind == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, block["text"])
        elif kind == "number":
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, block["text"])
        elif kind == "code":
            for code_line in block["text"].splitlines() or [""]:
                p = doc.add_paragraph(style="Manual Code")
                p.add_run(code_line)
        elif kind == "table":
            add_docx_table(doc, block["rows"])

    doc.save(path)


def configure_docx(doc: Document, running_label: str) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title_style = styles.add_style("Manual Title", 1)
    title_style.font.name = "Calibri"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title_style.font.size = Pt(24)
    title_style.font.bold = True
    title_style.font.color.rgb = DARK_BLUE
    title_style.paragraph_format.space_after = Pt(12)
    title_style.paragraph_format.line_spacing = 1.1

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    code_style = styles.add_style("Manual Code", 1)
    code_style.font.name = "Consolas"
    code_style._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code_style.font.size = Pt(9)
    code_style.font.color.rgb = RGBColor(30, 45, 60)
    code_style.paragraph_format.left_indent = Inches(0.25)
    code_style.paragraph_format.space_after = Pt(0)
    code_style.paragraph_format.line_spacing = 1.05

    for list_style_name in ("List Bullet", "List Number"):
        style = styles[list_style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.text = running_label
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.text = "Prudentia AI manuals"
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = MUTED


def add_docx_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    col_count = len(rows[0])
    widths = docx_widths_for_columns(col_count)
    table = doc.add_table(rows=0, cols=col_count)
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_table_width(table, sum(widths), TABLE_INDENT_DXA)

    for row_index, row_values in enumerate(rows):
        row = table.add_row()
        for cell_index, value in enumerate(row_values):
            cell = row.cells[cell_index]
            cell.width = Inches(widths[cell_index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=80, bottom=80, start=120, end=120)
            shading = LIGHT_BLUE if row_index == 0 else None
            if shading:
                set_cell_shading(cell, shading)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline_runs(p, value, bold=row_index == 0)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)


def docx_widths_for_columns(col_count: int) -> list[int]:
    if col_count == 2:
        return [2700, 6660]
    if col_count == 3:
        return [2100, 3260, 4000]
    if col_count == 4:
        return [1800, 2400, 2400, 2760]
    base = TABLE_WIDTH_DXA // col_count
    widths = [base] * col_count
    widths[-1] += TABLE_WIDTH_DXA - sum(widths)
    return widths


def set_table_width(table, width_dxa: int, indent_dxa: int) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top: int, bottom: int, start: int, end: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        element = tc_mar.find(qn(f"w:{tag}"))
        if element is None:
            element = OxmlElement(f"w:{tag}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_inline_runs(paragraph, text: str, size: int | None = None, color: RGBColor | None = None, bold: bool = False) -> None:
    for token, token_style in inline_tokens(text):
        run = paragraph.add_run(token)
        if token_style == "bold" or bold:
            run.bold = True
        if token_style == "code":
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
            run.font.size = Pt(9.5 if size is None else size)
        else:
            run.font.name = "Calibri"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
            if size is not None:
                run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color


def inline_tokens(text: str) -> list[tuple[str, str]]:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    tokens: list[tuple[str, str]] = []
    position = 0
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")
    for match in pattern.finditer(text):
        if match.start() > position:
            tokens.append((text[position:match.start()], "normal"))
        value = match.group(0)
        if value.startswith("`"):
            tokens.append((value[1:-1], "code"))
        else:
            tokens.append((value[2:-2], "bold"))
        position = match.end()
    if position < len(text):
        tokens.append((text[position:], "normal"))
    return tokens
