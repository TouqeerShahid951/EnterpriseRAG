from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
import json
import zipfile

from rag.artifact_jobs.adapters.generated_memory import (
    InMemoryGeneratedArtifactRepository,
)
from rag.artifact_jobs.composer import (
    COMPOSITION_RECORD_TEXT_LIMIT,
    _compose_section,
    compose_document_bundle,
    normalize_bundle_evidence_ids,
)
from rag.artifact_jobs.contracts import (
    ArtifactContentBundle,
    ContentBlock,
    ContentListItem,
    ContentSection,
    ContentTable,
    ContentTableRow,
    DocumentPlan,
    DocumentPlanSection,
    EvidenceBackedContent,
    EvidenceCitation,
    EvidenceManifest,
    EvidenceRecord,
    EvidenceSection,
    PaginatedDocumentSpec,
    PresentationSlide,
    PresentationSpec,
)
from rag.artifact_jobs.execution import ArtifactJobExecutor
from rag.artifact_jobs.llm_json import LlmContractError
from rag.artifact_jobs.layout_profiles import select_layout_profile
from rag.artifact_jobs.planner import plan_document
from rag.artifact_jobs.renderer import render_document
from rag.artifact_jobs.retrieval import retrieve_document_evidence
from rag.artifact_jobs.adapters.job_memory import InMemoryArtifactJobRepository
from rag.artifact_jobs.service import ArtifactJobService, _public_stage_timings
from rag.auth.context import UserContext
from rag.core.config import Settings
from rag.query.qdrant import SearchHit
from rag.documents.models import DocumentRecord
from rag.auth.identity_models import UserRecord
from rag.query.rag_config_models import RagConfigRecord
from rag.artifact_jobs.adapters.storage import LocalGeneratedArtifactStorage


def test_stage_timing_serializer_handles_missing_and_partial_timings() -> None:
    assert _public_stage_timings(None) == {}
    assert _public_stage_timings({"planning": {"duration_ms": 25}, "bad": {"started_at": "x"}}) == {
        "planning": {"duration_ms": 25, "started_at": None, "completed_at": None}
    }


def test_artifact_job_summary_exposes_user_facing_progress() -> None:
    job_repo = InMemoryArtifactJobRepository()
    job = _job(repo=job_repo, requested_formats=("docx", "pptx"))
    updated = job_repo.update_job(job.id, {
        "status": "rendering",
        "stage": "rendering_pptx_slide",
        "progress_pct": 90,
        "stage_progress": {"unit": "slides", "current": 4, "total": 11, "label": "Building slide 4 of 11"},
    })
    service = ArtifactJobService(
        repo_factory=lambda: job_repo,
        artifact_repo_factory=InMemoryGeneratedArtifactRepository,
        queue_factory=lambda: _Queue(),
        retention_days=1,
    )

    summary = service.summary(updated or job)

    assert summary.stage_label == "Rendering files"
    assert summary.stage_detail == "Building slide 4 of 11"
    assert summary.stage_progress is not None
    assert summary.stage_progress.unit == "slides"
    assert summary.attempt_count == (updated.attempt_count if updated else job.attempt_count)


def test_retry_resets_artifact_job_attempt_timestamps() -> None:
    job_repo = InMemoryArtifactJobRepository()
    queue = _Queue()
    job = _job(repo=job_repo)
    running, claimed = job_repo.start_attempt(job.id)
    assert claimed
    old_started_at = datetime(2026, 1, 1, tzinfo=UTC)
    failed = job_repo.update_job((running or job).id, {
        "status": "failed",
        "stage": "failed",
        "progress_pct": 100,
        "started_at": old_started_at,
        "completed_at": datetime(2026, 1, 1, 0, 5, tzinfo=UTC),
        "last_heartbeat_at": datetime(2026, 1, 1, 0, 4, tzinfo=UTC),
        "error_code": "artifact_generation_failed",
        "error_message_safe": "Generation failed.",
    })
    assert failed is not None
    service = ArtifactJobService(
        repo_factory=lambda: job_repo,
        artifact_repo_factory=InMemoryGeneratedArtifactRepository,
        queue_factory=lambda: queue,
        retention_days=1,
    )

    summary = service.retry(job.id, _user_context())

    assert summary.status == "queued"
    assert summary.started_at is None
    assert summary.completed_at is None
    assert summary.last_heartbeat_at is None
    assert summary.updated_at is not None
    assert queue.job_ids == [job.id]
    queued = job_repo.get_job(job.id)
    assert queued is not None
    assert queued.started_at is None
    assert queued.completed_at is None
    assert queued.last_heartbeat_at is None
    restarted, restarted_claimed = job_repo.start_attempt(job.id)
    assert restarted_claimed
    assert restarted is not None
    assert restarted.started_at is not None
    assert restarted.started_at != old_started_at


def test_composer_truncates_evidence_payload_for_llm_prompt() -> None:
    inference = _CapturingSectionInference()
    record = _evidence_record(text="A" * (COMPOSITION_RECORD_TEXT_LIMIT + 100))

    _compose_section(
        _job(),
        _plan(),
        "Facts",
        "Summarize the facts.",
        ["paragraph"],
        [record],
        inference=inference,
        model=None,
    )

    assert "A" * COMPOSITION_RECORD_TEXT_LIMIT in inference.prompts[0]
    assert "A" * (COMPOSITION_RECORD_TEXT_LIMIT + 1) not in inference.prompts[0]


def test_composer_failure_falls_back_with_evidence_ids_preserved() -> None:
    bundle = compose_document_bundle(
        _job(),
        _plan(),
        _evidence_manifest([
            _evidence_record(text="First fact."),
            _evidence_record(evidence_id="E2", text="Second fact."),
        ]),
        inference=_FailingInference(),
        model=None,
        section_timeout_seconds=0.01,
    )

    list_items = bundle.content.sections[0].blocks[0].list_items
    assert [item.evidence_ids for item in list_items] == [["E1"], ["E2"]]
    assert [citation.evidence_id for citation in bundle.content.citations] == ["E1", "E2"]


def test_bundle_evidence_normalization_repairs_missing_ev_prefix() -> None:
    evidence = _evidence_manifest([
        _evidence_record(evidence_id="ev_abc123", text="Evidence fact."),
    ])
    section = ContentSection(
        title="Facts",
        blocks=[ContentBlock(kind="paragraph", text="Evidence fact.", evidence_ids=["abc123"])],
    )
    bundle = ArtifactContentBundle(
        content=EvidenceBackedContent(
            title="FIR summary",
            purpose="Summarize alleged crimes.",
            sections=[section],
            citations=[_citation()],
        ),
        paginated=PaginatedDocumentSpec(title="FIR summary", sections=[section]),
        presentation=PresentationSpec(
            title="FIR summary",
            slides=[PresentationSlide(title="Facts", blocks=section.blocks)],
        ),
    )

    normalized = normalize_bundle_evidence_ids(bundle, evidence)

    assert normalized.content.sections[0].blocks[0].evidence_ids == ["ev_abc123"]
    assert normalized.paginated.sections[0].blocks[0].evidence_ids == ["ev_abc123"]
    assert normalized.presentation.slides[0].blocks[0].evidence_ids == ["ev_abc123"]
    assert [citation.evidence_id for citation in normalized.content.citations] == ["ev_abc123"]


def test_composer_failure_builds_concise_grouped_table() -> None:
    bundle = compose_document_bundle(
        _job(original_request="Create a presentation of all crimes in FIRs"),
        _structured_rows_plan(),
        _evidence_manifest_for_section("Details", [
            EvidenceRecord(
                evidence_id="E1",
                section_title="Details",
                query="all crimes in FIRs details",
                doc_id="fir-1",
                doc_title="FIR_02_kidnapping.pdf",
                chunk_id="fir-1:1",
                text="",
                structured_fields={
                    "1. Complainant Information": "Contact Numbers: 03000000000",
                    "Legal Action and Current Status": "This FIR is registered under Section 365-A PPC for kidnapping for ransom.",
                    "Forensic and Physical Evidence": "CCTV footage and call records were collected.",
                },
            ),
            EvidenceRecord(
                evidence_id="E2",
                section_title="Details",
                query="all crimes in FIRs details",
                doc_id="fir-2",
                doc_title="FIR_05_narcotics.pdf",
                chunk_id="fir-2:1",
                text="",
                structured_fields={
                    "2. Accused Information": "Contact Numbers: 03111111111",
                    "Legal Action and Investigation Status": "The case is registered under Section 9(c) of the Control of Narcotic Substances Act.",
                    "Evidence and Investigative Action": "Recovered packets were sent to PFSA for chemical analysis.",
                },
            ),
        ]),
        inference=_FailingInference(),
        model=None,
        section_timeout_seconds=0.01,
    )

    table = bundle.content.sections[0].blocks[0].table
    assert table is not None
    assert table.headers == ["Document", "Crime / offense", "Key details", "Evidence / status"]
    assert [row.values[1] for row in table.rows] == ["Kidnapping", "Narcotics"]
    assert "Contact Numbers" not in json.dumps(table.model_dump())


def test_composer_lets_llm_choose_table_structure() -> None:
    inference = _CapturingSectionInference()

    bundle = compose_document_bundle(
        _job(original_request="Create a presentation of all risks in policies"),
        _structured_rows_plan(title="Policy Risks", query="all risks in policies"),
        _evidence_manifest_for_section("Details", [
            EvidenceRecord(
                evidence_id="E1",
                section_title="Details",
                query="all risks in policies details",
                doc_id="policy-1",
                doc_title="Access-Control-Policy.pdf",
                chunk_id="policy-1:1",
                text="",
                structured_fields={
                    "Risk": "Privileged access without approval",
                    "Control": "Quarterly access review is required.",
                    "Status": "Open",
                },
            ),
            EvidenceRecord(
                evidence_id="E2",
                section_title="Details",
                query="all risks in policies details",
                doc_id="policy-2",
                doc_title="Vendor-Risk-Policy.pdf",
                chunk_id="policy-2:1",
                text="",
                structured_fields={
                    "Risk": "Vendor assessment missing before onboarding",
                    "Control": "Third-party due diligence must be completed.",
                    "Status": "Approved",
                },
            ),
        ]),
        inference=inference,
        model=None,
        section_timeout_seconds=0.01,
    )

    assert len(inference.prompts) == 2
    assert "Preferred blocks: [\"table\"]" in inference.prompts[0]
    assert "structured_fields" in inference.prompts[0]
    assert "Create final DOCX/PDF and presentation formatting specifications" in inference.prompts[1]
    assert bundle.content.sections[0].blocks[0].kind == "paragraph"


def test_formatter_llm_chooses_presentation_title_subtitle_and_slides() -> None:
    bundle = compose_document_bundle(
        _job(original_request="Create an executive presentation about policy risks", requested_formats=("pptx",)),
        _structured_rows_plan(title="Policy Risks", query="policy risks"),
        _evidence_manifest_for_section("Details", [
            EvidenceRecord(
                evidence_id="E1",
                section_title="Details",
                query="policy risks",
                doc_id="policy-1",
                doc_title="Access-Control-Policy.pdf",
                chunk_id="policy-1:1",
                text="Privileged access without approval is an open risk.",
                structured_fields={"Risk": "Privileged access without approval", "Status": "Open"},
            )
        ]),
        inference=_PresentationFormattingInference(),
        model=None,
        section_timeout_seconds=0.01,
    )

    assert bundle.presentation.title == "Policy Risk Briefing"
    assert bundle.presentation.subtitle == "Executive overview"
    assert [slide.title for slide in bundle.presentation.slides] == ["Priority Risks", "Actions"]


def test_planner_extracts_topic_from_detailed_presentation_request() -> None:
    plan = plan_document(
        _job(
            original_request="Create a detailed presentation of all crimes committed in the FIRs",
            requested_formats=("pptx",),
        ),
        inference=_FailingInference(),
        model=None,
    )

    assert plan.title == "All Crimes Committed In The FIRs"
    assert "Create a detailed presentation" not in plan.purpose


def test_planner_keeps_simple_request_on_fast_path() -> None:
    inference = _CountingPlannerInference()

    plan = plan_document(
        _job(original_request="Create a report about vendor policy", requested_formats=("pdf",)),
        inference=inference,
        model=None,
    )

    assert inference.calls == 0
    assert plan.title == "Vendor Policy"
    assert [section.title for section in plan.sections] == ["Summary", "Details"]


def test_planner_uses_single_llm_call_for_complex_artifact_request() -> None:
    inference = _CountingPlannerInference(_planner_plan_json())

    plan = plan_document(
        _job(
            original_request="Create a comprehensive executive briefing with a timeline and risk matrix for vendor policy changes",
            requested_formats=("pptx",),
        ),
        inference=inference,
        model="planner-model",
    )

    assert inference.calls == 1
    assert plan.title == "Vendor Policy Change Briefing"
    assert [section.retrieval_mode for section in plan.sections] == ["timeline", "structured_rows"]
    assert all(len(section.retrieval_queries) <= 4 for section in plan.sections)
    assert "pptx" in plan.format_requirements


def test_planner_falls_back_when_complex_llm_plan_fails() -> None:
    plan = plan_document(
        _job(
            original_request="Create a comprehensive timeline of all vendor policy changes",
            requested_formats=("pptx",),
        ),
        inference=_FailingInference(),
        model=None,
    )

    assert plan.title == "Timeline Of All Vendor Policy Changes"
    assert [section.title for section in plan.sections] == ["Summary", "Details", "Timeline"]


def test_retrieval_infers_document_family_scope(monkeypatch) -> None:
    monkeypatch.setattr("rag.artifact_jobs.retrieval._retrieve_query", lambda *_args, **_kwargs: [])
    qdrant = _ScopedQdrant()

    evidence = retrieve_document_evidence(
        _job(original_request="Create a presentation of all crimes in FIRs"),
        _structured_rows_plan(),
        inference=object(),
        qdrant=qdrant,  # type: ignore[arg-type]
        config=Settings(document_repository="memory"),
        rag_config=_rag_config(),
        document_repo=_CatalogDocumentRepo([
            _document("fir-1", "FIR_02_kidnapping.pdf"),
            _document("image-1", "preview16.jpg", summary="The image shows a tank firing."),
            _document("manual-1", "Dell PowerEdge R630 Technical Manual.pdf"),
        ]),
    )

    assert qdrant.document_scans == [["fir-1"]]
    assert not qdrant.authorized_scope_scanned
    assert [record.doc_title for record in evidence.records] == ["FIR_02_kidnapping.pdf"]


def test_retrieval_infers_policy_scope_from_doc_type(monkeypatch) -> None:
    monkeypatch.setattr("rag.artifact_jobs.retrieval._retrieve_query", lambda *_args, **_kwargs: [])
    qdrant = _ScopedQdrant({
        "policy-1": "Access-Control-Policy.pdf",
        "policy-2": "Vendor-Risk-Policy.pdf",
        "manual-1": "Dell PowerEdge R630 Technical Manual.pdf",
    })

    evidence = retrieve_document_evidence(
        _job(original_request="Create a presentation of all risks in policies"),
        _structured_rows_plan(title="Policy Risks", query="all risks in policies"),
        inference=object(),
        qdrant=qdrant,  # type: ignore[arg-type]
        config=Settings(document_repository="memory"),
        rag_config=_rag_config(),
        document_repo=_CatalogDocumentRepo([
            _document("policy-1", "Access-Control-Policy.pdf", doc_type="policy"),
            _document("policy-2", "Vendor-Risk-Policy.pdf", doc_type="policy"),
            _document("manual-1", "Dell PowerEdge R630 Technical Manual.pdf", doc_type="manual"),
        ]),
    )

    assert qdrant.document_scans == [["policy-1", "policy-2"]]
    assert [record.doc_title for record in evidence.records] == [
        "Access-Control-Policy.pdf",
        "Vendor-Risk-Policy.pdf",
    ]


def test_retrieval_does_not_scan_every_structured_row_without_scope(monkeypatch) -> None:
    monkeypatch.setattr("rag.artifact_jobs.retrieval._retrieve_query", lambda *_args, **_kwargs: [])
    qdrant = _ScopedQdrant()

    evidence = retrieve_document_evidence(
        _job(original_request="Create a presentation of all standards"),
        _structured_rows_plan(title="Standards", query="all standards"),
        inference=object(),
        qdrant=qdrant,  # type: ignore[arg-type]
        config=Settings(document_repository="memory"),
        rag_config=_rag_config(),
        document_repo=_CatalogDocumentRepo([
            _document("manual-1", "Dell PowerEdge R630 Technical Manual.pdf"),
        ]),
    )

    assert qdrant.document_scans == []
    assert not qdrant.authorized_scope_scanned
    assert evidence.records == []


def test_seeded_response_job_renders_without_artifact_composition(tmp_path) -> None:
    job_repo = InMemoryArtifactJobRepository()
    job = _job(repo=job_repo)
    plan = _plan()
    evidence = _evidence_manifest([_evidence_record(evidence_id="ev_E1")])
    bundle = normalize_bundle_evidence_ids(_bundle(), evidence)
    job_repo.update_job(job.id, {
        "plan_json": plan.model_dump(mode="json"),
        "evidence_manifest_json": evidence.model_dump(mode="json"),
        "content_spec_json": bundle.model_dump(mode="json"),
    })
    artifact_repo = InMemoryGeneratedArtifactRepository()
    executor = ArtifactJobExecutor(
        config=Settings(
            document_repository="memory",
            artifact_libreoffice_required=False,
        ),
        repo_factory=lambda: job_repo,
        artifact_repo_factory=lambda: artifact_repo,
        storage_factory=lambda: LocalGeneratedArtifactStorage(str(tmp_path)),
        identity_repo_factory=lambda: _IdentityRepo(job.user_id, job.permission_version),
        document_repo_factory=_DocumentRepo,
        inference=_FailingInference(),
        qdrant=object(),
        model_name="test-model",
        rag_config=RagConfigRecord(
            base_url="http://example.test",
            chat_model="test-chat",
            embed_model="test-embed",
            faithfulness_model=None,
            chat_timeout_seconds=1,
            embed_timeout_seconds=1,
        ),
    )

    completed = executor.execute(job.id)

    assert completed.status == "complete"
    assert artifact_repo.list_artifacts_for_job(job.id)[0].format == "docx"


def test_seeded_response_job_normalizes_recoverable_evidence_ids(tmp_path) -> None:
    job_repo = InMemoryArtifactJobRepository()
    job = _job(repo=job_repo)
    plan = _plan()
    evidence = _evidence_manifest([_evidence_record(evidence_id="ev_E1")])
    bundle = normalize_bundle_evidence_ids(_bundle(), evidence)
    bad_id = evidence.records[0].evidence_id.removeprefix("ev_")
    broken_block = bundle.content.sections[0].blocks[0].model_copy(update={"evidence_ids": [bad_id]})
    broken_section = bundle.content.sections[0].model_copy(update={"blocks": [broken_block]})
    bundle = bundle.model_copy(update={
        "content": bundle.content.model_copy(update={"sections": [broken_section]}),
        "paginated": bundle.paginated.model_copy(update={"sections": [broken_section]}),
        "presentation": bundle.presentation.model_copy(update={
            "slides": [bundle.presentation.slides[0].model_copy(update={"blocks": [broken_block]})],
        }),
    })
    job_repo.update_job(job.id, {
        "plan_json": plan.model_dump(mode="json"),
        "evidence_manifest_json": evidence.model_dump(mode="json"),
        "content_spec_json": bundle.model_dump(mode="json"),
    })
    artifact_repo = InMemoryGeneratedArtifactRepository()
    executor = ArtifactJobExecutor(
        config=Settings(
            document_repository="memory",
            artifact_libreoffice_required=False,
        ),
        repo_factory=lambda: job_repo,
        artifact_repo_factory=lambda: artifact_repo,
        storage_factory=lambda: LocalGeneratedArtifactStorage(str(tmp_path)),
        identity_repo_factory=lambda: _IdentityRepo(job.user_id, job.permission_version),
        document_repo_factory=_DocumentRepo,
        inference=_FailingInference(),
        qdrant=object(),
        model_name="test-model",
        rag_config=_rag_config(),
    )

    completed = executor.execute(job.id)

    assert completed.status == "complete"
    stored = ArtifactContentBundle.model_validate(job_repo.get_job(job.id).content_spec_json)
    assert stored.content.sections[0].blocks[0].evidence_ids == [evidence.records[0].evidence_id]


def test_seeded_response_job_falls_back_when_repair_json_is_malformed(tmp_path, monkeypatch) -> None:
    job_repo = InMemoryArtifactJobRepository()
    job = _job(repo=job_repo, requested_formats=("pptx",))
    evidence = _evidence_manifest([
        _evidence_record(evidence_id="E1", text="Supported evidence fact."),
    ])
    bad_section = ContentSection(
        title="Facts",
        blocks=[
            ContentBlock(
                kind="bullet_list",
                list_items=[
                    ContentListItem(text="Supported evidence fact.", evidence_ids=["E1"]),
                    ContentListItem(text="Ungrounded claim.", evidence_ids=["E1"]),
                ],
            )
        ],
    )
    bundle = ArtifactContentBundle(
        content=EvidenceBackedContent(
            title="FIR summary",
            purpose="Summarize alleged crimes.",
            sections=[bad_section],
            citations=[_citation()],
        ),
        paginated=PaginatedDocumentSpec(title="FIR summary", sections=[bad_section]),
        presentation=PresentationSpec(title="FIR summary", slides=[PresentationSlide(title="Facts", blocks=bad_section.blocks)]),
    )
    job_repo.update_job(job.id, {
        "plan_json": _plan().model_dump(mode="json"),
        "evidence_manifest_json": evidence.model_dump(mode="json"),
        "content_spec_json": bundle.model_dump(mode="json"),
    })
    monkeypatch.setattr(
        "rag.artifact_jobs.execution.repair_document_bundle",
        lambda *args, **kwargs: (_ for _ in ()).throw(LlmContractError("ArtifactContentBundle validation failed after repair: bad json")),
    )
    artifact_repo = InMemoryGeneratedArtifactRepository()
    executor = ArtifactJobExecutor(
        config=Settings(
            document_repository="memory",
            artifact_libreoffice_required=False,
        ),
        repo_factory=lambda: job_repo,
        artifact_repo_factory=lambda: artifact_repo,
        storage_factory=lambda: LocalGeneratedArtifactStorage(str(tmp_path)),
        identity_repo_factory=lambda: _IdentityRepo(job.user_id, job.permission_version),
        document_repo_factory=_DocumentRepo,
        inference=_FailingInference(),
        qdrant=object(),
        model_name="test-model",
        rag_config=_rag_config(),
    )

    completed = executor.execute(job.id)

    assert completed.status == "complete"
    stored = ArtifactContentBundle.model_validate(job_repo.get_job(job.id).content_spec_json)
    assert [item.text for item in stored.content.sections[0].blocks[0].list_items] == ["Supported evidence fact."]


def test_seeded_response_job_uses_deterministic_repair_before_llm(tmp_path, monkeypatch) -> None:
    job_repo = InMemoryArtifactJobRepository()
    job = _job(repo=job_repo, requested_formats=("pptx",))
    evidence = _evidence_manifest([
        _evidence_record(evidence_id="E1", text="Supported evidence fact."),
    ])
    bad_section = ContentSection(
        title="Facts",
        blocks=[
            ContentBlock(
                kind="bullet_list",
                list_items=[
                    ContentListItem(text="Supported evidence fact.", evidence_ids=["E1"]),
                    ContentListItem(text="Ungrounded claim.", evidence_ids=["E1"]),
                ],
            )
        ],
    )
    bundle = ArtifactContentBundle(
        content=EvidenceBackedContent(
            title="FIR summary",
            purpose="Summarize alleged crimes.",
            sections=[bad_section],
            citations=[_citation()],
        ),
        paginated=PaginatedDocumentSpec(title="FIR summary", sections=[bad_section]),
        presentation=PresentationSpec(title="FIR summary", slides=[PresentationSlide(title="Facts", blocks=bad_section.blocks)]),
    )
    job_repo.update_job(job.id, {
        "plan_json": _plan().model_dump(mode="json"),
        "evidence_manifest_json": evidence.model_dump(mode="json"),
        "content_spec_json": bundle.model_dump(mode="json"),
    })

    def fail_if_called(*args, **kwargs):
        raise AssertionError("LLM repair should not run when deterministic repair already passes validation")

    monkeypatch.setattr("rag.artifact_jobs.execution.repair_document_bundle", fail_if_called)
    artifact_repo = InMemoryGeneratedArtifactRepository()
    executor = ArtifactJobExecutor(
        config=Settings(
            document_repository="memory",
            artifact_libreoffice_required=False,
        ),
        repo_factory=lambda: job_repo,
        artifact_repo_factory=lambda: artifact_repo,
        storage_factory=lambda: LocalGeneratedArtifactStorage(str(tmp_path)),
        identity_repo_factory=lambda: _IdentityRepo(job.user_id, job.permission_version),
        document_repo_factory=_DocumentRepo,
        inference=_FailingInference(),
        qdrant=object(),
        model_name="test-model",
        rag_config=_rag_config(),
    )

    completed = executor.execute(job.id)

    assert completed.status == "complete"
    stored = ArtifactContentBundle.model_validate(job_repo.get_job(job.id).content_spec_json)
    assert [item.text for item in stored.content.sections[0].blocks[0].list_items] == ["Supported evidence fact."]


def test_composer_reports_progress_for_each_composition_batch() -> None:
    events: list[tuple[int, int, str]] = []

    compose_document_bundle(
        _job(),
        _plan(),
        _evidence_manifest([
            _evidence_record(evidence_id=f"E{index}", text=f"Evidence fact {index}.")
            for index in range(1, 82)
        ]),
        inference=_CapturingSectionInference(),
        model=None,
        section_timeout_seconds=0.01,
        progress_callback=lambda current, total, phase: events.append((current, total, phase)),
    )

    assert events == [
        (1, 3, "batch"),
        (1, 3, "batch_complete"),
        (2, 3, "batch"),
        (2, 3, "batch_complete"),
        (3, 3, "batch"),
        (3, 3, "batch_complete"),
        (3, 3, "formatting"),
    ]


def test_executor_maps_composition_progress_to_visible_job_stages() -> None:
    job_repo = InMemoryArtifactJobRepository()
    job = _job(repo=job_repo)
    executor = object.__new__(ArtifactJobExecutor)
    executor.repo_factory = lambda: job_repo

    executor._record_composition_progress(job.id, current=1, total=4, phase="batch")
    first = job_repo.get_job(job.id)
    assert first is not None
    assert first.status == "composing"
    assert first.stage == "composing_1_of_4_batch"
    assert first.progress_pct == 56
    assert first.stage_progress == {
        "unit": "batches",
        "current": 1,
        "total": 4,
        "label": "Composing content batch 1 of 4",
    }

    executor._record_composition_progress(job.id, current=2, total=4, phase="batch_complete")
    middle = job_repo.get_job(job.id)
    assert middle is not None
    assert middle.stage == "composing_2_of_4_batch"
    assert middle.progress_pct == 63

    executor._record_composition_progress(job.id, current=4, total=4, phase="formatting")
    final = job_repo.get_job(job.id)
    assert final is not None
    assert final.stage == "formatting_outputs"
    assert final.progress_pct == 70
    assert final.stage_progress == {
        "unit": "slides",
        "current": 0,
        "total": 1,
        "label": "Choosing title, sections, and slide layout",
    }


def test_executor_records_render_format_progress() -> None:
    job_repo = InMemoryArtifactJobRepository()
    job = _job(repo=job_repo)
    executor = object.__new__(ArtifactJobExecutor)
    executor.repo_factory = lambda: job_repo

    executor._record_render_format_progress(job.id, "pptx", 2, 3)

    updated = job_repo.get_job(job.id)
    assert updated is not None
    assert updated.stage == "rendering_pptx"
    assert updated.stage_progress == {
        "unit": "formats",
        "current": 2,
        "total": 3,
        "label": "Rendering PPTX (2 of 3)",
    }


def test_artifact_job_renderer_uses_shared_renderer_for_all_formats(tmp_path) -> None:
    job_repo = InMemoryArtifactJobRepository()
    job = _job(repo=job_repo, requested_formats=("docx", "pdf", "pptx"))
    artifact_repo = InMemoryGeneratedArtifactRepository()
    storage = LocalGeneratedArtifactStorage(str(tmp_path))
    executor = ArtifactJobExecutor(
        config=Settings(
            document_repository="memory",
            artifact_libreoffice_required=False,
        ),
        repo_factory=lambda: job_repo,
        artifact_repo_factory=lambda: artifact_repo,
        storage_factory=lambda: storage,
        identity_repo_factory=lambda: _IdentityRepo(job.user_id, job.permission_version),
        document_repo_factory=_DocumentRepo,
        inference=object(),
        qdrant=object(),
        model_name="test-model",
        rag_config=RagConfigRecord(
            base_url="http://example.test",
            chat_model="test-chat",
            embed_model="test-embed",
            faithfulness_model=None,
            chat_timeout_seconds=1,
            embed_timeout_seconds=1,
        ),
    )

    failures = executor._render_formats(job, _evidence_manifest([_evidence_record()]), _bundle())

    assert failures == []
    assert {artifact.format for artifact in artifact_repo.list_artifacts_for_job(job.id)} == {"docx", "pdf", "pptx"}
    assert set(job_repo.get_job(job.id).stage_timings_json) == {"storing"}  # type: ignore[union-attr]


def test_deterministic_pdf_renders_polished_primitives() -> None:
    rendered = render_document(
        artifact_format="pdf",
        bundle=_pdf_showcase_bundle(),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
        require_libreoffice=False,
    )

    assert rendered.content_type == "application/pdf"
    assert rendered.content.startswith(b"%PDF")
    assert len(rendered.content) > 2000


def test_adaptive_layout_profile_infers_document_archetypes() -> None:
    assert select_layout_profile(_pdf_showcase_bundle()).key == "evidence_brief"
    assert select_layout_profile(_timeline_bundle()).key == "timeline_report"
    assert select_layout_profile(_operator_guide_bundle()).key == "operator_guide"
    assert select_layout_profile(_standard_layout_bundle()).key == "standard_report"


def test_deterministic_docx_uses_adaptive_style_system() -> None:
    rendered = render_document(
        artifact_format="docx",
        bundle=_pdf_showcase_bundle(),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
        require_libreoffice=False,
    )

    with zipfile.ZipFile(BytesIO(rendered.content)) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        header_xml = "\n".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith("word/header")
        )
        footer_xml = "\n".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith("word/footer")
        )

    assert "EVIDENCE BRIEF" in document_xml
    assert "w:tblGrid" in document_xml
    assert "w:tblHeader" in document_xml
    assert "Prudentia AI" in header_xml
    assert " PAGE " in footer_xml


def test_deterministic_pptx_labels_timeline_profile() -> None:
    rendered = render_document(
        artifact_format="pptx",
        bundle=_timeline_bundle(),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
        require_libreoffice=False,
    )

    from pptx import Presentation

    presentation = Presentation(BytesIO(rendered.content))
    assert "TIMELINE REPORT" in _presentation_text(presentation)
    assert _slide_background_hex(presentation.slides[0]) == "1F3A5F"


def test_deterministic_pptx_table_geometry_uses_content_shape() -> None:
    rendered = render_document(
        artifact_format="pptx",
        bundle=_dynamic_table_bundle(row_count=2),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
        require_libreoffice=False,
    )

    from pptx import Presentation

    presentation = Presentation(BytesIO(rendered.content))
    table = _presentation_tables(presentation)[0]
    widths = [column.width for column in table.columns]
    row_heights = [row.height for row in table.rows]

    assert widths[1] > widths[0] * 2
    assert widths[1] > widths[2]
    assert row_heights[1] > row_heights[2]


def test_deterministic_pptx_splits_tables_by_estimated_height() -> None:
    rendered = render_document(
        artifact_format="pptx",
        bundle=_dynamic_table_bundle(row_count=8),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
        require_libreoffice=False,
    )

    from pptx import Presentation

    presentation = Presentation(BytesIO(rendered.content))

    assert len(_presentation_tables(presentation)) >= 2
    assert "Event 8" in _presentation_table_text(presentation)


def test_deterministic_pptx_splits_large_lists_without_truncation() -> None:
    rendered = render_document(
        artifact_format="pptx",
        bundle=_large_list_bundle(item_count=18),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
        require_libreoffice=False,
    )

    from pptx import Presentation

    presentation = Presentation(BytesIO(rendered.content))
    text = _presentation_text(presentation)
    assert "Item 18 evidence fact." in text
    assert len(presentation.slides) >= 4
    slides = list(presentation.slides)
    assert _slide_background_hex(slides[0]) == "173B3F"
    content_backgrounds = [_slide_background_hex(slide) for slide in slides[1:]]
    assert content_backgrounds[0] == "DCEBE8"
    assert len(set(content_backgrounds)) > 1


def test_deterministic_pptx_skips_generated_references_content_slide() -> None:
    bundle = _bundle()
    placeholder_slide = PresentationSlide(
        title="References",
        blocks=[ContentBlock(kind="heading", text="References")],
    )
    bundle = bundle.model_copy(update={
        "presentation": bundle.presentation.model_copy(update={
            "slides": [*bundle.presentation.slides, placeholder_slide],
        })
    })
    rendered = render_document(
        artifact_format="pptx",
        bundle=bundle,
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
        require_libreoffice=False,
    )

    from pptx import Presentation

    presentation = Presentation(BytesIO(rendered.content))
    reference_slide_texts = [
        _slide_text(slide)
        for slide in presentation.slides
        if "References" in _slide_text(slide).splitlines()
    ]
    assert len(reference_slide_texts) == 1
    assert "FIR_02_kidnapping.pdf" in reference_slide_texts[0]
    assert reference_slide_texts[0].splitlines().count("References") == 1


def test_normalizer_drops_generated_references_content_slide() -> None:
    bundle = _bundle()
    placeholder_slide = PresentationSlide(
        title="References",
        blocks=[ContentBlock(kind="heading", text="References")],
    )
    bundle = bundle.model_copy(update={
        "presentation": bundle.presentation.model_copy(update={
            "slides": [*bundle.presentation.slides, placeholder_slide],
        })
    })

    normalized = normalize_bundle_evidence_ids(bundle, _evidence_manifest([_evidence_record()]))

    assert [slide.title for slide in normalized.presentation.slides] == ["Facts"]


def test_deterministic_pptx_reports_slide_progress_without_changing_content() -> None:
    events: list[tuple[int, int, str]] = []
    rendered = render_document(
        artifact_format="pptx",
        bundle=_large_list_bundle(item_count=18),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
        require_libreoffice=False,
        progress_callback=lambda current, total, label: events.append((current, total, label)),
    )

    from pptx import Presentation

    presentation = Presentation(BytesIO(rendered.content))
    assert events[0] == (1, len(presentation.slides), "Title slide")
    assert events[-1] == (len(presentation.slides), len(presentation.slides), "References")
    assert len(events) == len(presentation.slides)


def _job(
    *,
    repo: InMemoryArtifactJobRepository | None = None,
    requested_formats: tuple[str, ...] = ("docx",),
    original_request: str = "Generate an evidence-backed artifact.",
):
    selected_repo = repo or InMemoryArtifactJobRepository()
    return selected_repo.create_job(
        client_request_id="request-1",
        user_id="user-1",
        permission_version=1,
        user_email="user@example.test",
        account_type="platform_admin",
        group_paths=["/"],
        clearance_level="NATO_RESTRICTED",
        session_id="session-1",
        trace_id="trace-1",
        original_request=original_request,
        requested_formats=list(requested_formats),
        group_path=None,
        document_ids=[],
        conversation_context=[],
        retention_days=1,
    )


def _user_context() -> UserContext:
    return UserContext(
        user_id="user-1",
        email="user@example.test",
        account_type="platform_admin",
        group_paths=("/",),
        clearance_level="NATO_RESTRICTED",
        permission_version=1,
    )


def _plan() -> DocumentPlan:
    return DocumentPlan(
        title="FIR summary",
        purpose="Summarize alleged crimes.",
        sections=[
            DocumentPlanSection(
                title="Facts",
                objective="Summarize the facts.",
                preferred_blocks=["paragraph"],
                retrieval_queries=["facts"],
            )
        ],
    )


def _structured_rows_plan(*, title: str = "All Crimes In FIRs", query: str = "all crimes in FIRs") -> DocumentPlan:
    return DocumentPlan(
        title=title,
        purpose=f"Create an evidence-grounded artifact about {query}.",
        sections=[
            DocumentPlanSection(
                title="Details",
                objective=f"Extract structured facts about {query}.",
                preferred_blocks=["table"],
                retrieval_queries=[query],
                retrieval_mode="structured_rows",
                coverage_requirement="complete_authorized_scope",
            )
        ],
    )


def _evidence_record(*, evidence_id: str = "E1", text: str = "Evidence fact.") -> EvidenceRecord:
    return EvidenceRecord(
        evidence_id=evidence_id,
        section_title="Facts",
        query="facts",
        doc_id="doc-1",
        doc_title="FIR_02_kidnapping.pdf",
        chunk_id=f"chunk-{evidence_id}",
        page_start=1,
        page_end=1,
        text=text,
    )


def _evidence_manifest(records: list[EvidenceRecord]) -> EvidenceManifest:
    return _evidence_manifest_for_section("Facts", records)


def _evidence_manifest_for_section(title: str, records: list[EvidenceRecord]) -> EvidenceManifest:
    return EvidenceManifest(
        sections=[
            EvidenceSection(
                title=title,
                objective="Summarize the facts.",
                retrieval_mode="focused_search",
                coverage_requirement="relevant_evidence",
                coverage_status="sufficient",
                records=records,
            )
        ]
    )


def _bundle() -> ArtifactContentBundle:
    citation = _citation()
    section = ContentSection(
        title="Facts",
        blocks=[ContentBlock(kind="paragraph", text="Evidence fact.", evidence_ids=["E1"])],
    )
    content = EvidenceBackedContent(
        title="FIR summary",
        purpose="Summarize alleged crimes.",
        sections=[section],
        citations=[citation],
    )
    return ArtifactContentBundle(
        content=content,
        paginated=PaginatedDocumentSpec(title="FIR summary", sections=[section]),
        presentation=PresentationSpec(
            title="FIR summary",
            slides=[PresentationSlide(title="Facts", blocks=section.blocks)],
        ),
    )


def _pdf_showcase_bundle() -> ArtifactContentBundle:
    citation = _citation()
    section = ContentSection(
        title="Highlights",
        blocks=[
            ContentBlock(kind="paragraph", text="Evidence fact.", evidence_ids=["E1"]),
            ContentBlock(
                kind="bullet_list",
                list_items=[
                    ContentListItem(text="First cited point.", evidence_ids=["E1"]),
                    ContentListItem(text="Second cited point.", evidence_ids=["E1"]),
                ],
            ),
            ContentBlock(
                kind="callout",
                text="Prioritize cited conclusions over unsupported details.",
                evidence_ids=["E1"],
            ),
            ContentBlock(
                kind="table",
                table=ContentTable(
                    headers=["Document", "Topic", "Outcome", "Status", "Owner"],
                    rows=[
                        ContentTableRow(
                            values=["FIR_02_kidnapping.pdf", "Kidnapping", "Evidence reviewed", "Open", "Investigator"],
                            evidence_ids=["E1"],
                        )
                    ],
                ),
            ),
        ],
    )
    return ArtifactContentBundle(
        content=EvidenceBackedContent(
            title="PDF layout showcase",
            purpose="Exercise PDF layout primitives.",
            sections=[section],
            citations=[citation],
        ),
        paginated=PaginatedDocumentSpec(
            title="PDF layout showcase",
            subtitle="Evidence-backed output",
            sections=[section],
            include_references=True,
        ),
        presentation=PresentationSpec(
            title="PDF layout showcase",
            slides=[PresentationSlide(title="Highlights", blocks=section.blocks)],
        ),
    )


def _timeline_bundle() -> ArtifactContentBundle:
    citation = _citation()
    section = ContentSection(
        title="Timeline",
        blocks=[
            ContentBlock(
                kind="numbered_list",
                list_items=[
                    ContentListItem(text="Complaint registered at 21:30.", evidence_ids=["E1"]),
                    ContentListItem(text="Recovered property logged after midnight.", evidence_ids=["E1"]),
                ],
            ),
            ContentBlock(
                kind="table",
                table=ContentTable(
                    headers=["Time", "Event"],
                    rows=[
                        ContentTableRow(values=["21:30", "Complaint registered."], evidence_ids=["E1"]),
                        ContentTableRow(values=["00:25", "Medical note completed."], evidence_ids=["E1"]),
                    ],
                ),
            ),
        ],
    )
    return ArtifactContentBundle(
        content=EvidenceBackedContent(
            title="Incident timeline",
            purpose="Build a chronological timeline of events.",
            sections=[section],
            citations=[citation],
        ),
        paginated=PaginatedDocumentSpec(title="Incident timeline", sections=[section]),
        presentation=PresentationSpec(
            title="Incident timeline",
            slides=[PresentationSlide(title="Timeline", blocks=section.blocks)],
        ),
    )


def _operator_guide_bundle() -> ArtifactContentBundle:
    citation = _citation()
    section = ContentSection(
        title="SOP checklist",
        blocks=[
            ContentBlock(
                kind="numbered_list",
                list_items=[
                    ContentListItem(text="Review authorized source records.", evidence_ids=["E1"]),
                    ContentListItem(text="Confirm each generated claim has a citation.", evidence_ids=["E1"]),
                ],
            )
        ],
    )
    return ArtifactContentBundle(
        content=EvidenceBackedContent(
            title="Artifact generation SOP",
            purpose="Provide a workflow checklist for artifact review.",
            sections=[section],
            citations=[citation],
        ),
        paginated=PaginatedDocumentSpec(title="Artifact generation SOP", sections=[section]),
        presentation=PresentationSpec(
            title="Artifact generation SOP",
            slides=[PresentationSlide(title="SOP checklist", blocks=section.blocks)],
        ),
    )


def _standard_layout_bundle() -> ArtifactContentBundle:
    citation = _citation()
    section = ContentSection(
        title="Overview",
        blocks=[
            ContentBlock(
                kind="paragraph",
                text="A neutral layout showcase should remain on the standard report profile.",
                evidence_ids=["E1"],
            )
        ],
    )
    return ArtifactContentBundle(
        content=EvidenceBackedContent(
            title="Layout showcase",
            purpose="Summarize available materials in a neutral document.",
            sections=[section],
            citations=[citation],
        ),
        paginated=PaginatedDocumentSpec(title="Layout showcase", sections=[section]),
        presentation=PresentationSpec(
            title="Layout showcase",
            slides=[PresentationSlide(title="Overview", blocks=section.blocks)],
        ),
    )


def _dynamic_table_bundle(*, row_count: int) -> ArtifactContentBundle:
    citation = _citation()
    long_event = (
        "Event {index}: custody metadata reconciliation requires a careful comparison of intake notes, "
        "handoff records, and ledger timestamps before this row can be treated as a final conclusion. "
        "The reviewer should keep the source caveat visible, preserve the operational sequence, and avoid "
        "collapsing unresolved chain-of-custody gaps into a confirmed finding."
    )
    rows = [
        ContentTableRow(
            values=[
                f"{8 + index:02d}:15",
                long_event.format(index=index) if index == 1 or row_count > 2 else f"Event {index}: short verified update.",
                "Needs timestamp reconciliation" if index == 1 else "Verified",
            ],
            evidence_ids=["E1"],
        )
        for index in range(1, row_count + 1)
    ]
    section = ContentSection(
        title="Timeline",
        blocks=[
            ContentBlock(
                kind="table",
                table=ContentTable(
                    headers=["Time", "Event Description", "Status"],
                    rows=rows,
                ),
            )
        ],
    )
    return ArtifactContentBundle(
        content=EvidenceBackedContent(
            title="Incident timeline",
            purpose="Build a chronological timeline of events.",
            sections=[section],
            citations=[citation],
        ),
        paginated=PaginatedDocumentSpec(title="Incident timeline", sections=[section]),
        presentation=PresentationSpec(
            title="Incident timeline",
            slides=[PresentationSlide(title="Timeline", blocks=section.blocks)],
        ),
    )


def _large_list_bundle(*, item_count: int) -> ArtifactContentBundle:
    citation = _citation()
    section = ContentSection(
        title="Facts",
        blocks=[
            ContentBlock(
                kind="bullet_list",
                list_items=[
                    ContentListItem(text=f"Item {index} evidence fact.", evidence_ids=["E1"])
                    for index in range(1, item_count + 1)
                ],
            )
        ],
    )
    content = EvidenceBackedContent(
        title="FIR summary",
        purpose="Summarize alleged crimes.",
        sections=[section],
        citations=[citation],
    )
    return ArtifactContentBundle(
        content=content,
        paginated=PaginatedDocumentSpec(title="FIR summary", sections=[section]),
        presentation=PresentationSpec(
            title="FIR summary",
            slides=[PresentationSlide(title="Facts", blocks=section.blocks)],
        ),
    )


def _rag_config() -> RagConfigRecord:
    return RagConfigRecord(
        base_url="http://example.test",
        chat_model="test-chat",
        embed_model="test-embed",
        faithfulness_model=None,
        chat_timeout_seconds=1,
        embed_timeout_seconds=1,
    )


def _document(
    document_id: str,
    title: str,
    *,
    doc_type: str = "report",
    summary: str | None = None,
) -> DocumentRecord:
    return DocumentRecord(
        id=document_id,
        title=title,
        source_id=title,
        group_path="/admin",
        clearance_level="NATO_RESTRICTED",
        doc_type=doc_type,
        language="en",
        effective_date=None,
        expiry_date=None,
        description=None,
        summary=summary,
        topics=(),
        llm_topics=(),
        auto_doc_type=doc_type,
        extracted_dates={},
        metadata_flags={},
        is_current=True,
        superseded_by=None,
        pending_supersedes=(),
        content_hash=None,
        uploaded_by=None,
        file_path=None,
        ingest_status="complete",
        deleted_at=None,
        created_at=None,
        updated_at=None,
    )


def _citation() -> EvidenceCitation:
    citation = EvidenceCitation(
        evidence_id="E1",
        doc_id="doc-1",
        doc_title="FIR_02_kidnapping.pdf",
        chunk_id="chunk-E1",
        page_start=1,
        page_end=1,
    )
    return citation


def _presentation_text(presentation) -> str:
    return "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )


def _slide_text(slide) -> str:
    return "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))


def _presentation_tables(presentation):
    return [
        shape.table
        for slide in presentation.slides
        for shape in slide.shapes
        if getattr(shape, "has_table", False)
    ]


def _presentation_table_text(presentation) -> str:
    return "\n".join(
        cell.text
        for table in _presentation_tables(presentation)
        for row in table.rows
        for cell in row.cells
    )


def _slide_background_hex(slide) -> str:
    return str(slide.background.fill.fore_color.rgb)


class _CapturingSectionInference:
    def __init__(self) -> None:
        self.prompts: list[str] = []

    def generate_json(self, *, prompt: str, **_kwargs) -> str:
        self.prompts.append(prompt)
        if "FormatSpecifications" in prompt or "formatting specifications" in prompt:
            return _format_spec_json()
        return json.dumps({
            "title": "Facts",
            "blocks": [{"kind": "paragraph", "text": "Evidence fact.", "evidence_ids": ["E1"]}],
        })


class _PresentationFormattingInference:
    def generate_json(self, *, prompt: str, **_kwargs) -> str:
        if "formatting specifications" in prompt:
            return _format_spec_json(
                title="Policy Risk Briefing",
                subtitle="Executive overview",
                slide_titles=("Priority Risks", "Actions"),
            )
        return json.dumps({
            "title": "Details",
            "blocks": [
                {
                    "kind": "bullet_list",
                    "list_items": [
                        {"text": "Privileged access without approval is open.", "evidence_ids": ["E1"]}
                    ],
                }
            ],
        })


class _CountingPlannerInference:
    def __init__(self, response: str | None = None) -> None:
        self.response = response or "{}"
        self.calls = 0
        self.prompts: list[str] = []

    def generate_json(self, *, prompt: str, **_kwargs) -> str:
        self.calls += 1
        self.prompts.append(prompt)
        return self.response


class _FailingInference:
    def generate_json(self, **_kwargs) -> str:
        raise TimeoutError("composition call failed")


def _planner_plan_json() -> str:
    return json.dumps({
        "title": "Vendor Policy Change Briefing",
        "purpose": "Create an evidence-grounded briefing on vendor policy changes.",
        "audience": "Executive leadership",
        "language": "English",
        "tone": "professional",
        "detail_level": "detailed",
        "document_type": "presentation",
        "assumptions": ["Use only authorized retrieved evidence."],
        "clarification_questions": [],
        "sections": [
            {
                "title": "Timeline",
                "objective": "Build a chronological view of vendor policy changes.",
                "preferred_blocks": ["numbered_list", "table"],
                "retrieval_queries": [
                    "vendor policy changes timeline",
                    "vendor policy amendments dates",
                    "vendor policy version history",
                    "vendor policy chronology",
                ],
                "retrieval_mode": "timeline",
                "coverage_requirement": "relevant_evidence",
            },
            {
                "title": "Risk Matrix",
                "objective": "Extract policy risks and classify their status.",
                "preferred_blocks": ["table", "bullet_list"],
                "retrieval_queries": ["vendor policy risks", "vendor risk matrix"],
                "retrieval_mode": "structured_rows",
                "coverage_requirement": "complete_authorized_scope",
            },
        ],
        "format_requirements": {
            "pptx": ["Use executive slide titles.", "Use tables for risks."],
        },
    })


def _format_spec_json(
    *,
    title: str = "Formatted Artifact",
    subtitle: str = "Evidence-backed output",
    slide_titles: tuple[str, ...] = ("Facts",),
) -> str:
    block = {"kind": "paragraph", "text": "Evidence fact.", "evidence_ids": ["E1"]}
    return json.dumps({
        "paginated": {
            "title": title,
            "subtitle": subtitle,
            "sections": [{"title": "Facts", "blocks": [block]}],
            "include_references": True,
            "include_coverage_notes": True,
        },
        "presentation": {
            "title": title,
            "subtitle": subtitle,
            "slides": [{"title": slide_title, "blocks": [block]} for slide_title in slide_titles],
            "include_references_slide": True,
        },
    })


class _CatalogDocumentRepo:
    def __init__(self, documents: list[DocumentRecord]) -> None:
        self.documents = documents

    def list_documents(self, *, state: str = "active") -> list[DocumentRecord]:
        assert state == "active"
        return self.documents


class _ScopedQdrant:
    def __init__(self, titles: dict[str, str] | None = None) -> None:
        self.document_scans: list[list[str]] = []
        self.authorized_scope_scanned = False
        self.titles = titles or {"fir-1": "FIR_02_kidnapping.pdf"}

    def retrieve_document_chunks(
        self,
        *,
        document_ids: list[str],
        qdrant_filter: dict[str, object],
        structured_only: bool,
        limit: int,
    ) -> list[SearchHit]:
        _ = qdrant_filter, structured_only, limit
        self.document_scans.append(list(document_ids))
        return [
            _search_hit(doc_id, self.titles[doc_id], f"Evidence for {self.titles[doc_id]}.")
            for doc_id in document_ids
            if doc_id in self.titles
        ]

    def retrieve_authorized_chunks(self, **_kwargs) -> list[SearchHit]:
        self.authorized_scope_scanned = True
        return [_search_hit("manual-1", "Dell PowerEdge R630 Technical Manual.pdf", "Standard: Ethernet.")]


class _Queue:
    def __init__(self) -> None:
        self.job_ids: list[str] = []

    def enqueue(self, job_id: str) -> None:
        self.job_ids.append(job_id)


def _search_hit(doc_id: str, title: str, text: str) -> SearchHit:
    return SearchHit(
        point_id=f"{doc_id}:1",
        score=1.0,
        payload={
            "doc_id": doc_id,
            "doc_title": title,
            "chunk_id": f"{doc_id}:1",
            "page_start": 1,
            "page_end": 1,
            "chunk_type": "table_row",
            "text": text,
            "structured_fields": [{"label": "Crime", "value": text}],
        },
    )


class _IdentityRepo:
    def __init__(self, user_id: str, permission_version: int) -> None:
        self.user = UserRecord(
            id=user_id,
            email="user@example.test",
            name="Test User",
            password_hash="",
            is_active=True,
            account_type="platform_admin",
            clearance_level="NATO_RESTRICTED",
            must_change_password=False,
            permission_version=permission_version,
            group_paths=("/",),
        )

    def get_user_by_id(self, user_id: str) -> UserRecord | None:
        return self.user if user_id == self.user.id else None


class _DocumentRepo:
    def __init__(self) -> None:
        self.events: list[dict[str, object]] = []

    def get_document(self, document_id: str):
        return _document(document_id, "Authorized source")

    def append_audit_event(self, **kwargs):
        self.events.append(kwargs)
