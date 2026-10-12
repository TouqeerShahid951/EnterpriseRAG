"""Artifact publication requires support in every adapted format and bounded repair."""

from datetime import UTC, datetime
from io import BytesIO
import json
from time import perf_counter
from unittest.mock import Mock

import pytest

from rag.artifact_jobs.adapters.query_generation import (
    QueryRuntimeArtifactJsonGenerator,
)
from rag.artifact_jobs.contracts import (
    ContentBlock,
    ContentListItem,
    ContentTable,
    ContentTableRow,
    KeyValueEntry,
)
from rag.artifact_jobs.generation.citation_validation import (
    bundle_claims,
    verify_bundle_citations,
)
from rag.artifact_jobs.renderer import render_document
from rag.artifact_jobs.execution.errors import (
    ArtifactJobCancelled,
    ArtifactPermissionChanged,
)
from generation_optimization_support import (
    artifact_job,
    content_bundle,
    document_plan,
    evidence_manifest,
    evidence_record,
)
from generation_execution_cases import seeded_executor
from rag.artifact_jobs.adapters.job_memory import InMemoryArtifactJobRepository
from rag.artifact_jobs.adapters.generated_memory import (
    InMemoryGeneratedArtifactRepository,
)
from rag.shared.contracts.evidence import PARTIAL_NOTICE


class Model:
    def __init__(self, *, repair=None, reject=False):
        self.repair = repair
        self.reject = reject
        self.repairs = 0
        self.checks = 0

    def judge_faithfulness(self, *, prompt, **kwargs):
        self.checks += 1
        data = json.loads(prompt[prompt.index('{"question"') :])
        return json.dumps(
            {
                "responsive": True,
                "missing_aspects": [],
                "claims": [
                    {
                        "claim_id": c["claim_id"],
                        "verdict": "unsupported"
                        if self.reject or "Unsupported" in c["text"]
                        else "supported",
                        "quotes": [
                            {"source_label": label, "quote": data["evidence"][label][0]}
                            for label in c["sources"]
                        ],
                    }
                    for c in data["claims"]
                ],
            }
        )

    def generate_json(self, **kwargs):
        self.repairs += 1
        assert kwargs["timeout_seconds"] > 0 and kwargs["json_schema"]
        assert "Original evidence" in kwargs["prompt"]
        return self.repair.model_dump_json() if self.repair else "malformed"


def manifest():
    return evidence_manifest(
        [
            evidence_record().model_copy(
                update={"index_generation_id": "generation-1", "content_hash": "a" * 64}
            )
        ]
    )


def check(bundle, *, model=None):
    return verify_bundle_citations(
        bundle,
        manifest(),
        inference=QueryRuntimeArtifactJsonGenerator(model or Model()),
        model="judge",
        deadline=perf_counter() + 10,
        check_cancelled=lambda: None,
    )


def test_inventory_covers_each_block_type_and_all_format_headings():
    bundle = content_bundle()
    blocks = [
        ContentBlock(kind="paragraph", text="Paragraph", evidence_ids=["E1"]),
        ContentBlock(kind="heading", text="Factual heading", evidence_ids=["E1"]),
        ContentBlock(kind="quotation", text="Quotation", evidence_ids=["E1"]),
        ContentBlock(kind="callout", text="Callout", evidence_ids=["E1"]),
        ContentBlock(
            kind="bullet_list",
            list_items=[ContentListItem(text="List claim", evidence_ids=["E1"])],
        ),
        ContentBlock(
            kind="key_value",
            entries=[KeyValueEntry(key="Key", value="Value", evidence_ids=["E1"])],
        ),
        ContentBlock(
            kind="table",
            table=ContentTable(
                headers=["Column"],
                rows=[ContentTableRow(values=["Cell"], evidence_ids=["E1"])],
            ),
        ),
    ]
    bundle.presentation.slides[0].blocks = blocks
    bundle.presentation.slides[0].title = "Factual slide title"
    bundle.paginated.subtitle = "Factual subtitle"
    texts = [c.text for c in bundle_claims(bundle, {"E1": "[doc:1]"})]
    for expected in (
        "Paragraph",
        "Factual heading",
        "Quotation",
        "Callout",
        "List claim",
        "Key: Value",
        '{"Column": "Cell"}',
        "Factual slide title",
        "Factual subtitle",
    ):
        assert expected in texts


def test_standard_notices_are_not_facts_but_substantive_additions_are_checked():
    bundle = content_bundle()
    bundle.content.warnings = [
        "Rendered from the previously displayed grounded answer.",
        "Rendered from the previously displayed grounded answer. Retention is 90 days.",
    ]
    bundle.content.sections[0].title = "Grounded answer"
    bundle.content.sections[0].blocks.append(
        ContentBlock(kind="paragraph", text=PARTIAL_NOTICE)
    )
    texts = [c.text for c in bundle_claims(bundle, {"E1": "[doc:1]"})]
    assert PARTIAL_NOTICE not in texts and "Grounded answer" not in texts
    assert bundle.content.warnings[0] not in texts
    assert bundle.content.warnings[1] in texts


def test_supported_partial_artifact_preserves_coverage_notice():
    bundle = content_bundle()
    for spec in (bundle.content, bundle.paginated, bundle.presentation):
        group = spec.slides[0] if hasattr(spec, "slides") else spec.sections[0]
        group.blocks.append(ContentBlock(kind="paragraph", text=PARTIAL_NOTICE))

    class PartialModel(Model):
        def judge_faithfulness(self, *, prompt, **kwargs):
            data = json.loads(super().judge_faithfulness(prompt=prompt, **kwargs))
            data["missing_aspects"] = ["Another requested aspect has no evidence."]
            return json.dumps(data)

    final, errors, _ = check(bundle, model=PartialModel())
    assert not errors
    assert final.citation_verification["coverage"] == "partial"
    assert PARTIAL_NOTICE in final.paginated.sections[0].blocks[-1].text

    bundle.presentation.slides[0].blocks[0].text = "Unsupported fact."
    assert check(bundle, model=PartialModel())[1]


@pytest.mark.parametrize("format_name", ["content", "paginated", "presentation"])
def test_unsupported_adaptation_blocks_whole_bundle(format_name):
    bundle = content_bundle()
    spec = getattr(bundle, format_name)
    group = spec.slides[0] if format_name == "presentation" else spec.sections[0]
    group.blocks = [
        ContentBlock(kind="paragraph", text="Unsupported fact.", evidence_ids=["E1"])
    ]
    checked, errors, _ = check(bundle)
    assert errors and checked.citation_verification is None


def test_missing_provenance_cannot_publish():
    evidence = manifest()
    evidence.records[0].provenance_kind = "unknown"
    checked, errors, calls = verify_bundle_citations(
        content_bundle(),
        evidence,
        inference=QueryRuntimeArtifactJsonGenerator(Model()),
        model="judge",
        deadline=perf_counter() + 10,
        check_cancelled=lambda: None,
    )
    assert errors and calls == 0 and checked.citation_verification is None


@pytest.mark.parametrize("succeeds", [True, False])
def test_executor_repairs_once_then_rechecks_all_formats(tmp_path, succeeds):
    jobs = InMemoryArtifactJobRepository()
    job = artifact_job(repo=jobs)
    artifacts = InMemoryGeneratedArtifactRepository()
    executor = seeded_executor(
        tmp_path=tmp_path, job_repo=jobs, artifact_repo=artifacts, job=job
    )
    executor.config = executor.config.model_copy(update={"ollama_num_ctx": 65536})
    candidate = content_bundle()
    candidate.presentation.slides[0].blocks = [
        ContentBlock(kind="paragraph", text="Unsupported fact.", evidence_ids=["E1"])
    ]
    model = Model(repair=content_bundle() if succeeds else candidate)
    executor.inference = QueryRuntimeArtifactJsonGenerator(model)
    if succeeds:
        final = executor._verify_citation_support(
            job, candidate, manifest(), document_plan()
        )
        assert final.citation_verification["status"] == "verified"
        assert final.citation_verification["repair_count"] == 1
        assert "Unsupported" not in final.model_dump_json()
    else:
        with pytest.raises(RuntimeError, match="artifact_citation_verification_failed"):
            executor._verify_citation_support(
                job, candidate, manifest(), document_plan()
            )
        assert not artifacts.list_artifacts_for_job(job.id)
    assert model.repairs == 1
    assert 2 <= model.checks <= 8


def test_cancellation_stops_before_model_call(tmp_path):
    jobs = InMemoryArtifactJobRepository()
    job = artifact_job(repo=jobs)
    executor = seeded_executor(
        tmp_path=tmp_path,
        job_repo=jobs,
        artifact_repo=InMemoryGeneratedArtifactRepository(),
        job=job,
    )
    executor._check_cancelled = Mock(side_effect=ArtifactJobCancelled("cancelled"))
    with pytest.raises(ArtifactJobCancelled):
        executor._verify_citation_support(
            job, content_bundle(), manifest(), document_plan()
        )


def test_changed_generation_blocks_storage(tmp_path):
    jobs = InMemoryArtifactJobRepository()
    job = artifact_job(repo=jobs)
    executor = seeded_executor(
        tmp_path=tmp_path,
        job_repo=jobs,
        artifact_repo=InMemoryGeneratedArtifactRepository(),
        job=job,
    )
    executor._citation_evidence = manifest()
    # Repository fixture has no matching current generation.
    with pytest.raises(ArtifactPermissionChanged, match="revision"):
        executor._guard_publication(job, ["doc-1"])


@pytest.mark.parametrize("format_name", ["pdf", "docx", "pptx"])
def test_export_contains_full_verified_quote_and_revision(format_name, monkeypatch):
    if format_name == "pdf":
        import reportlab.rl_config

        monkeypatch.setattr(reportlab.rl_config, "pageCompression", 0)
    bundle, errors, _ = check(content_bundle())
    assert errors == []
    rendered = render_document(
        artifact_format=format_name,
        bundle=bundle,
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
        require_libreoffice=False,
    )
    if format_name == "pdf":
        text = rendered.content.decode("latin1")
    elif format_name == "docx":
        from docx import Document

        text = "\n".join(p.text for p in Document(BytesIO(rendered.content)).paragraphs)
    else:
        from pptx import Presentation

        text = "\n".join(
            shape.text
            for slide in Presentation(BytesIO(rendered.content)).slides
            for shape in slide.shapes
            if shape.has_text_frame
        )
    assert "Evidence fact." in text
    assert "generation-1" in text
    assert "a" * 64 in text


def test_pptx_evidence_appendix_never_truncates_quotes_or_source_count():
    from pptx import Presentation

    bundle, _, _ = check(content_bundle())
    quote = "Long supporting evidence " * 24 + "FINAL_SENTINEL"
    bundle.content.citations = [
        bundle.content.citations[0].model_copy(
            update={
                "evidence_id": f"E{i}",
                "doc_title": f"Source {i}",
                "quotes": [quote],
            }
        )
        for i in range(21)
    ]
    rendered = render_document(
        artifact_format="pptx",
        bundle=bundle,
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
        require_libreoffice=False,
    )
    slides = Presentation(BytesIO(rendered.content)).slides
    text = "\n".join(
        shape.text for slide in slides for shape in slide.shapes if shape.has_text_frame
    )
    assert text.count("FINAL_SENTINEL") == 21
    assert "Source 20" in text
    assert all(
        shape.top + shape.height <= 7.5 * 914400
        for slide in slides
        for shape in slide.shapes
    )
