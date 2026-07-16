from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO

from pptx import Presentation

from rag.artifact_jobs.contracts import ContentBlock, PresentationSlide
from rag.artifact_jobs.renderer import render_document

from generation_optimization_support import content_bundle
from generation_rendering_support import (
    dynamic_table_bundle,
    large_list_bundle,
    presentation_table_text,
    presentation_tables,
    presentation_text,
    slide_background_hex,
    slide_text,
    timeline_bundle,
)


def test_deterministic_pptx_labels_timeline_profile() -> None:
    rendered = render_document(
        artifact_format="pptx",
        bundle=timeline_bundle(),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
        require_libreoffice=False,
    )

    presentation = Presentation(BytesIO(rendered.content))
    assert "TIMELINE REPORT" in presentation_text(presentation)
    assert slide_background_hex(presentation.slides[0]) == "1F3A5F"


def test_deterministic_pptx_table_geometry_uses_content_shape() -> None:
    rendered = render_document(
        artifact_format="pptx",
        bundle=dynamic_table_bundle(row_count=2),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
        require_libreoffice=False,
    )

    presentation = Presentation(BytesIO(rendered.content))
    table = presentation_tables(presentation)[0]
    widths = [column.width for column in table.columns]
    row_heights = [row.height for row in table.rows]

    assert widths[1] > widths[0] * 2
    assert widths[1] > widths[2]
    assert row_heights[1] > row_heights[2]


def test_deterministic_pptx_splits_tables_by_estimated_height() -> None:
    rendered = render_document(
        artifact_format="pptx",
        bundle=dynamic_table_bundle(row_count=8),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
        require_libreoffice=False,
    )

    presentation = Presentation(BytesIO(rendered.content))

    assert len(presentation_tables(presentation)) >= 2
    assert "Event 8" in presentation_table_text(presentation)


def test_deterministic_pptx_splits_large_lists_without_truncation() -> None:
    rendered = render_document(
        artifact_format="pptx",
        bundle=large_list_bundle(item_count=18),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
        require_libreoffice=False,
    )

    presentation = Presentation(BytesIO(rendered.content))
    text = presentation_text(presentation)
    assert "Item 18 evidence fact." in text
    assert len(presentation.slides) >= 4
    slides = list(presentation.slides)
    assert slide_background_hex(slides[0]) == "173B3F"
    content_backgrounds = [slide_background_hex(slide) for slide in slides[1:]]
    assert content_backgrounds[0] == "DCEBE8"
    assert len(set(content_backgrounds)) > 1


def test_deterministic_pptx_skips_generated_references_content_slide() -> None:
    bundle = content_bundle()
    placeholder_slide = PresentationSlide(
        title="References",
        blocks=[ContentBlock(kind="heading", text="References")],
    )
    bundle = bundle.model_copy(
        update={
            "presentation": bundle.presentation.model_copy(
                update={
                    "slides": [*bundle.presentation.slides, placeholder_slide],
                }
            )
        }
    )
    rendered = render_document(
        artifact_format="pptx",
        bundle=bundle,
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
        require_libreoffice=False,
    )

    presentation = Presentation(BytesIO(rendered.content))
    reference_slide_texts = [
        slide_text(slide)
        for slide in presentation.slides
        if "References" in slide_text(slide).splitlines()
    ]
    assert len(reference_slide_texts) == 1
    assert "FIR_02_kidnapping.pdf" in reference_slide_texts[0]
    assert reference_slide_texts[0].splitlines().count("References") == 1


def test_deterministic_pptx_reports_slide_progress_without_changing_content() -> None:
    events: list[tuple[int, int, str]] = []
    rendered = render_document(
        artifact_format="pptx",
        bundle=large_list_bundle(item_count=18),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
        require_libreoffice=False,
        progress_callback=lambda current, total, label: events.append(
            (current, total, label)
        ),
    )

    presentation = Presentation(BytesIO(rendered.content))
    assert events[0] == (1, len(presentation.slides), "Title slide")
    assert events[-1] == (
        len(presentation.slides),
        len(presentation.slides),
        "References",
    )
    assert len(events) == len(presentation.slides)
