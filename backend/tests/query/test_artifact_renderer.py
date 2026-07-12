from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
import zipfile

from pptx import Presentation

from rag.schemas.query import RAGResponse, SourceAnchor
from rag.query.artifact_models import (
    ArtifactCitation,
    ArtifactContent,
    ArtifactFieldDefinition,
    ArtifactTable,
    ArtifactTableRow,
    CoverageReport,
    SemanticSection,
    SupportedClaim,
)
from rag.query.artifact_renderer import _artifact_bundle, render_artifact


def test_query_pptx_uses_shared_adaptive_layout_renderer() -> None:
    rendered = render_artifact(
        artifact_format="pptx",
        title="Incident timeline",
        prompt="Create a timeline presentation.",
        artifact_content=_timeline_content(),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
    )

    presentation = Presentation(BytesIO(rendered.content))

    assert "TIMELINE REPORT" in _presentation_text(presentation)
    assert _slide_background_hex(presentation.slides[0]) == "1F3A5F"
    assert "References" in _presentation_text(presentation)


def test_query_docx_uses_shared_adaptive_layout_renderer() -> None:
    rendered = render_artifact(
        artifact_format="docx",
        title="Incident timeline",
        prompt="Create a timeline document.",
        artifact_content=_timeline_content(),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
    )

    with zipfile.ZipFile(BytesIO(rendered.content)) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        header_xml = "\n".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith("word/header")
        )

    assert "TIMELINE REPORT" in document_xml
    assert "Prudentia AI" in header_xml
    assert "w:tblGrid" in document_xml


def test_query_renderer_keeps_legacy_response_adapter_on_shared_renderer() -> None:
    rendered = render_artifact(
        artifact_format="pptx",
        title="Legacy Evidence Report",
        prompt="Create an evidence report.",
        response=_rag_response(),
        generated_at=datetime(2026, 1, 1, tzinfo=UTC),
    )

    presentation = Presentation(BytesIO(rendered.content))

    assert "EVIDENCE BRIEF" in _presentation_text(presentation)
    assert "Legacy Evidence Report" in _presentation_text(presentation)


def test_query_renderer_preserves_claim_level_evidence_ids() -> None:
    base = _timeline_content()
    content = ArtifactContent(
        title=base.title,
        purpose=base.purpose,
        sections=(
            SemanticSection(
                heading="Findings",
                paragraphs=("First finding.", "Second finding."),
                claim_ids=("claim-1", "claim-2"),
            ),
        ),
        claims=(
            SupportedClaim(claim_id="claim-1", text="First finding.", evidence_ids=("E1",)),
            SupportedClaim(claim_id="claim-2", text="Second finding.", evidence_ids=("E2",)),
        ),
        citations=(
            base.citations[0],
            ArtifactCitation(
                evidence_id="E2",
                doc_id="doc-2",
                doc_title="Second report.pdf",
                chunk_id="chunk-2",
                page_start=2,
                page_end=2,
            ),
        ),
        coverage=base.coverage,
    )

    bundle = _artifact_bundle(content)

    assert bundle.content.sections[0].blocks[0].evidence_ids == ["E1"]
    assert bundle.content.sections[0].blocks[1].evidence_ids == ["E2"]


def _timeline_content() -> ArtifactContent:
    citation = ArtifactCitation(
        evidence_id="E1",
        doc_id="doc-1",
        doc_title="Operation Arch Light.pdf",
        chunk_id="chunk-1",
        page_start=1,
        page_end=1,
    )
    claim = SupportedClaim(
        claim_id="claim-1",
        text="Complaint registered before evidence logging.",
        evidence_ids=("E1",),
    )
    table = ArtifactTable(
        title="Timeline",
        fields=(
            ArtifactFieldDefinition(name="time", label="Time"),
            ArtifactFieldDefinition(name="event", label="Event"),
        ),
        rows=(
            ArtifactTableRow(
                values=(("time", "09:00"), ("event", "Complaint registered.")),
                evidence_ids=("E1",),
            ),
        ),
    )
    return ArtifactContent(
        title="Incident timeline",
        purpose="Build a chronological timeline of events.",
        sections=(
            SemanticSection(
                heading="Timeline",
                bullets=("Complaint registered before evidence logging.",),
                tables=(table,),
                claim_ids=("claim-1",),
            ),
        ),
        claims=(claim,),
        citations=(citation,),
        coverage=CoverageReport(
            status="high_confidence",
            evidence_count=1,
            relevant_evidence_count=1,
            documents_searched=1,
            documents_expected=1,
        ),
    )


def _rag_response() -> RAGResponse:
    return RAGResponse(
        trace_id="trace-1",
        answer="The report contains one supported evidence statement.",
        sources=[
            SourceAnchor(
                doc_id="doc-1",
                doc_title="Operation Arch Light.pdf",
                chunk_id="chunk-1",
                page=1,
                page_start=1,
                page_end=1,
                excerpt="The report contains one supported evidence statement.",
                group_path="/",
            )
        ],
        artifacts=[],
        artifact_job=None,
        conflict_flag=False,
        conflict_detail=None,
        faithfulness_score=1.0,
        faithfulness_status="checked",
        unfounded_claims=[],
        intent="factual_simple",
        session_id="session-1",
        latency_ms=1,
        node_timings=[],
        degraded=False,
        degraded_reason=None,
    )


def _presentation_text(presentation) -> str:
    return "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )


def _slide_background_hex(slide) -> str:
    return str(slide.background.fill.fore_color.rgb)
