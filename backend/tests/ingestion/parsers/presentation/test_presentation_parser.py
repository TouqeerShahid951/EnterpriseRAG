from io import BytesIO
from dataclasses import replace

import pytest

from rag.ingestion.errors import WorkerStepError
from rag.ingestion.chunking import chunk_items
from rag.ingestion.parsers.images import ImageAnalysisResult, PNG_CONTENT_TYPE
from rag.ingestion.parsers.images.normalization import normalized_image_source
from rag.ingestion.parsers.models import ParsedPdfItem
from rag.ingestion.parsers.presentation import parser


def test_pptx_combines_structured_chart_and_required_slide_visual_chunks(monkeypatch) -> None:
    content = _presentation_with_chart()
    monkeypatch.setattr(parser, "parse_docling_pptx", lambda _content: [_docling_item()])
    monkeypatch.setattr(parser, "render_pptx_slides", lambda _content: [_slide_source()])
    analyzer = FakeAnalyzer()

    result = parser.parse_presentation_document(
        content,
        analyze_slide_visuals=True,
        image_context=("doc-1", FakeStore(), analyzer),
    )

    assert result.provenance["document_kind"] == "pptx"
    assert result.provenance["slide_count"] == 1
    assert any("2026" in item.text and item.item_type == "table" for item in result.items)
    assert any("flow from intake to approval" in item.text for item in result.items)
    assert all(item.page_start == 1 for item in result.items)
    assert any(asset.source_kind == "pptx_slide" for asset in result.assets)


def test_disabled_visual_analysis_stores_slide_without_calling_vision(monkeypatch) -> None:
    content = _presentation_with_chart()
    monkeypatch.setattr(parser, "parse_docling_pptx", lambda _content: [_docling_item()])
    monkeypatch.setattr(parser, "render_pptx_slides", lambda _content: [_slide_source()])
    analyzer = FakeAnalyzer()

    result = parser.parse_presentation_document(
        content,
        analyze_slide_visuals=False,
        image_context=("doc-1", FakeStore(), analyzer),
    )

    assert analyzer.slide_calls == 0
    assert result.items[0].image_source_kind == "pptx_slide"
    assert result.provenance["config"] == {"analyze_slide_visuals": False}


def test_required_slide_visual_analysis_fails_when_vision_returns_empty(monkeypatch) -> None:
    content = _presentation_with_chart()
    monkeypatch.setattr(parser, "parse_docling_pptx", lambda _content: [_docling_item()])
    monkeypatch.setattr(parser, "render_pptx_slides", lambda _content: [_slide_source()])

    with pytest.raises(WorkerStepError, match="every PPTX slide") as exc_info:
        parser.parse_presentation_document(
            content,
            analyze_slide_visuals=True,
            image_context=("doc-1", FakeStore(), EmptyAnalyzer()),
        )

    assert exc_info.value.code == "pptx_visual_analysis_failed"


def test_pptx_citations_use_rendered_slide_pixels_for_text_charts_and_pictures(monkeypatch) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation(BytesIO(_presentation_with_chart()))
    presentation.slides[0].shapes.add_picture(
        BytesIO(_slide_source().content), Inches(2), Inches(3), Inches(4), Inches(2),
    )
    output = BytesIO()
    presentation.save(output)
    # The shared Docling adapter has flipped the backend's mislabeled EMU Y axis.
    bbox = (float(Inches(1)), float(Inches(5.5)), float(Inches(4)), float(Inches(6.5)))
    parsed = replace(_docling_item(), text="Records must be retained.", item_type="text", bbox=bbox, source_spans=[{
        "page": 1, "bbox": list(bbox), "exact_text": "Records must be retained.",
    }])
    monkeypatch.setattr(parser, "parse_docling_pptx", lambda _content: [parsed])
    monkeypatch.setattr(parser, "render_pptx_slides", lambda _content: [_slide_source()])

    result = parser.parse_presentation_document(
        output.getvalue(), analyze_slide_visuals=True,
        image_context=("deck-1", FakeStore(), FakeAnalyzer()),
    )
    chunks = chunk_items(result.items, doc_id="deck-1", target_tokens=500, overlap_tokens=0, parent_max_tokens=1000)
    targets = [target for chunk in chunks for target in chunk.citation_targets]
    assert targets and all(target["kind"] == "slide" and target["page"] == 1 for target in targets)
    by_text = {target["exact_text"]: target for target in targets}
    assert by_text["Records must be retained."]["bbox"] == pytest.approx([32, 24, 128, 48])
    chart = next(item for item in result.items if "source:pptx_chart" in item.quality_flags)
    assert chart.bbox == pytest.approx([32, 36, 288, 132])
    picture = next(item for item in result.items if item.image_source_kind == "pptx_media")
    assert picture.bbox == pytest.approx([64, 72, 192, 120])
    picture_asset = next(asset for asset in result.assets if asset.source_kind == "pptx_media")
    assert picture_asset.bbox == picture.bbox
    full_slide = next(item for item in result.items if "full_slide_visual" in item.quality_flags)
    assert full_slide.bbox == (0, 0, 320, 180)

    asset = next(asset for asset in result.assets if asset.source_kind == "pptx_slide")
    with pytest.raises(WorkerStepError, match="dimensions are missing"):
        parser._attach_slide_assets(
            [replace(parsed, quality_flags=["source:pptx_structured"])],
            [replace(asset, width=0)], slide_size=(9144000, 6858000),
        )


class FakeStore:
    def put_image_asset(self, *, doc_id: str, asset_id: str, filename: str, content: bytes, content_type: str) -> str:
        return f"memory://{doc_id}/{asset_id}/{filename}"


class FakeAnalyzer:
    def __init__(self) -> None:
        self.slide_calls = 0

    def analyze_slide(self, *, content: bytes, content_type: str) -> ImageAnalysisResult:
        self.slide_calls += 1
        return ImageAnalysisResult(caption="The diagram shows a flow from intake to approval.", confidence=0.95)

    def analyze_image(self, *, content: bytes, content_type: str) -> ImageAnalysisResult:
        return ImageAnalysisResult(caption="Embedded illustration", confidence=0.9)


class EmptyAnalyzer(FakeAnalyzer):
    def analyze_slide(self, *, content: bytes, content_type: str) -> ImageAnalysisResult:
        return ImageAnalysisResult(quality_flags=["vision_empty_response"])


def _presentation_with_chart() -> bytes:
    from pptx import Presentation
    from pptx.chart.data import ChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Operational trend"
    data = ChartData()
    data.categories = ["2025", "2026"]
    data.add_series("Cases", (10, 14))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(8), Inches(4), data)
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def _docling_item() -> ParsedPdfItem:
    return ParsedPdfItem(
        index=0,
        text="Operational trend",
        item_type="heading",
        page_start=1,
        page_end=1,
        parser="docling",
    )


def _slide_source():
    from PIL import Image

    output = BytesIO()
    Image.new("RGB", (320, 180), "white").save(output, format="PNG")
    return normalized_image_source(
        output.getvalue(),
        filename="slide-1.png",
        source_kind="pptx_slide",
        quality_flags=["source:pptx_slide", "full_slide_visual"],
        page=1,
        bbox=(0.0, 0.0, 320.0, 180.0),
        page_area_ratio=1.0,
        target_content_type=PNG_CONTENT_TYPE,
    )


def test_xy_and_bubble_charts_keep_real_axes_and_all_plots():
    from types import SimpleNamespace
    from pptx import Presentation
    from pptx.chart.data import XyChartData, BubbleChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    xy = XyChartData()
    series = xy.add_series("Pressure")
    series.add_data_point(100, 5)
    series.add_data_point(200, 6)
    bubble = BubbleChartData()
    bubble.add_series("Volume").add_data_point(300, 7, 9)
    plots = []
    for data, kind in [(xy, XL_CHART_TYPE.XY_SCATTER), (bubble, XL_CHART_TYPE.BUBBLE)]:
        chart = slide.shapes.add_chart(kind, Inches(1), Inches(1), Inches(4), Inches(3), data).chart
        plots.extend(chart.plots)
    rows = parser._chart_rows(SimpleNamespace(plots=plots))
    assert rows[0] == ["Plot", "Series", "Category", "X", "Y", "Bubble size"]
    assert rows[1] == ["1", "Pressure", "", "100", "5.0", ""]
    assert rows[2] == ["1", "Pressure", "", "200", "6.0", ""]
    assert rows[3] == ["2", "Volume", "", "300", "7.0", "9"]
    point = plots[0].series[0]._element.xpath("./c:xVal/c:numRef/c:numCache/c:pt")[0]
    point.getparent().remove(point)
    with pytest.raises(WorkerStepError, match="axis values are missing"):
        parser._chart_rows(SimpleNamespace(plots=plots))
