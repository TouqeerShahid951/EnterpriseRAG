from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook

from rag.ingestion.chunking import chunk_items
from rag.ingestion.errors import WorkerStepError
from rag.ingestion.parsers.document import parse_document
from rag.ingestion.parsers.spreadsheet import CSV_CONTENT_TYPE, XLSX_CONTENT_TYPE


def test_csv_parser_emits_schema_and_structured_rows() -> None:
    parsed = parse_document(
        b"AC_ID,LOC_CD,STS\nA-101,ISB-01,OP\nA-102,RWP-02,MT\n",
        content_type=CSV_CONTENT_TYPE,
        file_path="memory://aircraft.csv",
        min_chars_per_page=10,
    )

    chunks = chunk_items(parsed.items, doc_id="doc-csv", target_tokens=128, overlap_tokens=0, parent_max_tokens=512)
    schema = next(chunk for chunk in chunks if "spreadsheet_schema" in chunk.quality_flags)
    rows = [chunk for chunk in chunks if chunk.chunk_type == "table_row"]

    assert parsed.provenance["document_kind"] == "csv"
    assert schema.table_title == "aircraft"
    assert "LOC_CD" in schema.text
    assert len(rows) == 2
    assert {"label": "LOC_CD", "value": "ISB-01"} in rows[0].structured_fields


def test_xlsx_parser_reads_visible_sheets_and_preserves_formulas() -> None:
    workbook = Workbook()
    aircraft = workbook.active
    aircraft.title = "Aircraft"
    aircraft.append(["AC_ID", "Inspection_Date", "Next_Code"])
    aircraft.append(["A-101", date(2026, 8, 26), '=A2&"-NEXT"'])
    hidden = workbook.create_sheet("Hidden")
    hidden.append(["secret"])
    hidden.sheet_state = "hidden"

    parsed = parse_document(
        _workbook_bytes(workbook),
        content_type=XLSX_CONTENT_TYPE,
        file_path="memory://aircraft.xlsx",
        min_chars_per_page=10,
    )

    text = "\n".join(item.text for item in parsed.items)
    assert parsed.provenance["document_kind"] == "xlsx"
    assert parsed.provenance["config"]["hidden_sheets_skipped"] == 1
    assert "Inspection_Date (date" in text
    assert '=A2&"-NEXT"' in text
    assert "secret" not in text


def test_spreadsheet_parser_rejects_more_than_the_row_limit() -> None:
    content = "header\n" + "\n".join(str(index) for index in range(10_001))

    with pytest.raises(WorkerStepError) as exc_info:
        parse_document(
            content.encode(),
            content_type=CSV_CONTENT_TYPE,
            file_path="memory://large.csv",
            min_chars_per_page=10,
        )

    assert exc_info.value.code == "spreadsheet_limit_exceeded"


def test_xlsx_review_combines_header_rows_and_applies_column_edits() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Capabilities"
    sheet.append(["Capability matrix"])
    sheet.append([])
    sheet.append(["Team", "Team", "Mission", "Mission"])
    sheet.append(["Code", "Name", "Sensor", "Notes"])
    sheet.append(["R-1", "Alpha", "Thermal", "Night"])
    sheet.append(["R-2", "Bravo", "Radar", "All weather"])

    parsed = parse_document(
        _workbook_bytes(workbook),
        content_type=XLSX_CONTENT_TYPE,
        file_path="memory://capabilities.xlsx",
        min_chars_per_page=10,
        spreadsheet_review={
            "sheets": [
                {
                    "source_name": "Capabilities",
                    "table_name": "Mission capabilities",
                    "header_rows": [3, 4],
                    "excluded": False,
                    "columns": [
                        {"index": 0, "name": "Team code", "description": "Stable team identifier", "aliases": ["unit code"]}
                    ],
                }
            ]
        },
    )

    schema = parsed.items[0]
    assert schema.section_title == "Mission capabilities"
    assert schema.table_json["rows"][0] == ["Team code", "Team / Name", "Mission / Sensor", "Mission / Notes"]
    assert "Stable team identifier" in schema.text
    assert "Capability matrix" not in "\n".join(item.text for item in parsed.items)


def test_xlsx_parser_detects_title_and_grouped_header_rows() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Report title"])
    sheet.append([])
    sheet.append(["Team", "Team", "Mission"])
    sheet.append(["Code", "Name", "Sensor", None])
    sheet.append(["R-1", "Alpha", "Thermal", "Night"])
    sheet.append(["R-2", "Bravo", "Radar", "All weather"])

    parsed = parse_document(
        _workbook_bytes(workbook),
        content_type=XLSX_CONTENT_TYPE,
        file_path="memory://grouped.xlsx",
        min_chars_per_page=10,
    )

    metadata = parsed.items[0].table_json["spreadsheet"]
    assert metadata["header_rows"] == [3, 4]
    assert parsed.items[0].table_json["rows"][0] == ["Team / Code", "Team / Name", "Mission / Sensor", "Mission"]


def _workbook_bytes(workbook: Workbook) -> bytes:
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_titled_csv_keeps_actual_headers_and_data():
    parsed = parse_document(b"Sales report\nProduct,Revenue\nWidget,100\nGadget,200\n",
        content_type=CSV_CONTENT_TYPE, file_path="sales.csv", min_chars_per_page=10)
    assert parsed.items[0].table_json["rows"][0] == ["Product", "Revenue"]
    assert parsed.items[1].table_json["rows"][1:] == [["Widget", "100"], ["Gadget", "200"]]
    assert "spreadsheet_header_inferred" in parsed.items[0].quality_flags


def test_xlsx_keeps_formats_raw_values_and_cached_formulas():
    import zipfile
    from xml.etree import ElementTree as ET

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["ID", "Rate", "Amount", "Calculation"])
    sheet.append([123, .25, 1200, "=1+1"])
    sheet["A2"].number_format = "000000"
    sheet["B2"].number_format = "0%"
    sheet["C2"].number_format = '"$"#,##0.00'
    content = BytesIO()
    with zipfile.ZipFile(BytesIO(_workbook_bytes(workbook))) as original, zipfile.ZipFile(content, "w") as modified:
        for entry in original.infolist():
            data = original.read(entry)
            if entry.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(data)
                ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                root.find('.//s:c[@r="D2"]/s:v', ns).text = "2"
                data = ET.tostring(root)
            modified.writestr(entry, data)
    parsed = parse_document(content.getvalue(), content_type=XLSX_CONTENT_TYPE, file_path="values.xlsx", min_chars_per_page=10)
    assert parsed.items[1].table_json["rows"][1] == ["000123", "25%", "$1,200.00", "=1+1 [cached result: 2]"]
    cells = {cell["cell"]: cell for cell in parsed.items[1].table_json["spreadsheet"]["cell_values"]}
    assert cells["A2"]["raw_value"] == "123"
    assert cells["B2"]["number_format"] == "0%"
    assert cells["D2"]["formula"] == "=1+1"
    assert cells["D2"]["cached_value"] == "2"
