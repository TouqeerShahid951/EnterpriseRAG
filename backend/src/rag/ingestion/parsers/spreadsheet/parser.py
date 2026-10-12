"""Deterministic CSV and XLSX parsing into the existing table model."""

from __future__ import annotations

import csv
from dataclasses import dataclass, field, replace
from datetime import date, datetime, time
from io import BytesIO, StringIO
from pathlib import PurePosixPath
from typing import Any, Iterable
import zipfile

from rag.ingestion.errors import WorkerStepError
from .cells import xlsx_rows
from rag.ingestion.parsers.models import DocumentParseResult, ParsedPdfItem
from rag.ingestion.parsers.provenance import base_report

CSV_CONTENT_TYPE = "text/csv"
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

MAX_WORKSHEETS = 20
MAX_ROWS = 10_000
MAX_COLUMNS = 256
MAX_CELLS = 500_000
MAX_CELL_CHARS = 10_000
MAX_XLSX_ENTRIES = 10_000
MAX_XLSX_UNCOMPRESSED_BYTES = 250 * 1024 * 1024
SAMPLE_ROWS = 5


@dataclass(frozen=True)
class SpreadsheetTable:
    source_name: str
    title: str
    headers: list[str]
    rows: list[list[str]]
    data_row_numbers: list[int]
    column_types: list[str]
    source: str
    header_rows: list[int]
    header_candidates: list[int]
    excluded: bool = False
    description: str = ""
    column_metadata: list[dict[str, object]] | None = None
    cell_metadata: list[dict[str, object]] = field(default_factory=list)
    quality_flags: list[str] = field(default_factory=list)


def parse_spreadsheet_document(
    file_bytes: bytes,
    *,
    content_type: str | None,
    file_path: str,
    spreadsheet_review: dict[str, Any] | None = None,
) -> DocumentParseResult:
    kind = _spreadsheet_kind(content_type, file_path)
    tables, skipped_sheets = (
        (_parse_csv(file_bytes, file_path, spreadsheet_review), 0)
        if kind == "csv"
        else _parse_xlsx(file_bytes, spreadsheet_review)
    )
    if not tables:
        raise WorkerStepError("spreadsheet_empty", "Spreadsheet does not contain any readable tables.")
    included_tables = [table for table in tables if not table.excluded]
    if not included_tables:
        raise WorkerStepError("spreadsheet_all_sheets_excluded", "Spreadsheet review must include at least one sheet.")
    _validate_table_totals(included_tables)

    items = _table_items(included_tables)
    return DocumentParseResult(
        items=items,
        provenance=base_report(
            document_kind=kind,
            page_count=None,
            primary_parser="csv" if kind == "csv" else "openpyxl",
            secondary_parser=None,
            routing_mode="spreadsheet_tables",
            config={
                "table_count": len(included_tables),
                "hidden_sheets_skipped": skipped_sheets,
                "header_detection": "shape_transition",
                "formula_mode": "source_expression",
            },
            items=items,
        ),
    )


def _spreadsheet_kind(content_type: str | None, file_path: str) -> str:
    normalized = (content_type or "").split(";", 1)[0].strip().lower()
    if normalized == CSV_CONTENT_TYPE or file_path.lower().endswith(".csv"):
        return "csv"
    if normalized == XLSX_CONTENT_TYPE or file_path.lower().endswith(".xlsx"):
        return "xlsx"
    raise WorkerStepError("unsupported_document_type", "Spreadsheet format is not supported for ingestion.")


def _parse_csv(
    file_bytes: bytes,
    file_path: str,
    spreadsheet_review: dict[str, Any] | None,
) -> list[SpreadsheetTable]:
    try:
        text = file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise WorkerStepError("csv_parse_failed", "CSV must be valid UTF-8.") from exc
    if "\x00" in text:
        raise WorkerStepError("csv_parse_failed", "CSV contains invalid null bytes.")
    try:
        dialect = csv.Sniffer().sniff(text[:65_536], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    try:
        rows = list(csv.reader(StringIO(text), dialect, strict=True))
    except csv.Error as exc:
        raise WorkerStepError("csv_parse_failed", "CSV could not be parsed.") from exc
    title = PurePosixPath(file_path).stem or "CSV"
    table = _table_from_rows(
        title,
        rows,
        source="csv",
        config=_sheet_config(spreadsheet_review, title),
    )
    return [table] if table is not None else []


def _parse_xlsx(
    file_bytes: bytes,
    spreadsheet_review: dict[str, Any] | None,
) -> tuple[list[SpreadsheetTable], int]:
    _validate_xlsx_archive(file_bytes)
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(
            BytesIO(file_bytes),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as exc:
        raise WorkerStepError("xlsx_parse_failed", "XLSX workbook could not be parsed.") from exc

    tables: list[SpreadsheetTable] = []
    skipped_sheets = 0
    cached_workbook = None
    try:
        cached_workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True, keep_links=False)
        if len(workbook.worksheets) > MAX_WORKSHEETS:
            _raise_limit(f"XLSX cannot contain more than {MAX_WORKSHEETS} worksheets.")
        for worksheet in workbook.worksheets:
            if worksheet.sheet_state != "visible":
                skipped_sheets += 1
                continue
            cell_metadata: list[dict[str, object]] = []
            table = _table_from_rows(
                worksheet.title,
                xlsx_rows(worksheet, cached_workbook[worksheet.title], cell_metadata),
                source="xlsx",
                config=_sheet_config(spreadsheet_review, worksheet.title),
            )
            if table is not None:
                tables.append(replace(table, cell_metadata=cell_metadata))
    finally:
        workbook.close()
        if cached_workbook is not None:
            cached_workbook.close()
    return tables, skipped_sheets


def _validate_xlsx_archive(file_bytes: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(file_bytes)) as archive:
            entries = archive.infolist()
            names = {entry.filename for entry in entries}
            if len(entries) > MAX_XLSX_ENTRIES:
                _raise_limit("XLSX archive contains too many files.")
            if sum(entry.file_size for entry in entries) > MAX_XLSX_UNCOMPRESSED_BYTES:
                _raise_limit("XLSX uncompressed content exceeds 250 MB.")
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise WorkerStepError("xlsx_parse_failed", "Uploaded file is not a valid XLSX workbook.")
            if archive.testzip() is not None:
                raise WorkerStepError("xlsx_parse_failed", "XLSX workbook contains corrupt data.")
    except zipfile.BadZipFile as exc:
        raise WorkerStepError("xlsx_parse_failed", "Uploaded file is not a valid XLSX workbook.") from exc


def _table_from_rows(
    title: str,
    source_rows: Iterable[Iterable[Any]],
    *,
    source: str,
    config: dict[str, Any] | None,
) -> SpreadsheetTable | None:
    rows: list[tuple[int, list[str]]] = []
    for row_number, source_row in enumerate(source_rows, start=1):
        row = [_cell_text(value) for value in source_row]
        while row and not row[-1]:
            row.pop()
        if not any(row):
            continue
        if len(row) > MAX_COLUMNS:
            _raise_limit(f"Spreadsheet table cannot contain more than {MAX_COLUMNS} columns.")
        rows.append((row_number, row))
        if len(rows) > MAX_ROWS + 3:
            _raise_limit(f"Spreadsheet cannot contain more than {MAX_ROWS} data rows.")
    if not rows:
        return None

    width = max(len(row) for _, row in rows)
    detected_header_rows = _detected_header_rows(rows, width)
    header_rows = _configured_header_rows(config, rows) or detected_header_rows
    raw_headers = _composed_headers(rows, header_rows, width)
    headers = _normalized_headers(raw_headers, width)
    column_metadata = _column_metadata(config, headers)
    headers = [str(column["name"]) for column in column_metadata]
    last_header_row = max(header_rows)
    data_entries = [(number, _padded(row, width)) for number, row in rows if number > last_header_row]
    return SpreadsheetTable(
        source_name=" ".join(title.split()) or "Sheet",
        title=_configured_table_name(config, title),
        headers=headers,
        rows=[row for _, row in data_entries],
        data_row_numbers=[number for number, _ in data_entries],
        column_types=_column_types([row for _, row in data_entries], width),
        source=source,
        header_rows=header_rows,
        header_candidates=[number for number, _ in rows[:10]],
        excluded=bool(config and config.get("excluded")),
        description=_clean_text(config.get("description") if config else "", 500),
        column_metadata=column_metadata,
        quality_flags=[] if config and config.get("header_rows") else ["spreadsheet_header_inferred"],
    )


def _detected_header_rows(rows: list[tuple[int, list[str]]], width: int) -> list[int]:
    # A sparse title is not a header. Choose the first substantial row, then
    # combine an explicit repeated group row with the following header row.
    candidates = rows[:20]
    for index, (number, row) in enumerate(candidates):
        filled = [value for value in row if value]
        if width > 1 and len(filled) <= 1:
            continue
        if len(filled) < max(1, width // 2):
            continue
        if index + 1 < len(candidates) and len(set(filled)) < len(filled):
            next_number, next_row = candidates[index + 1]
            if all(_value_type(value) == "text" for value in next_row if value):
                return [number, next_number]
        return [number]
    return [rows[0][0]]


def _configured_header_rows(
    config: dict[str, Any] | None,
    rows: list[tuple[int, list[str]]],
) -> list[int]:
    if not config or not isinstance(config.get("header_rows"), list):
        return []
    available = {number for number, _ in rows}
    selected = sorted({value for value in config["header_rows"] if isinstance(value, int) and not isinstance(value, bool)})
    if not selected or len(selected) > 3 or any(value not in available for value in selected):
        raise WorkerStepError("spreadsheet_review_invalid", "Spreadsheet header rows are invalid for the source sheet.")
    return selected


def _composed_headers(
    rows: list[tuple[int, list[str]]],
    header_rows: list[int],
    width: int,
) -> list[str]:
    by_number = {number: _padded(row, width) for number, row in rows}
    levels: list[list[str]] = []
    for position, number in enumerate(header_rows):
        level = by_number[number]
        if position < len(header_rows) - 1:
            carried = ""
            level = [carried := value or carried for value in level]
        levels.append(level)
    return [
        " / ".join(dict.fromkeys(value for level in levels if (value := level[index])))
        for index in range(width)
    ]


def _sheet_config(review: dict[str, Any] | None, source_name: str) -> dict[str, Any] | None:
    sheets = review.get("sheets") if isinstance(review, dict) else None
    if not isinstance(sheets, list):
        return None
    return next(
        (
            dict(sheet)
            for sheet in sheets
            if isinstance(sheet, dict) and sheet.get("source_name") == source_name
        ),
        None,
    )


def _configured_table_name(config: dict[str, Any] | None, source_name: str) -> str:
    configured = _clean_text(config.get("table_name") if config else "", 200)
    return configured or " ".join(source_name.split()) or "Sheet"


def _column_metadata(config: dict[str, Any] | None, headers: list[str]) -> list[dict[str, object]]:
    configured = config.get("columns") if config else None
    by_index = {
        column["index"]: column
        for column in configured
        if isinstance(configured, list)
        and isinstance(column, dict)
        and isinstance(column.get("index"), int)
        and not isinstance(column.get("index"), bool)
    } if isinstance(configured, list) else {}
    result: list[dict[str, object]] = []
    for index, header in enumerate(headers):
        override = by_index.get(index, {})
        name = _clean_text(override.get("name"), 200) or header
        aliases = override.get("aliases")
        result.append(
            {
                "index": index,
                "name": name,
                "description": _clean_text(override.get("description"), 500),
                "aliases": [
                    text
                    for value in aliases[:5]
                    if (text := _clean_text(value, 80))
                ] if isinstance(aliases, list) else [],
            }
        )
    return result


def _clean_text(value: object, limit: int) -> str:
    return " ".join(str(value or "").split())[:limit].rstrip()


def _normalized_headers(raw_headers: list[str], width: int) -> list[str]:
    headers: list[str] = []
    counts: dict[str, int] = {}
    for index in range(width):
        base = raw_headers[index].strip() if index < len(raw_headers) else ""
        base = base or f"Column {index + 1}"
        counts[base] = counts.get(base, 0) + 1
        headers.append(base if counts[base] == 1 else f"{base} ({counts[base]})")
    return headers


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        text = value.isoformat(sep=" ")
    elif isinstance(value, (date, time)):
        text = value.isoformat()
    elif isinstance(value, bool):
        text = "true" if value else "false"
    else:
        text = str(value)
    text = " ".join(text.split())
    if len(text) > MAX_CELL_CHARS:
        _raise_limit(f"Spreadsheet cells cannot exceed {MAX_CELL_CHARS} characters.")
    return text


def _column_types(rows: list[list[str]], width: int) -> list[str]:
    return [_inferred_type(row[index] for row in rows if index < len(row)) for index in range(width)]


def _inferred_type(values: Iterable[str]) -> str:
    kinds = {_value_type(value) for value in values if value}
    if not kinds:
        return "empty"
    if len(kinds) == 1:
        return kinds.pop()
    if kinds <= {"integer", "number"}:
        return "number"
    return "mixed"


def _value_type(value: str) -> str:
    if value.startswith("="):
        return "formula"
    if value.lower() in {"true", "false"}:
        return "boolean"
    try:
        datetime.fromisoformat(value)
        return "datetime" if "T" in value or " " in value else "date"
    except ValueError:
        pass
    try:
        int(value)
        return "integer"
    except ValueError:
        pass
    try:
        float(value)
        return "number"
    except ValueError:
        return "text"


def _table_items(tables: list[SpreadsheetTable]) -> list[ParsedPdfItem]:
    items: list[ParsedPdfItem] = []
    for table in tables:
        common = {
            "page_start": None,
            "page_end": None,
            "section_title": table.title,
            "section_path": [table.title],
            "parser": table.source,
        }
        items.append(
            ParsedPdfItem(
                index=len(items),
                text=_schema_text(table),
                item_type="table",
                quality_flags=[f"source:{table.source}", "spreadsheet_schema", *table.quality_flags],
                table_json={
                    "caption": table.title,
                    "rows": [table.headers],
                    "spreadsheet": {
                        "source_name": table.source_name,
                        "header_rows": table.header_rows,
                        "header_candidates": table.header_candidates,
                        "data_row_numbers": table.data_row_numbers,
                        "column_count": len(table.headers),
                        "description": table.description,
                        "columns": [
                            {**column, "type": table.column_types[index]}
                            for index, column in enumerate(table.column_metadata or [])
                        ],
                    },
                },
                **common,
            )
        )
        if table.rows:
            items.append(
                ParsedPdfItem(
                    index=len(items),
                    text=_markdown_table(table.headers, table.rows),
                    item_type="table",
                    quality_flags=[f"source:{table.source}", "spreadsheet_table", *table.quality_flags],
                    table_json={
                        "caption": table.title,
                        "rows": [table.headers, *table.rows],
                        "spreadsheet": {
                            "source_name": table.source_name,
                            "cell_values": table.cell_metadata,
                            "header_rows": table.header_rows,
                            "data_row_numbers": table.data_row_numbers,
                            "column_count": len(table.headers),
                        },
                    },
                    **common,
                )
            )
    return items


def _schema_text(table: SpreadsheetTable) -> str:
    samples = _sample_rows(table.rows)
    lines = [f"Spreadsheet table: {table.title}", "Columns:"]
    if table.description:
        lines.insert(1, f"Table description: {table.description}")
    for index, (header, column_type) in enumerate(zip(table.headers, table.column_types, strict=True)):
        examples = [row[index] for row in samples if index < len(row) and row[index]][:3]
        suffix = f"; examples: {' | '.join(examples)}" if examples else ""
        metadata = (table.column_metadata or [])[index]
        description = str(metadata.get("description") or "")
        aliases = metadata.get("aliases")
        semantic = [description] if description else []
        if isinstance(aliases, list) and aliases:
            semantic.append(f"aliases: {', '.join(str(alias) for alias in aliases)}")
        detail = f" — {'; '.join(semantic)}" if semantic else ""
        lines.append(f"- {header} ({column_type}{suffix}){detail}")
    return "\n".join(lines)


def _sample_rows(rows: list[list[str]]) -> list[list[str]]:
    if len(rows) <= SAMPLE_ROWS:
        return rows
    indexes = sorted({round(index * (len(rows) - 1) / (SAMPLE_ROWS - 1)) for index in range(SAMPLE_ROWS)})
    return [rows[index] for index in indexes]


def _markdown_table(headers: list[str], rows: list[list[str]]) -> str:
    escaped_headers = [_markdown_cell(value) for value in headers]
    lines = [
        "| " + " | ".join(escaped_headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    lines.extend("| " + " | ".join(_markdown_cell(value) for value in row) + " |" for row in rows)
    return "\n".join(lines)


def _markdown_cell(value: str) -> str:
    return value.replace("|", "\\|")


def _padded(row: list[str], width: int) -> list[str]:
    return [*row, *([""] * (width - len(row)))]


def _validate_table_totals(tables: list[SpreadsheetTable]) -> None:
    total_rows = sum(len(table.rows) for table in tables)
    total_cells = sum(len(table.headers) * (len(table.rows) + 1) for table in tables)
    if total_rows > MAX_ROWS:
        _raise_limit(f"Spreadsheet cannot contain more than {MAX_ROWS} data rows.")
    if total_cells > MAX_CELLS:
        _raise_limit(f"Spreadsheet cannot contain more than {MAX_CELLS} cells.")


def _raise_limit(message: str) -> None:
    raise WorkerStepError("spreadsheet_limit_exceeded", message)
