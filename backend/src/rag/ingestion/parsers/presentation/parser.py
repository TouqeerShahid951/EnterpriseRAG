"""PPTX parsing, native asset extraction, and slide-level visual analysis."""

from __future__ import annotations

from collections.abc import Callable
from dataclasses import replace
from io import BytesIO
from pathlib import Path
import math
import shutil
import subprocess
import tempfile

from rag.ingestion.errors import UnsupportedDocumentError, WorkerStepError
from rag.ingestion.parsers.docling.adapter import parse_docling_pptx
from rag.ingestion.parsers.hierarchy import section_id
from rag.ingestion.parsers.images import (
    ImageAnalyzer,
    ImageAssetStore,
    ImageProgressCallback,
    ImageSource,
    JPEG_CONTENT_TYPE,
    PNG_CONTENT_TYPE,
    image_sources_to_items,
)
from rag.ingestion.parsers.images.normalization import normalized_image_source
from rag.ingestion.parsers.models import BBox, DocumentParseResult, ParsedImageAsset, ParsedPdfItem
from rag.ingestion.parsers.pdf.image_processing import renumber_items
from rag.ingestion.parsers.provenance import base_report

PPTX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
ImageItemsFactory = Callable[..., tuple[list[ParsedPdfItem], list[ParsedImageAsset]]]


def parse_presentation_document(
    file_bytes: bytes,
    *,
    analyze_slide_visuals: bool,
    image_context: tuple[str, ImageAssetStore, ImageAnalyzer] | None = None,
    image_items_factory: ImageItemsFactory = image_sources_to_items,
    image_progress_callback: ImageProgressCallback | None = None,
) -> DocumentParseResult:
    presentation, slide_titles = _presentation(file_bytes)
    slides = presentation.slides
    slide_size = (presentation.slide_width, presentation.slide_height)
    items = _docling_items(file_bytes, slide_titles)
    items.extend(_chart_items(slides, slide_titles, start_index=len(items)))
    assets: list[ParsedImageAsset] = []

    if image_context is not None:
        rendered_sources = render_pptx_slides(file_bytes)
        if len(rendered_sources) != len(slides):
            raise WorkerStepError("pptx_render_incomplete", "Rendered PPTX slide count did not match the presentation.")
        visual_items, slide_assets = image_items_factory(
            rendered_sources,
            doc_id=image_context[0],
            store=image_context[1],
            analyzer=_SlideAnalyzer(image_context[2]),
            start_index=len(items),
            progress_callback=image_progress_callback,
            analyze=analyze_slide_visuals,
        )
        assets.extend(slide_assets)
        if analyze_slide_visuals and len(visual_items) != len(slides):
            raise WorkerStepError("pptx_visual_analysis_failed", "Visual analysis did not produce a result for every PPTX slide.")
        items.extend(visual_items)

        media_items, media_assets = image_items_factory(
            _picture_sources(slides, slide_size=slide_size, slide_assets=slide_assets),
            doc_id=image_context[0],
            store=image_context[1],
            analyzer=image_context[2],
            start_index=len(items),
            progress_callback=image_progress_callback,
            analyze=analyze_slide_visuals,
        )
        items.extend(media_items)
        assets.extend(media_assets)
        items = _attach_slide_assets(items, slide_assets, slide_size=slide_size)
    elif analyze_slide_visuals:
        raise WorkerStepError("pptx_visual_analysis_unavailable", "PPTX slide visual analysis is not configured.")

    items = renumber_items(sorted(items, key=_item_order))
    if not items or not "".join(item.text for item in items).strip():
        raise UnsupportedDocumentError()
    return DocumentParseResult(
        items=items,
        provenance=base_report(
            document_kind="pptx",
            page_count=len(slides),
            primary_parser="docling",
            secondary_parser="vision" if analyze_slide_visuals else "python-pptx",
            routing_mode="docling_pptx_with_slide_visuals" if analyze_slide_visuals else "docling_pptx",
            config={"analyze_slide_visuals": analyze_slide_visuals},
            items=items,
        )
        | {"image_asset_count": len(assets), "slide_count": len(slides)},
        assets=assets,
    )


def render_pptx_slides(file_bytes: bytes) -> list[ImageSource]:
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if executable is None:
        raise WorkerStepError("pptx_renderer_missing", "LibreOffice is required to render PPTX slides.")
    try:
        with tempfile.TemporaryDirectory(prefix="pptx-render-") as temp_dir:
            source = Path(temp_dir) / "presentation.pptx"
            source.write_bytes(file_bytes)
            result = subprocess.run(
                [executable, "--headless", "--convert-to", "pdf", "--outdir", temp_dir, str(source)],
                capture_output=True,
                text=True,
                timeout=180,
                check=False,
            )
            pdf_path = Path(temp_dir) / "presentation.pdf"
            if result.returncode != 0 or not pdf_path.exists():
                raise WorkerStepError("pptx_render_failed", "LibreOffice failed to render the PPTX presentation.")
            return _pdf_slide_sources(pdf_path.read_bytes())
    except subprocess.TimeoutExpired as exc:
        raise WorkerStepError("pptx_render_timeout", "PPTX slide rendering exceeded 180 seconds.") from exc


def _pdf_slide_sources(pdf_bytes: bytes) -> list[ImageSource]:
    try:
        import fitz
    except ImportError as exc:
        raise WorkerStepError("pptx_render_dependency_missing", "PyMuPDF is required to rasterize PPTX slides.") from exc
    sources: list[ImageSource] = []
    try:
        with fitz.open(stream=pdf_bytes, filetype="pdf") as document:
            for index, page in enumerate(document, start=1):
                pixmap = page.get_pixmap(dpi=150, alpha=False)
                sources.append(
                    normalized_image_source(
                        pixmap.tobytes("png"),
                        filename=f"slide-{index}.png",
                        source_kind="pptx_slide",
                        quality_flags=["source:pptx_slide", "full_slide_visual"],
                        page=index,
                        bbox=(0.0, 0.0, float(pixmap.width), float(pixmap.height)),
                        page_area_ratio=1.0,
                        target_content_type=PNG_CONTENT_TYPE,
                    )
                )
    except WorkerStepError:
        raise
    except Exception as exc:
        raise WorkerStepError("pptx_render_failed", "Rendered PPTX slides could not be read.") from exc
    return sources


def _presentation(file_bytes: bytes):
    try:
        from pptx import Presentation

        presentation = Presentation(BytesIO(file_bytes))
    except ImportError as exc:
        raise WorkerStepError("pptx_dependency_missing", "python-pptx is not installed.") from exc
    except Exception as exc:
        raise WorkerStepError("pptx_parse_failed", "The PPTX package could not be opened.") from exc
    if not presentation.slides:
        raise UnsupportedDocumentError()
    titles = {index: _slide_title(slide, index) for index, slide in enumerate(presentation.slides, start=1)}
    return presentation, titles


def _docling_items(file_bytes: bytes, titles: dict[int, str]) -> list[ParsedPdfItem]:
    try:
        parsed = parse_docling_pptx(file_bytes)
    except ImportError as exc:
        raise WorkerStepError("pptx_dependency_missing", "Docling PPTX support is not installed.") from exc
    except Exception as exc:
        raise WorkerStepError("pptx_parse_failed", "Docling failed to parse PPTX.") from exc
    items: list[ParsedPdfItem] = []
    for item in parsed:
        slide = item.page_start
        if slide not in titles:
            continue
        title = titles[slide]
        path = [title]
        items.append(
            replace(
                item,
                page_end=slide,
                section_title=title,
                section_path=path,
                parent_section_id=section_id([f"Slide {slide}", title]),
                quality_flags=sorted({*item.quality_flags, "source:pptx_structured"}),
            )
        )
    return items


def _chart_items(slides, titles: dict[int, str], *, start_index: int) -> list[ParsedPdfItem]:
    items: list[ParsedPdfItem] = []
    for slide_no, slide in enumerate(slides, start=1):
        for chart_no, shape in enumerate((shape for shape in slide.shapes if getattr(shape, "has_chart", False)), start=1):
            rows = _chart_rows(shape.chart)
            if len(rows) < 2:
                continue
            chart_title = _chart_title(shape.chart) or f"Chart {chart_no}"
            text = "\n".join(" | ".join(str(value) for value in row) for row in rows)
            items.append(
                ParsedPdfItem(
                    index=start_index + len(items),
                    text=f"{chart_title}\n{text}",
                    item_type="table",
                    page_start=slide_no,
                    page_end=slide_no,
                    bbox=_shape_bbox(shape),
                    section_title=titles[slide_no],
                    section_path=[titles[slide_no]],
                    parent_section_id=section_id([f"Slide {slide_no}", titles[slide_no]]),
                    parser="python-pptx",
                    quality_flags=["source:pptx_chart", "structured_chart_data"],
                    table_json={"rows": rows, "row_count": len(rows) - 1, "caption": chart_title, "structured_kind": "chart"},
                )
            )
    return items


def _chart_rows(chart) -> list[list[str]]:
    from pptx.chart.series import BubbleSeries, XySeries

    rows = [["Plot", "Series", "Category", "X", "Y", "Bubble size"]]
    for plot_index, plot in enumerate(chart.plots, 1):
        for series_index, series in enumerate(plot.series, 1):
            name = str(series.name or f"Series {series_index}")
            values = series.values
            if isinstance(series, XySeries):
                xs = _axis_values(series, "xVal")
                sizes = _axis_values(series, "bubbleSize") if isinstance(series, BubbleSeries) else {}
                expected = set(range(len(values)))
                if set(xs) != expected or (isinstance(series, BubbleSeries) and set(sizes) != expected):
                    raise WorkerStepError("pptx_chart_incomplete", "Chart axis values are missing; repair the chart cache before ingestion.")
                for index, y in enumerate(values):
                    rows.append([str(plot_index), name, "", xs[index], "" if y is None else str(y), sizes.get(index, "")])
            else:
                categories = [str(category) for category in getattr(plot, "categories", [])]
                if len(categories) != len(values):
                    raise WorkerStepError("pptx_chart_unsupported", "Chart categories do not match its values; explicit source repair is required.")
                for category, value in zip(categories, values):
                    rows.append([str(plot_index), name, category, "", "" if value is None else str(value), ""])
    return rows


def _axis_values(series, axis: str) -> dict[int, str]:
    # python-pptx exposes Y values only; read X/bubble caches from the same OOXML series.
    points = series._element.xpath(f"./c:{axis}/c:numRef/c:numCache/c:pt | ./c:{axis}/c:numLit/c:pt")
    values: dict[int, str] = {}
    for point in points:
        value = point.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}v")
        if value is None or value.text is None:
            raise WorkerStepError("pptx_chart_incomplete", "Chart contains an empty axis value.")
        try:
            index = int(point.get("idx"))
            valid_value = math.isfinite(float(value.text))
        except (TypeError, ValueError) as exc:
            raise WorkerStepError("pptx_chart_incomplete", "Chart contains an invalid numeric axis value.") from exc
        if not valid_value or index < 0 or index in values:
            raise WorkerStepError("pptx_chart_incomplete", "Chart contains invalid axis indexes.")
        values[index] = value.text
    return values


def _chart_title(chart) -> str:
    if not getattr(chart, "has_title", False):
        return ""
    return " ".join(str(getattr(chart.chart_title.text_frame, "text", "")).split())


def _picture_sources(slides, *, slide_size: tuple[int, int], slide_assets: list[ParsedImageAsset]) -> list[ImageSource]:
    try:
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        return []
    sources: list[ImageSource] = []
    by_slide = {asset.page: asset for asset in slide_assets}
    for slide_no, slide in enumerate(slides, start=1):
        for image_no, shape in enumerate((shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE), start=1):
            image = shape.image
            content_type = str(image.content_type)
            if content_type not in {JPEG_CONTENT_TYPE, PNG_CONTENT_TYPE}:
                continue
            bbox = _slide_pixel_bbox(_shape_bbox(shape), slide_size, by_slide[slide_no])
            try:
                sources.append(
                    normalized_image_source(
                        image.blob,
                        filename=f"slide-{slide_no}-image-{image_no}.{image.ext}",
                        source_kind="pptx_media",
                        quality_flags=["source:pptx_media"],
                        page=slide_no,
                        bbox=bbox,
                        target_content_type=content_type,
                    )
                )
            except WorkerStepError:
                continue
    return sources


def _attach_slide_assets(
    items: list[ParsedPdfItem], assets: list[ParsedImageAsset], *, slide_size: tuple[int, int],
) -> list[ParsedPdfItem]:
    by_slide = {asset.page: asset for asset in assets if asset.source_kind == "pptx_slide" and asset.page is not None}
    located: list[ParsedPdfItem] = []
    for item in items:
        asset = by_slide.get(item.page_start)
        if asset is None:
            raise WorkerStepError("pptx_source_location_missing", "A PPTX item has no rendered source slide.")
        bbox, spans = item.bbox, item.source_spans
        if {"source:pptx_structured", "source:pptx_chart"}.intersection(item.quality_flags):
            docling = "source:pptx_structured" in item.quality_flags
            bbox = _slide_pixel_bbox(bbox, slide_size, asset, docling=docling)
            spans = [
                {**span, "bbox": list(box)}
                for span in spans
                if (box := _slide_pixel_bbox(span["bbox"], slide_size, asset, docling=docling)) is not None
            ]
        located.append(replace(
            item,
            bbox=bbox,
            source_spans=spans,
            image_asset_id=item.image_asset_id or asset.id,
            image_source_kind=item.image_source_kind or "pptx_slide",
            image_content_type=item.image_content_type or asset.content_type,
        ))
    return located


def _shape_bbox(shape) -> BBox:
    return (float(shape.left), float(shape.top), float(shape.left + shape.width), float(shape.top + shape.height))


def _slide_pixel_bbox(
    bbox: BBox | None, slide_size: tuple[int, int], asset: ParsedImageAsset, *, docling: bool = False,
) -> BBox | None:
    if bbox is None:
        return None
    width, height = slide_size
    if width <= 0 or height <= 0 or (asset.width or 0) <= 0 or (asset.height or 0) <= 0:
        raise WorkerStepError("pptx_source_location_missing", "PPTX slide dimensions are missing.")
    left, top, right, bottom = bbox
    if not all(math.isfinite(value) for value in bbox):
        raise WorkerStepError("pptx_source_location_missing", "A PPTX source region has invalid coordinates.")
    if docling:
        # Docling 2.97 labels native top-left EMUs as bottom-left; undo the
        # shared adapter's Y flip before scaling to the rendered slide pixels.
        top, bottom = height - bottom, height - top
    if right <= left or bottom <= top:
        return None
    return (left * asset.width / width, top * asset.height / height,
            right * asset.width / width, bottom * asset.height / height)


def _slide_title(slide, slide_no: int) -> str:
    title_shape = slide.shapes.title
    title = " ".join(str(getattr(title_shape, "text", "")).split()) if title_shape is not None else ""
    return title or f"Slide {slide_no}"


def _item_order(item: ParsedPdfItem) -> tuple[int, int, int]:
    priority = {"heading": 0, "text": 1, "table": 2, "image_text": 3}
    return (item.page_start or 0, priority.get(item.item_type, 4), item.index)


class _SlideAnalyzer:
    def __init__(self, analyzer: ImageAnalyzer) -> None:
        self._analyzer = analyzer

    def analyze_image(self, *, content: bytes, content_type: str):
        return self._analyzer.analyze_slide(content=content, content_type=content_type)
