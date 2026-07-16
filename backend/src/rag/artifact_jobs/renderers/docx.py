"""Render validated artifact bundles as DOCX packages."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from ..contracts import ArtifactContentBundle, ContentBlock, EvidenceCitation
from rag.artifact_jobs.generation.layout_profiles import ArtifactLayoutProfile
from .shared import _list_items, _reference_line, _source_summary, _with_sources


def render_docx(bundle: ArtifactContentBundle, generated_at: datetime, profile: ArtifactLayoutProfile) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches

    document = Document()
    section = document.sections[0]
    if _has_wide_table(bundle):
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.85 if _has_wide_table(bundle) else 1.0)
    section.right_margin = Inches(0.85 if _has_wide_table(bundle) else 1.0)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)
    _apply_docx_style_defaults(document, profile)
    _set_docx_page_chrome(section, profile)
    _add_docx_opening_block(document, bundle, generated_at, profile)

    citations = {item.evidence_id: item for item in bundle.content.citations}
    if bundle.paginated.include_coverage_notes and bundle.content.warnings:
        document.add_heading("Coverage Notes", level=1)
        for warning in bundle.content.warnings:
            paragraph = document.add_paragraph(warning, style="List Bullet")
            _format_docx_paragraph_runs(paragraph, size=10.5, color=profile.text_hex)
    for semantic_section in bundle.paginated.sections:
        document.add_heading(semantic_section.title, level=1)
        for block in semantic_section.blocks:
            _append_docx_block(document, block, citations, profile)
    if bundle.paginated.include_references:
        document.add_page_break()
        document.add_heading("References", level=1)
        for index, citation in enumerate(bundle.content.citations, start=1):
            paragraph = document.add_paragraph(_reference_line(index, citation))
            _set_docx_paragraph_spacing(paragraph, after=3, line=1.05)
            _format_docx_paragraph_runs(paragraph, size=9, color=profile.muted_hex)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _apply_docx_style_defaults(document, profile: ArtifactLayoutProfile) -> None:
    from docx.shared import Pt, RGBColor

    def rgb(value: str):
        return RGBColor.from_string(value)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(profile.text_hex)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for style_name, size, color, before, after in (
        ("Heading 1", 16, profile.brand_hex, 16, 8),
        ("Heading 2", 13, profile.brand_hex, 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def _set_docx_page_chrome(section, profile: ArtifactLayoutProfile) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    header = section.header.paragraphs[0]
    header.text = f"Prudentia AI | {profile.label}"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _format_docx_paragraph_runs(header, size=8.5, color=profile.muted_hex, bold=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Evidence-backed artifact | Page ")
    _format_docx_run(run, size=8, color=profile.muted_hex)
    _append_docx_page_field(footer.add_run())


def _add_docx_opening_block(
    document,
    bundle: ArtifactContentBundle,
    generated_at: datetime,
    profile: ArtifactLayoutProfile,
) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    paragraph = document.add_paragraph()
    _set_docx_paragraph_spacing(paragraph, after=1)
    run = paragraph.add_run(profile.kicker)
    _format_docx_run(run, size=10.5, bold=True, color=profile.accent_hex)

    paragraph = document.add_paragraph()
    _set_docx_paragraph_spacing(paragraph, after=4)
    run = paragraph.add_run(bundle.paginated.title)
    _format_docx_run(run, size=24, bold=True, color="000000")
    if bundle.paginated.subtitle:
        paragraph = document.add_paragraph()
        _set_docx_paragraph_spacing(paragraph, after=8)
        run = paragraph.add_run(bundle.paginated.subtitle)
        _format_docx_run(run, size=13, color=profile.muted_hex)
    meta = document.add_paragraph(f"Generated {generated_at.strftime('%Y-%m-%d %H:%M UTC')}")
    meta.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _set_docx_paragraph_spacing(meta, after=12)
    _format_docx_paragraph_runs(meta, size=8.5, color=profile.muted_hex)
    rule = document.add_paragraph()
    _set_docx_paragraph_spacing(rule, before=2, after=10)
    _set_docx_bottom_border(rule, profile.accent_hex)


def _has_wide_table(bundle: ArtifactContentBundle) -> bool:
    return any(
        block.table is not None and len(block.table.headers) > 5
        for section in bundle.paginated.sections
        for block in section.blocks
    )


def _format_docx_run(run, *, size: float | None = None, color: str | None = None, bold: bool | None = None) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def _format_docx_paragraph_runs(
    paragraph,
    *,
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
) -> None:
    for run in paragraph.runs:
        _format_docx_run(run, size=size, color=color, bold=bold)


def _set_docx_paragraph_spacing(paragraph, *, before: float = 0, after: float = 6, line: float = 1.1) -> None:
    from docx.shared import Pt

    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def _append_docx_page_field(run) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def _set_docx_bottom_border(paragraph, color: str, *, width: str = "8") -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), width)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def _set_docx_cell_shading(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_docx_cell_margins(cell, *, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_docx_table_geometry(table, widths_inches: list[float], *, indent_dxa: int = 120) -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches

    widths_dxa = [max(1, int(width * 1440)) for width in widths_inches]
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    old_grid = tbl.find(qn("w:tblGrid"))
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl.insert(0, grid)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_inches[index])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_docx_cell_margins(cell)


def _mark_docx_repeat_header(row) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def _docx_content_width_inches(document) -> float:
    section = document.sections[0]
    return float(section.page_width - section.left_margin - section.right_margin) / 914400


def _docx_table_widths(headers: list[str], available_width: float) -> list[float]:
    if not headers:
        return [available_width]
    source_column = headers[-1].strip().lower() == "source"
    body_headers = headers[:-1] if source_column and len(headers) > 1 else headers
    source_width = min(max(available_width * 0.18, 1.05), 1.35) if source_column and len(headers) > 2 else 0
    remaining = max(available_width - source_width, 1.0)
    weights = [_docx_column_weight(header) for header in body_headers]
    total_weight = sum(weights) or 1
    widths = [remaining * weight / total_weight for weight in weights]
    if source_column:
        widths.append(source_width)
    return widths


def _docx_column_weight(header: str) -> float:
    normalized = header.strip().lower()
    if normalized in {"id", "ref", "date", "time", "status", "level", "score", "owner", "source", "confidence"}:
        return 0.75
    if any(token in normalized for token in ("finding", "summary", "description", "evidence", "fact", "rationale", "action", "recommendation")):
        return 2.4
    return 1.2


def _style_docx_table(document, table, headers: list[str], profile: ArtifactLayoutProfile) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    _set_docx_table_geometry(table, _docx_table_widths(headers, _docx_content_width_inches(document)))
    _mark_docx_repeat_header(table.rows[0])
    for cell in table.rows[0].cells:
        _set_docx_cell_shading(cell, profile.brand_hex)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _set_docx_paragraph_spacing(paragraph, after=0, line=1.0)
            _format_docx_paragraph_runs(paragraph, size=8.8, color="FFFFFF", bold=True)
    for row_index, row in enumerate(table.rows[1:], start=1):
        if row_index % 2 == 0:
            for cell in row.cells:
                _set_docx_cell_shading(cell, profile.row_hex)
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                _set_docx_paragraph_spacing(paragraph, after=0, line=1.05)
                _format_docx_paragraph_runs(paragraph, size=8.6, color=profile.text_hex)


def _append_docx_callout(
    document,
    text: str,
    evidence_ids: list[str],
    citations: dict[str, EvidenceCitation],
    profile: ArtifactLayoutProfile,
) -> None:
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    _set_docx_table_geometry(table, [_docx_content_width_inches(document)])
    cell = table.cell(0, 0)
    cell.text = _with_sources(text, evidence_ids, citations)
    _set_docx_cell_shading(cell, profile.callout_hex)
    for paragraph in cell.paragraphs:
        _set_docx_paragraph_spacing(paragraph, after=0, line=1.15)
        _format_docx_paragraph_runs(paragraph, size=10, color=profile.text_hex, bold=True)
    spacer = document.add_paragraph()
    _set_docx_paragraph_spacing(spacer, after=2)


def _append_docx_block(
    document,
    block: ContentBlock,
    citations: dict[str, EvidenceCitation],
    profile: ArtifactLayoutProfile,
) -> None:
    if block.kind == "heading":
        document.add_heading(block.text or "", level=block.level)
        return
    if block.kind in {"page_break", "section_break"}:
        document.add_page_break()
        return
    if block.kind == "paragraph":
        paragraph = document.add_paragraph(_with_sources(block.text or "", block.evidence_ids, citations))
        _set_docx_paragraph_spacing(paragraph, after=6, line=1.12)
        _format_docx_paragraph_runs(paragraph, size=10.5, color=profile.text_hex)
        return
    if block.kind in {"bullet_list", "numbered_list"}:
        style = "List Bullet" if block.kind == "bullet_list" else "List Number"
        for item_text, evidence_ids in _list_items(block):
            paragraph = document.add_paragraph(_with_sources(item_text, evidence_ids, citations), style=style)
            _set_docx_paragraph_spacing(paragraph, after=4, line=1.15)
            _format_docx_paragraph_runs(paragraph, size=10.2, color=profile.text_hex)
        return
    if block.kind == "quotation":
        paragraph = document.add_paragraph(_with_sources(block.text or "", block.evidence_ids, citations), style="Quote")
        _set_docx_paragraph_spacing(paragraph, after=6, line=1.12)
        _format_docx_paragraph_runs(paragraph, size=10, color=profile.muted_hex)
        return
    if block.kind == "callout":
        _append_docx_callout(document, block.text or "", block.evidence_ids, citations, profile)
        return
    if block.kind == "key_value":
        table = document.add_table(rows=0, cols=3)
        table.style = "Table Grid"
        header_cells = table.add_row().cells
        for index, header in enumerate(("Field", "Value", "Source")):
            header_cells[index].text = header
        for entry in block.entries:
            cells = table.add_row().cells
            cells[0].text = entry.key
            cells[1].text = entry.value
            cells[2].text = _source_summary(entry.evidence_ids, citations)
        _style_docx_table(document, table, ["Field", "Value", "Source"], profile)
        return
    if block.kind == "table" and block.table is not None:
        table = document.add_table(rows=1, cols=len(block.table.headers) + 1)
        table.style = "Table Grid"
        headers = [*block.table.headers, "Source"]
        for index, header in enumerate(headers):
            table.rows[0].cells[index].text = header
        for row in block.table.rows:
            cells = table.add_row().cells
            for index, value in enumerate(row.values):
                cells[index].text = value
            cells[-1].text = _source_summary(row.evidence_ids, citations)
        _style_docx_table(document, table, headers, profile)
