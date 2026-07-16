"""Render validated artifact bundles as PPTX presentations."""

from __future__ import annotations

from collections.abc import Callable
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
import re

from ..contracts import ArtifactContentBundle, ContentBlock, EvidenceCitation
from rag.artifact_jobs.generation.layout_profiles import ArtifactLayoutProfile
from .shared import _compact, _list_items, _reference_line, _source_summary


PPTX_MAX_BULLETS_PER_SLIDE = 8
PPTX_MAX_TEXT_UNITS_PER_SLIDE = 7
PPTX_TEXT_UNIT_CHARS = 150
PPTX_BODY_PANEL_LEFT_IN = 0.65
PPTX_BODY_PANEL_TOP_IN = 1.12
PPTX_BODY_PANEL_WIDTH_IN = 12.02
PPTX_BODY_PANEL_HEIGHT_IN = 5.48
PPTX_TABLE_LEFT_IN = 0.92
PPTX_TABLE_TOP_IN = 1.52
PPTX_TABLE_WIDTH_IN = 11.35
PPTX_TABLE_MAX_HEIGHT_IN = 4.72
PPTX_TABLE_HEADER_MIN_HEIGHT_IN = 0.42
PPTX_TABLE_ROW_MIN_HEIGHT_IN = 0.34
PPTX_TABLE_ROW_MAX_HEIGHT_IN = 1.1
PPTX_TABLE_BODY_FONT_MAX = 9.3
PPTX_TABLE_BODY_FONT_MIN = 7.2
_REFERENCE_SLIDE_TITLES = {"reference", "references", "source", "sources", "citation", "citations", "bibliography"}


@dataclass(frozen=True)
class PptxTableLayout:
    column_widths: list[float]
    row_heights: list[float]
    header_height: float
    body_font_size: float

    @property
    def height(self) -> float:
        return self.header_height + sum(self.row_heights)


def render_pptx(
    bundle: ArtifactContentBundle,
    generated_at: datetime,
    profile: ArtifactLayoutProfile,
    *,
    progress_callback: Callable[[int, int, str], None] | None = None,
) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    citations = {item.evidence_id: item for item in bundle.content.citations}
    total_slides = _pptx_render_slide_count(bundle)
    rendered_slides = 0
    content_slide_index = 0

    def record_slide(label: str) -> None:
        nonlocal rendered_slides
        rendered_slides += 1
        if progress_callback:
            progress_callback(rendered_slides, total_slides, label)

    def add_content_slide(title: str):
        nonlocal content_slide_index
        slide = _add_base_content_slide(presentation, title, content_slide_index, profile)
        content_slide_index += 1
        return slide

    title_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _set_slide_background(title_slide, RGBColor(*profile.pptx_title_background))
    _decorate_title_slide(title_slide, bundle, generated_at, profile)
    record_slide("Title slide")
    for slide_spec in _content_slides(bundle.presentation.slides):
        table_block = next((block for block in slide_spec.blocks if block.kind == "table" and block.table), None)
        if table_block is not None:
            chunks = _table_block_chunks(table_block)
            for chunk_index, chunk in enumerate(chunks, start=1):
                title = slide_spec.title if len(chunks) == 1 else f"{slide_spec.title} ({chunk_index}/{len(chunks)})"
                slide = add_content_slide(title)
                _add_slide_table(slide, chunk, profile)
                _add_slide_sources(slide, _block_evidence_ids(chunk), citations, profile)
                record_slide(title)
            continue

        evidence_ids = list(dict.fromkeys(
            evidence_id
            for block in slide_spec.blocks
            for evidence_id in _block_evidence_ids(block)
        ))
        line_chunks = _slide_line_chunks(_slide_lines(slide_spec.blocks) or ["No supported content was generated."])
        for chunk_index, line_chunk in enumerate(line_chunks, start=1):
            title = slide_spec.title if len(line_chunks) == 1 else f"{slide_spec.title} ({chunk_index}/{len(line_chunks)})"
            slide = add_content_slide(title)
            _add_slide_bullets(slide, line_chunk, profile)
            _add_slide_sources(slide, evidence_ids, citations, profile)
            record_slide(title)
    if bundle.presentation.include_references_slide:
        slide = _add_base_content_slide(presentation, "References", content_slide_index, profile)
        _add_reference_lines(
            slide,
            [_reference_line(index, citation) for index, citation in enumerate(bundle.content.citations[:18], start=1)],
            profile,
        )
        record_slide("References")
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _pptx_render_slide_count(bundle: ArtifactContentBundle) -> int:
    total = 1
    for slide_spec in _content_slides(bundle.presentation.slides):
        table_block = next((block for block in slide_spec.blocks if block.kind == "table" and block.table), None)
        if table_block is not None:
            total += len(_table_block_chunks(table_block))
        else:
            total += len(_slide_line_chunks(_slide_lines(slide_spec.blocks) or ["No supported content was generated."]))
    if bundle.presentation.include_references_slide:
        total += 1
    return max(total, 1)


def _decorate_title_slide(
    slide,
    bundle: ArtifactContentBundle,
    generated_at: datetime,
    profile: ArtifactLayoutProfile,
) -> None:
    title_background = _rgb_tuple_hex(profile.pptx_title_background)
    panel_hex = _scale_hex(title_background, 0.82)
    _add_slide_rect(slide, 0, 0, 0.18, 7.5, profile.accent_hex)
    _add_slide_rect(slide, 9.25, 0, 4.08, 7.5, panel_hex)
    _add_slide_rect(slide, 0.85, 0.85, 1.0, 0.05, profile.accent_hex)
    _add_slide_text(slide, profile.kicker, 0.85, 0.98, 5.2, 0.3, 11, profile.accent_hex, bold=True)
    _add_slide_text(
        slide,
        bundle.presentation.title,
        0.82,
        1.45,
        7.85,
        1.7,
        _title_font_size(bundle.presentation.title),
        "FFFFFF",
        bold=True,
    )
    if bundle.presentation.subtitle:
        _add_slide_text(slide, bundle.presentation.subtitle, 0.88, 3.22, 7.2, 0.82, 15, "D6E4E5")
    _add_slide_text(slide, "Prudentia AI", 9.75, 0.88, 2.7, 0.32, 11, "FFFFFF", bold=True)
    _add_slide_text(slide, profile.label, 9.75, 1.32, 2.7, 0.28, 9, "D6E4E5")
    _add_title_metric(slide, "Generated", generated_at.date().isoformat(), 9.75, 2.25, profile)
    _add_title_metric(slide, "Sources", str(len(bundle.content.citations)), 9.75, 3.22, profile)
    _add_title_metric(slide, "Sections", str(len(bundle.content.sections)), 9.75, 4.19, profile)
    _add_slide_text(slide, "Evidence-backed artifact", 9.75, 6.52, 2.6, 0.25, 8, "AFC7C9")


def _add_title_metric(slide, label: str, value: str, left: float, top: float, profile: ArtifactLayoutProfile) -> None:
    _add_slide_rect(slide, left, top, 2.55, 0.62, _scale_hex(_rgb_tuple_hex(profile.pptx_title_background), 0.92), profile.accent_hex)
    _add_slide_text(slide, label.upper(), left + 0.18, top + 0.12, 1.1, 0.16, 7, "AFC7C9", bold=True)
    _add_slide_text(slide, value, left + 1.42, top + 0.1, 0.85, 0.26, 14, "FFFFFF", bold=True)


def _add_base_content_slide(presentation, title: str, content_slide_index: int, profile: ArtifactLayoutProfile):
    from pptx.dml.color import RGBColor

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _set_slide_background(slide, RGBColor(*_content_background(content_slide_index, profile)))
    _add_slide_rect(slide, 0, 0, 0.15, 7.5, profile.brand_hex)
    _add_slide_rect(slide, 0.15, 0, 0.045, 7.5, profile.accent_hex)
    _add_slide_text(slide, profile.kicker, 0.65, 0.18, 4.0, 0.22, 8, profile.accent_hex, bold=True)
    _add_slide_text(slide, title, 0.65, 0.42, 10.75, 0.62, _heading_font_size(title), profile.brand_hex, bold=True)
    _add_slide_text(slide, f"{content_slide_index + 2:02d}", 12.08, 0.32, 0.48, 0.24, 9, profile.muted_hex, bold=True)
    _add_slide_rect(
        slide,
        PPTX_BODY_PANEL_LEFT_IN,
        PPTX_BODY_PANEL_TOP_IN,
        PPTX_BODY_PANEL_WIDTH_IN,
        PPTX_BODY_PANEL_HEIGHT_IN,
        "FFFFFF",
        profile.rule_hex,
    )
    _add_slide_rect(slide, PPTX_BODY_PANEL_LEFT_IN, PPTX_BODY_PANEL_TOP_IN, 0.08, PPTX_BODY_PANEL_HEIGHT_IN, profile.accent_hex)
    _add_slide_rect(slide, 0.68, 6.71, 11.95, 0.015, profile.rule_hex)
    return slide


def _content_background(index: int, profile: ArtifactLayoutProfile) -> tuple[int, int, int]:
    return profile.pptx_backgrounds[index % len(profile.pptx_backgrounds)]


def _content_slides(slides):
    return [slide for slide in slides if not _is_generated_references_slide(slide)]


def _is_generated_references_slide(slide) -> bool:
    return _normalized_reference_title(slide.title) in _REFERENCE_SLIDE_TITLES


def _normalized_reference_title(value: str) -> str:
    return re.sub(r"[^a-z]+", " ", value.lower()).strip()


def _add_slide_sources(
    slide,
    evidence_ids: list[str],
    citations: dict[str, EvidenceCitation],
    profile: ArtifactLayoutProfile,
) -> None:
    source_text = _source_summary(list(dict.fromkeys(evidence_ids)), citations)
    if source_text:
        _add_slide_text(slide, source_text, 0.78, 6.86, 10.9, 0.27, 7.6, profile.muted_hex)


def _set_slide_background(slide, color) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_slide_rect(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    fill_hex: str,
    line_hex: str | None = None,
):
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill_hex)
    if line_hex:
        shape.line.color.rgb = RGBColor.from_string(line_hex)
    else:
        shape.line.fill.background()
    return shape


def _add_slide_text(slide, text, left, top, width, height, size, color, *, bold=False) -> None:
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor.from_string(color)


def _title_font_size(text: str) -> int:
    length = len(_compact(text, 180))
    if length <= 45:
        return 38
    if length <= 70:
        return 32
    if length <= 95:
        return 27
    return 23


def _heading_font_size(text: str) -> int:
    return 25 if len(text) <= 60 else 21


def _add_slide_bullets(slide, lines: list[str], profile: ArtifactLayoutProfile) -> None:
    display_lines = lines or ["No supported content was generated."]
    max_length = max((len(line) for line in display_lines), default=0)
    row_gap = 0.08 if len(display_lines) <= 6 else 0.05
    available_height = 4.78
    font_size = 14.2 if len(display_lines) <= 4 and max_length <= 220 else 12.8 if len(display_lines) <= 6 else 11.4
    row_heights = _pptx_bullet_row_heights(display_lines, font_size, available_height, row_gap)
    top = 1.42
    left = 0.95
    width = 11.18
    row_top = top
    for index, line in enumerate(display_lines):
        row_height = row_heights[index]
        if len(display_lines) <= 8:
            _add_slide_rect(slide, left, row_top, width, row_height, profile.row_hex if index % 2 else "FFFFFF", profile.rule_hex)
        _add_slide_rect(slide, left + 0.18, row_top + 0.12, 0.055, max(0.22, row_height - 0.24), profile.accent_hex)
        _add_slide_text(slide, f"{index + 1:02d}", left + 0.36, row_top + 0.13, 0.38, 0.18, 7.8, profile.muted_hex, bold=True)
        limit = 360 if row_height >= 0.72 else 285 if row_height >= 0.55 else 220
        _add_slide_text(
            slide,
            _compact(line, limit),
            left + 0.82,
            row_top + 0.1,
            width - 1.05,
            max(0.22, row_height - 0.1),
            font_size,
            profile.text_hex,
        )
        row_top += row_height + row_gap


def _add_reference_lines(slide, lines: list[str], profile: ArtifactLayoutProfile) -> None:
    if not lines:
        _add_slide_text(slide, "No references were available.", 0.95, 1.48, 10.8, 0.35, 12, profile.text_hex)
        return
    columns = (lines[:9], lines[9:18])
    for column_index, column_lines in enumerate(columns):
        left = 0.95 + column_index * 5.72
        for row_index, line in enumerate(column_lines):
            top = 1.46 + row_index * 0.49
            _add_slide_text(slide, _compact(line, 155), left, top, 5.15, 0.35, 8.1, profile.text_hex)


def _pptx_bullet_row_heights(lines: list[str], font_size: float, available_height: float, row_gap: float) -> list[float]:
    if not lines:
        return [available_height]
    usable_height = max(0.44, available_height - row_gap * max(0, len(lines) - 1))
    line_counts = [
        max(1, _pptx_estimated_line_count(line, 9.7, font_size))
        for line in lines
    ]
    desired = [
        min(1.78, max(0.48, 0.34 + line_count * (font_size / 72 * 1.55)))
        for line_count in line_counts
    ]
    total_desired = sum(desired)
    if total_desired <= usable_height:
        return desired
    scale = usable_height / total_desired if total_desired else 1
    scaled = [max(0.44, height * scale) for height in desired]
    overflow = sum(scaled) - usable_height
    if overflow > 0:
        reducible = [
            (index, height - 0.44)
            for index, height in enumerate(scaled)
            if height > 0.44
        ]
        capacity = sum(amount for _index, amount in reducible)
        if capacity > 0:
            for index, amount in reducible:
                scaled[index] -= overflow * (amount / capacity)
    return scaled


def _rgb_tuple_hex(value: tuple[int, int, int]) -> str:
    return "".join(f"{part:02X}" for part in value)


def _scale_hex(value: str, factor: float) -> str:
    normalized = value.strip().lstrip("#")
    if len(normalized) != 6:
        return value
    channels = [
        max(0, min(255, int(int(normalized[index : index + 2], 16) * factor)))
        for index in (0, 2, 4)
    ]
    return "".join(f"{channel:02X}" for channel in channels)


def _add_slide_table(slide, block: ContentBlock, profile: ArtifactLayoutProfile) -> None:
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    assert block.table is not None
    body_rows = block.table.rows
    rows = len(body_rows) + 1
    cols = len(block.table.headers)
    layout = _pptx_table_layout(block, fit_to_frame=True)
    shape = slide.shapes.add_table(
        rows,
        cols,
        Inches(PPTX_TABLE_LEFT_IN),
        Inches(PPTX_TABLE_TOP_IN),
        Inches(PPTX_TABLE_WIDTH_IN),
        Inches(layout.height),
    )
    table = shape.table
    for column, width in enumerate(layout.column_widths):
        table.columns[column].width = Inches(width)
    table.rows[0].height = Inches(layout.header_height)
    for column, header in enumerate(block.table.headers):
        cell = table.cell(0, column)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(profile.brand_hex)
        cell.margin_left = Inches(0.1)
        cell.margin_right = Inches(0.1)
        cell.margin_top = Inches(0.05)
        cell.margin_bottom = Inches(0.05)
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.bold = True
        run.font.size = Pt(10 if cols <= 4 else 9)
    for row_index, row in enumerate(body_rows, start=1):
        table.rows[row_index].height = Inches(layout.row_heights[row_index - 1])
        for column, value in enumerate(row.values):
            cell = table.cell(row_index, column)
            cell.text = _compact(value, _pptx_cell_text_limit(layout.column_widths[column], layout.body_font_size))
            cell.margin_left = Inches(0.1)
            cell.margin_right = Inches(0.1)
            cell.margin_top = Inches(0.05)
            cell.margin_bottom = Inches(0.05)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(profile.row_hex if row_index % 2 == 1 else "FFFFFF")
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = "Aptos"
                paragraph.font.size = Pt(layout.body_font_size)
                paragraph.font.color.rgb = RGBColor.from_string(profile.text_hex)


def _slide_lines(blocks: list[ContentBlock]) -> list[str]:
    lines: list[str] = []
    for block in blocks:
        if block.kind in {"paragraph", "quotation", "callout", "heading"} and block.text:
            lines.append(block.text)
        elif block.kind in {"bullet_list", "numbered_list"}:
            lines.extend(item_text for item_text, _evidence_ids in _list_items(block))
        elif block.kind == "key_value":
            lines.extend(f"{entry.key}: {entry.value}" for entry in block.entries)
    return lines


def _block_evidence_ids(block: ContentBlock) -> list[str]:
    ids = list(block.evidence_ids)
    ids.extend(evidence_id for item in block.list_items for evidence_id in item.evidence_ids)
    ids.extend(evidence_id for entry in block.entries for evidence_id in entry.evidence_ids)
    if block.table is not None:
        ids.extend(evidence_id for row in block.table.rows for evidence_id in row.evidence_ids)
    return ids


def _table_block_chunks(block: ContentBlock, *, max_body_rows: int | None = None) -> list[ContentBlock]:
    if block.table is None:
        return [block]
    rows = block.table.rows
    if not rows:
        return [block]
    chunks: list[ContentBlock] = []
    current: list = []
    for row in rows:
        if max_body_rows is not None and current and len(current) >= max_body_rows:
            chunks.append(block.model_copy(update={"table": block.table.model_copy(update={"rows": current})}))
            current = []
        candidate = [*current, row]
        candidate_block = block.model_copy(update={"table": block.table.model_copy(update={"rows": candidate})})
        candidate_height = _pptx_table_layout(candidate_block, fit_to_frame=False).height
        if current and candidate_height > PPTX_TABLE_MAX_HEIGHT_IN:
            chunks.append(block.model_copy(update={"table": block.table.model_copy(update={"rows": current})}))
            current = [row]
        else:
            current = candidate
    if current:
        chunks.append(block.model_copy(update={"table": block.table.model_copy(update={"rows": current})}))
    return chunks or [block]


def _pptx_table_layout(block: ContentBlock, *, fit_to_frame: bool) -> PptxTableLayout:
    assert block.table is not None
    headers = list(block.table.headers)
    rows = block.table.rows
    column_widths = _pptx_table_col_widths(headers, rows)
    body_font_size = _pptx_table_body_font_size(headers, rows)
    header_height = max(
        PPTX_TABLE_HEADER_MIN_HEIGHT_IN,
        0.22 + 0.18 * max(
            _pptx_estimated_line_count(header, column_widths[index], 10)
            for index, header in enumerate(headers)
        ),
    )
    row_heights = [
        _pptx_table_row_height(row.values, column_widths, body_font_size)
        for row in rows
    ]
    layout = PptxTableLayout(
        column_widths=column_widths,
        row_heights=row_heights,
        header_height=min(header_height, 0.72),
        body_font_size=body_font_size,
    )
    return _fit_pptx_table_layout(layout) if fit_to_frame else layout


def _fit_pptx_table_layout(layout: PptxTableLayout) -> PptxTableLayout:
    if layout.height <= PPTX_TABLE_MAX_HEIGHT_IN or not layout.row_heights:
        return layout
    available_body_height = max(PPTX_TABLE_MAX_HEIGHT_IN - layout.header_height, PPTX_TABLE_ROW_MIN_HEIGHT_IN)
    total_body_height = sum(layout.row_heights)
    scale = available_body_height / total_body_height if total_body_height else 1
    fitted_rows = [max(PPTX_TABLE_ROW_MIN_HEIGHT_IN, height * scale) for height in layout.row_heights]
    if layout.header_height + sum(fitted_rows) > PPTX_TABLE_MAX_HEIGHT_IN:
        overflow = layout.header_height + sum(fitted_rows) - PPTX_TABLE_MAX_HEIGHT_IN
        reducible = [
            (index, height - PPTX_TABLE_ROW_MIN_HEIGHT_IN)
            for index, height in enumerate(fitted_rows)
            if height > PPTX_TABLE_ROW_MIN_HEIGHT_IN
        ]
        total_reducible = sum(value for _index, value in reducible)
        if total_reducible > 0:
            for index, value in reducible:
                fitted_rows[index] -= overflow * (value / total_reducible)
    font_size = layout.body_font_size if scale > 0.9 else max(PPTX_TABLE_BODY_FONT_MIN, layout.body_font_size - 0.5)
    return PptxTableLayout(
        column_widths=layout.column_widths,
        row_heights=fitted_rows,
        header_height=layout.header_height,
        body_font_size=font_size,
    )


def _pptx_table_col_widths(headers: list[str], rows) -> list[float]:
    if not headers:
        return [PPTX_TABLE_WIDTH_IN]
    values_by_column = [
        [
            str(row.values[index]) if index < len(row.values) else ""
            for row in rows
        ]
        for index in range(len(headers))
    ]
    weights = [
        _pptx_column_weight(header, values_by_column[index])
        for index, header in enumerate(headers)
    ]
    total_weight = sum(weights) or 1
    widths = [PPTX_TABLE_WIDTH_IN * weight / total_weight for weight in weights]
    minimums = [_pptx_column_min_width(header) for header in headers]
    if sum(minimums) > PPTX_TABLE_WIDTH_IN:
        scale = PPTX_TABLE_WIDTH_IN / sum(minimums)
        minimums = [width * scale for width in minimums]
    maximums = [_pptx_column_max_width(header, len(headers)) for header in headers]
    maximums = [max(maximum, minimums[index]) for index, maximum in enumerate(maximums)]
    return _redistribute_widths(widths, minimums, maximums, PPTX_TABLE_WIDTH_IN)


def _pptx_column_weight(header: str, values: list[str]) -> float:
    normalized = header.strip().lower()
    lengths = [len(_normalized_cell_text(header)), *(len(_normalized_cell_text(value)) for value in values)]
    average = sum(min(length, 160) for length in lengths) / max(len(lengths), 1)
    longest = max(lengths, default=0)
    content_weight = 0.55 + average / 34 + min(longest, 220) / 115
    if normalized in {"id", "ref", "date", "time", "status", "score", "level", "owner"}:
        return max(0.75, content_weight * 0.72)
    if any(token in normalized for token in ("event", "finding", "summary", "description", "evidence", "item", "rationale", "action", "recommendation", "note")):
        return max(1.75, content_weight * 1.25)
    return max(1.0, content_weight)


def _pptx_column_min_width(header: str) -> float:
    normalized = header.strip().lower()
    if normalized in {"id", "ref", "date", "time", "score", "level"}:
        return 0.72
    if normalized in {"status", "owner", "confidence"}:
        return 1.05
    return 1.25


def _pptx_column_max_width(header: str, column_count: int) -> float:
    normalized = header.strip().lower()
    if normalized in {"id", "ref", "date", "time", "score", "level"}:
        return 1.6
    if column_count <= 2:
        return PPTX_TABLE_WIDTH_IN * 0.72
    if any(token in normalized for token in ("event", "description", "summary", "finding", "recommendation", "note")):
        return PPTX_TABLE_WIDTH_IN * 0.58
    return PPTX_TABLE_WIDTH_IN * 0.45


def _redistribute_widths(widths: list[float], minimums: list[float], maximums: list[float], total_width: float) -> list[float]:
    adjusted = [min(max(widths[index], minimums[index]), maximums[index]) for index in range(len(widths))]
    for _iteration in range(8):
        delta = total_width - sum(adjusted)
        if abs(delta) < 0.01:
            break
        if delta > 0:
            candidates = [(index, maximums[index] - adjusted[index]) for index in range(len(adjusted)) if adjusted[index] < maximums[index]]
        else:
            candidates = [(index, adjusted[index] - minimums[index]) for index in range(len(adjusted)) if adjusted[index] > minimums[index]]
        capacity = sum(value for _index, value in candidates)
        if capacity <= 0:
            break
        for index, value in candidates:
            adjusted[index] += delta * (value / capacity)
            adjusted[index] = min(max(adjusted[index], minimums[index]), maximums[index])
    if adjusted:
        adjusted[-1] += total_width - sum(adjusted)
    return adjusted


def _pptx_table_body_font_size(headers: list[str], rows) -> float:
    cell_lengths = [
        len(_normalized_cell_text(value))
        for row in rows
        for value in row.values
    ]
    longest = max(cell_lengths, default=0)
    row_count = len(rows)
    column_count = len(headers)
    font_size = PPTX_TABLE_BODY_FONT_MAX
    if longest > 160 or row_count > 7:
        font_size -= 1.3
    elif longest > 95 or row_count > 5:
        font_size -= 0.6
    if column_count >= 5:
        font_size -= 0.6
    return max(PPTX_TABLE_BODY_FONT_MIN, font_size)


def _pptx_table_row_height(values: list[str], column_widths: list[float], font_size: float) -> float:
    line_count = max(
        (
            _pptx_estimated_line_count(value, column_widths[index], font_size)
            for index, value in enumerate(values)
            if index < len(column_widths)
        ),
        default=1,
    )
    height = 0.20 + line_count * (font_size / 72 * 1.28)
    return min(max(height, PPTX_TABLE_ROW_MIN_HEIGHT_IN), PPTX_TABLE_ROW_MAX_HEIGHT_IN)


def _pptx_estimated_line_count(text: str, column_width: float, font_size: float) -> int:
    normalized = _normalized_cell_text(text)
    if not normalized:
        return 1
    chars_per_line = max(6, int(column_width * 13.0 * (9.0 / max(font_size, 1))))
    lines = 1
    current = 0
    for word in normalized.split(" "):
        length = len(word)
        if length > chars_per_line:
            extra_lines = (length + chars_per_line - 1) // chars_per_line
            if current:
                lines += 1
                current = 0
            lines += max(0, extra_lines - 1)
            current = length % chars_per_line
            continue
        next_length = length if current == 0 else current + 1 + length
        if next_length <= chars_per_line:
            current = next_length
        else:
            lines += 1
            current = length
    return lines


def _pptx_cell_text_limit(column_width: float, font_size: float) -> int:
    lines_that_fit = max(2, int(PPTX_TABLE_ROW_MAX_HEIGHT_IN / (font_size / 72 * 1.28)))
    chars_per_line = max(10, int(column_width * 13.0 * (9.0 / max(font_size, 1))))
    return max(80, min(360, chars_per_line * lines_that_fit))


def _normalized_cell_text(value: str) -> str:
    return re.sub(r"\s+", " ", str(value)).strip()


def _slide_line_chunks(lines: list[str]) -> list[list[str]]:
    chunks: list[list[str]] = []
    current: list[str] = []
    current_units = 0
    for line in lines:
        units = _pptx_text_units(line)
        if current and (
            len(current) >= PPTX_MAX_BULLETS_PER_SLIDE
            or current_units + units > PPTX_MAX_TEXT_UNITS_PER_SLIDE
        ):
            chunks.append(current)
            current = []
            current_units = 0
        current.append(line)
        current_units += units
    if current:
        chunks.append(current)
    return chunks or [[]]


def _pptx_text_units(text: str) -> int:
    return max(1, (len(text.strip()) + PPTX_TEXT_UNIT_CHARS - 1) // PPTX_TEXT_UNIT_CHARS)
