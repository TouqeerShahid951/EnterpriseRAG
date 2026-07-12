"""Deterministic renderers for validated v2 document specifications."""

from __future__ import annotations

from collections.abc import Callable
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import zipfile
from xml.sax.saxutils import escape

from ..schemas.query import ArtifactFormat
from .contracts import ArtifactContentBundle, ContentBlock, EvidenceCitation
from .layout_profiles import ArtifactLayoutProfile, select_layout_profile


CONTENT_TYPES: dict[ArtifactFormat, str] = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "pdf": "application/pdf",
}
PPTX_MAX_BULLETS_PER_SLIDE = 8
PPTX_MAX_TEXT_UNITS_PER_SLIDE = 7
PPTX_TEXT_UNIT_CHARS = 150
PPTX_MAX_SOURCE_LABELS = 5
PPTX_BODY_PANEL_LEFT_IN = 0.65
PPTX_BODY_PANEL_TOP_IN = 1.12
PPTX_BODY_PANEL_WIDTH_IN = 12.02
PPTX_BODY_PANEL_HEIGHT_IN = 5.48
PPTX_TABLE_LEFT_IN = 0.92
PPTX_TABLE_TOP_IN = 1.52
PPTX_TABLE_WIDTH_IN = 11.35
PPTX_TABLE_MAX_HEIGHT_IN = 4.72
PPTX_TABLE_HEADER_MIN_HEIGHT_IN = 0.42
PPTX_TABLE_ROW_MIN_HEIGHT_IN = 0.34
PPTX_TABLE_ROW_MAX_HEIGHT_IN = 1.1
PPTX_TABLE_BODY_FONT_MAX = 9.3
PPTX_TABLE_BODY_FONT_MIN = 7.2
_REFERENCE_SLIDE_TITLES = {"reference", "references", "source", "sources", "citation", "citations", "bibliography"}
PDF_WIDE_TABLE_COLUMNS = 5


@dataclass(frozen=True)
class RenderedV2Artifact:
    filename: str
    format: ArtifactFormat
    content_type: str
    content: bytes
    smoke_warnings: tuple[str, ...] = ()


@dataclass(frozen=True)
class PptxTableLayout:
    column_widths: list[float]
    row_heights: list[float]
    header_height: float
    body_font_size: float

    @property
    def height(self) -> float:
        return self.header_height + sum(self.row_heights)


def render_document(
    *,
    artifact_format: ArtifactFormat,
    bundle: ArtifactContentBundle,
    generated_at: datetime,
    require_libreoffice: bool,
    progress_callback: Callable[[int, int, str], None] | None = None,
) -> RenderedV2Artifact:
    profile = select_layout_profile(bundle)
    if artifact_format == "docx":
        content = _render_docx(bundle, generated_at, profile)
    elif artifact_format == "pdf":
        content = _render_pdf(bundle, generated_at, profile)
    elif artifact_format == "pptx":
        content = _render_pptx(bundle, generated_at, profile, progress_callback=progress_callback)
    else:
        raise ValueError(f"unsupported artifact format: {artifact_format}")
    _validate_package(artifact_format, content)
    warnings = _office_smoke_check(
        artifact_format=artifact_format,
        content=content,
        require_libreoffice=require_libreoffice,
    )
    return RenderedV2Artifact(
        filename=f"{_safe_filename(bundle.content.title)}.{artifact_format}",
        format=artifact_format,
        content_type=CONTENT_TYPES[artifact_format],
        content=content,
        smoke_warnings=warnings,
    )


def _render_docx(bundle: ArtifactContentBundle, generated_at: datetime, profile: ArtifactLayoutProfile) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches

    document = Document()
    section = document.sections[0]
    if _has_wide_table(bundle):
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.85 if _has_wide_table(bundle) else 1.0)
    section.right_margin = Inches(0.85 if _has_wide_table(bundle) else 1.0)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)
    _apply_docx_style_defaults(document, profile)
    _set_docx_page_chrome(section, profile)
    _add_docx_opening_block(document, bundle, generated_at, profile)

    citations = {item.evidence_id: item for item in bundle.content.citations}
    if bundle.paginated.include_coverage_notes and bundle.content.warnings:
        document.add_heading("Coverage Notes", level=1)
        for warning in bundle.content.warnings:
            paragraph = document.add_paragraph(warning, style="List Bullet")
            _format_docx_paragraph_runs(paragraph, size=10.5, color=profile.text_hex)
    for semantic_section in bundle.paginated.sections:
        document.add_heading(semantic_section.title, level=1)
        for block in semantic_section.blocks:
            _append_docx_block(document, block, citations, profile)
    if bundle.paginated.include_references:
        document.add_page_break()
        document.add_heading("References", level=1)
        for index, citation in enumerate(bundle.content.citations, start=1):
            paragraph = document.add_paragraph(_reference_line(index, citation))
            _set_docx_paragraph_spacing(paragraph, after=3, line=1.05)
            _format_docx_paragraph_runs(paragraph, size=9, color=profile.muted_hex)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _apply_docx_style_defaults(document, profile: ArtifactLayoutProfile) -> None:
    from docx.shared import Pt, RGBColor

    def rgb(value: str):
        return RGBColor.from_string(value)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(profile.text_hex)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for style_name, size, color, before, after in (
        ("Heading 1", 16, profile.brand_hex, 16, 8),
        ("Heading 2", 13, profile.brand_hex, 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def _set_docx_page_chrome(section, profile: ArtifactLayoutProfile) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    header = section.header.paragraphs[0]
    header.text = f"Prudentia AI | {profile.label}"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _format_docx_paragraph_runs(header, size=8.5, color=profile.muted_hex, bold=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Evidence-backed artifact | Page ")
    _format_docx_run(run, size=8, color=profile.muted_hex)
    _append_docx_page_field(footer.add_run())


def _add_docx_opening_block(
    document,
    bundle: ArtifactContentBundle,
    generated_at: datetime,
    profile: ArtifactLayoutProfile,
) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    paragraph = document.add_paragraph()
    _set_docx_paragraph_spacing(paragraph, after=1)
    run = paragraph.add_run(profile.kicker)
    _format_docx_run(run, size=10.5, bold=True, color=profile.accent_hex)

    paragraph = document.add_paragraph()
    _set_docx_paragraph_spacing(paragraph, after=4)
    run = paragraph.add_run(bundle.paginated.title)
    _format_docx_run(run, size=24, bold=True, color="000000")
    if bundle.paginated.subtitle:
        paragraph = document.add_paragraph()
        _set_docx_paragraph_spacing(paragraph, after=8)
        run = paragraph.add_run(bundle.paginated.subtitle)
        _format_docx_run(run, size=13, color=profile.muted_hex)
    meta = document.add_paragraph(f"Generated {generated_at.strftime('%Y-%m-%d %H:%M UTC')}")
    meta.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _set_docx_paragraph_spacing(meta, after=12)
    _format_docx_paragraph_runs(meta, size=8.5, color=profile.muted_hex)
    rule = document.add_paragraph()
    _set_docx_paragraph_spacing(rule, before=2, after=10)
    _set_docx_bottom_border(rule, profile.accent_hex)


def _has_wide_table(bundle: ArtifactContentBundle) -> bool:
    return any(
        block.table is not None and len(block.table.headers) > 5
        for section in bundle.paginated.sections
        for block in section.blocks
    )


def _has_wide_pdf_table(bundle: ArtifactContentBundle) -> bool:
    return any(
        block.table is not None and len(block.table.headers) >= PDF_WIDE_TABLE_COLUMNS
        for section in bundle.paginated.sections
        for block in section.blocks
    )


def _format_docx_run(run, *, size: float | None = None, color: str | None = None, bold: bool | None = None) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def _format_docx_paragraph_runs(
    paragraph,
    *,
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
) -> None:
    for run in paragraph.runs:
        _format_docx_run(run, size=size, color=color, bold=bold)


def _set_docx_paragraph_spacing(paragraph, *, before: float = 0, after: float = 6, line: float = 1.1) -> None:
    from docx.shared import Pt

    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def _append_docx_page_field(run) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def _set_docx_bottom_border(paragraph, color: str, *, width: str = "8") -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), width)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def _set_docx_cell_shading(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_docx_cell_margins(cell, *, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_docx_table_geometry(table, widths_inches: list[float], *, indent_dxa: int = 120) -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches

    widths_dxa = [max(1, int(width * 1440)) for width in widths_inches]
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    old_grid = tbl.find(qn("w:tblGrid"))
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl.insert(0, grid)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_inches[index])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_docx_cell_margins(cell)


def _mark_docx_repeat_header(row) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def _docx_content_width_inches(document) -> float:
    section = document.sections[0]
    return float(section.page_width - section.left_margin - section.right_margin) / 914400


def _docx_table_widths(headers: list[str], available_width: float) -> list[float]:
    if not headers:
        return [available_width]
    source_column = headers[-1].strip().lower() == "source"
    body_headers = headers[:-1] if source_column and len(headers) > 1 else headers
    source_width = min(max(available_width * 0.18, 1.05), 1.35) if source_column and len(headers) > 2 else 0
    remaining = max(available_width - source_width, 1.0)
    weights = [_docx_column_weight(header) for header in body_headers]
    total_weight = sum(weights) or 1
    widths = [remaining * weight / total_weight for weight in weights]
    if source_column:
        widths.append(source_width)
    return widths


def _docx_column_weight(header: str) -> float:
    normalized = header.strip().lower()
    if normalized in {"id", "ref", "date", "time", "status", "level", "score", "owner", "source", "confidence"}:
        return 0.75
    if any(token in normalized for token in ("finding", "summary", "description", "evidence", "fact", "rationale", "action", "recommendation")):
        return 2.4
    return 1.2


def _style_docx_table(document, table, headers: list[str], profile: ArtifactLayoutProfile) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    _set_docx_table_geometry(table, _docx_table_widths(headers, _docx_content_width_inches(document)))
    _mark_docx_repeat_header(table.rows[0])
    for cell in table.rows[0].cells:
        _set_docx_cell_shading(cell, profile.brand_hex)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _set_docx_paragraph_spacing(paragraph, after=0, line=1.0)
            _format_docx_paragraph_runs(paragraph, size=8.8, color="FFFFFF", bold=True)
    for row_index, row in enumerate(table.rows[1:], start=1):
        if row_index % 2 == 0:
            for cell in row.cells:
                _set_docx_cell_shading(cell, profile.row_hex)
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                _set_docx_paragraph_spacing(paragraph, after=0, line=1.05)
                _format_docx_paragraph_runs(paragraph, size=8.6, color=profile.text_hex)


def _append_docx_callout(
    document,
    text: str,
    evidence_ids: list[str],
    citations: dict[str, EvidenceCitation],
    profile: ArtifactLayoutProfile,
) -> None:
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    _set_docx_table_geometry(table, [_docx_content_width_inches(document)])
    cell = table.cell(0, 0)
    cell.text = _with_sources(text, evidence_ids, citations)
    _set_docx_cell_shading(cell, profile.callout_hex)
    for paragraph in cell.paragraphs:
        _set_docx_paragraph_spacing(paragraph, after=0, line=1.15)
        _format_docx_paragraph_runs(paragraph, size=10, color=profile.text_hex, bold=True)
    spacer = document.add_paragraph()
    _set_docx_paragraph_spacing(spacer, after=2)


def _append_docx_block(
    document,
    block: ContentBlock,
    citations: dict[str, EvidenceCitation],
    profile: ArtifactLayoutProfile,
) -> None:
    if block.kind == "heading":
        document.add_heading(block.text or "", level=block.level)
        return
    if block.kind in {"page_break", "section_break"}:
        document.add_page_break()
        return
    if block.kind == "paragraph":
        paragraph = document.add_paragraph(_with_sources(block.text or "", block.evidence_ids, citations))
        _set_docx_paragraph_spacing(paragraph, after=6, line=1.12)
        _format_docx_paragraph_runs(paragraph, size=10.5, color=profile.text_hex)
        return
    if block.kind in {"bullet_list", "numbered_list"}:
        style = "List Bullet" if block.kind == "bullet_list" else "List Number"
        for item_text, evidence_ids in _list_items(block):
            paragraph = document.add_paragraph(_with_sources(item_text, evidence_ids, citations), style=style)
            _set_docx_paragraph_spacing(paragraph, after=4, line=1.15)
            _format_docx_paragraph_runs(paragraph, size=10.2, color=profile.text_hex)
        return
    if block.kind == "quotation":
        paragraph = document.add_paragraph(_with_sources(block.text or "", block.evidence_ids, citations), style="Quote")
        _set_docx_paragraph_spacing(paragraph, after=6, line=1.12)
        _format_docx_paragraph_runs(paragraph, size=10, color=profile.muted_hex)
        return
    if block.kind == "callout":
        _append_docx_callout(document, block.text or "", block.evidence_ids, citations, profile)
        return
    if block.kind == "key_value":
        table = document.add_table(rows=0, cols=3)
        table.style = "Table Grid"
        header_cells = table.add_row().cells
        for index, header in enumerate(("Field", "Value", "Source")):
            header_cells[index].text = header
        for entry in block.entries:
            cells = table.add_row().cells
            cells[0].text = entry.key
            cells[1].text = entry.value
            cells[2].text = _source_summary(entry.evidence_ids, citations)
        _style_docx_table(document, table, ["Field", "Value", "Source"], profile)
        return
    if block.kind == "table" and block.table is not None:
        table = document.add_table(rows=1, cols=len(block.table.headers) + 1)
        table.style = "Table Grid"
        headers = [*block.table.headers, "Source"]
        for index, header in enumerate(headers):
            table.rows[0].cells[index].text = header
        for row in block.table.rows:
            cells = table.add_row().cells
            for index, value in enumerate(row.values):
                cells[index].text = value
            cells[-1].text = _source_summary(row.evidence_ids, citations)
        _style_docx_table(document, table, headers, profile)


def _render_pdf(bundle: ArtifactContentBundle, generated_at: datetime, profile: ArtifactLayoutProfile) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    wide = _has_wide_pdf_table(bundle)
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4) if wide else A4,
        title=bundle.paginated.title,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=22 * mm,
        bottomMargin=17 * mm,
    )
    styles = getSampleStyleSheet()
    _apply_pdf_style_defaults(styles, colors, profile)
    styles.add(ParagraphStyle(name="KickerV2", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=8.5, leading=10, textColor=colors.HexColor("#" + profile.accent_hex), spaceAfter=2))
    styles.add(ParagraphStyle(name="CenteredMetaV2", parent=styles["Normal"], alignment=TA_CENTER, textColor=colors.HexColor("#" + profile.muted_hex), fontSize=8.5, leading=11))
    styles.add(ParagraphStyle(name="CalloutV2", parent=styles["BodyText"], textColor=colors.HexColor("#" + profile.text_hex), fontSize=9.2, leading=12.5))
    styles.add(ParagraphStyle(name="QuoteV2", parent=styles["BodyText"], fontName="Helvetica-Oblique", textColor=colors.HexColor("#" + profile.muted_hex), leftIndent=10, rightIndent=6))
    styles.add(ParagraphStyle(name="ReferenceV2", parent=styles["BodyText"], fontSize=8, leading=10, textColor=colors.HexColor("#" + profile.muted_hex)))
    styles.add(ParagraphStyle(name="TableCellV2", parent=styles["BodyText"], fontSize=7.7, leading=9.5, wordWrap="CJK"))
    styles.add(ParagraphStyle(name="TableHeaderV2", parent=styles["TableCellV2"], textColor=colors.white, fontName="Helvetica-Bold", leading=9.8))
    citations = {item.evidence_id: item for item in bundle.content.citations}
    story = [Paragraph(escape(profile.kicker), styles["KickerV2"]), Paragraph(escape(bundle.paginated.title), styles["Title"])]
    if bundle.paginated.subtitle:
        story.append(Paragraph(escape(bundle.paginated.subtitle), styles["CenteredMetaV2"]))
    story.extend([
        Paragraph(escape(f"Generated {generated_at.strftime('%Y-%m-%d %H:%M UTC')}"), styles["CenteredMetaV2"]),
        Spacer(1, 14),
    ])
    if bundle.paginated.include_coverage_notes and bundle.content.warnings:
        story.append(Paragraph("Coverage Notes", styles["Heading2"]))
        story.extend(_pdf_list_flowables([(warning, []) for warning in bundle.content.warnings], citations, styles, bullet_type="bullet", profile=profile))
    for semantic_section in bundle.paginated.sections:
        story.append(Paragraph(escape(semantic_section.title), styles["Heading1"]))
        for block in semantic_section.blocks:
            story.extend(_pdf_block(block, citations, styles, colors, Table, TableStyle, Paragraph, Spacer, PageBreak, document.width, profile))
    if bundle.paginated.include_references:
        story.extend([PageBreak(), Paragraph("References", styles["Heading1"])])
        for index, citation in enumerate(bundle.content.citations, start=1):
            story.extend([
                Paragraph(escape(_reference_line(index, citation)), styles["ReferenceV2"]),
                Spacer(1, 3),
            ])
    document.build(
        story,
        onFirstPage=lambda canvas, doc: _pdf_footer(canvas, doc, profile),
        onLaterPages=lambda canvas, doc: _pdf_footer(canvas, doc, profile),
    )
    return buffer.getvalue()


def _apply_pdf_style_defaults(styles, colors, profile: ArtifactLayoutProfile) -> None:
    styles["Normal"].fontName = "Helvetica"
    styles["Normal"].fontSize = 9.6
    styles["Normal"].leading = 13
    styles["Normal"].textColor = colors.HexColor("#" + profile.text_hex)
    styles["BodyText"].fontName = "Helvetica"
    styles["BodyText"].fontSize = 9.6
    styles["BodyText"].leading = 13
    styles["BodyText"].spaceAfter = 2
    styles["BodyText"].textColor = colors.HexColor("#" + profile.text_hex)
    styles["Title"].fontName = "Helvetica-Bold"
    styles["Title"].fontSize = 22
    styles["Title"].leading = 26
    styles["Title"].spaceAfter = 7
    styles["Title"].textColor = colors.HexColor("#" + profile.brand_hex)
    for name, size, leading, before, after in (
        ("Heading1", 14, 17, 12, 7),
        ("Heading2", 11.5, 14, 9, 5),
        ("Heading3", 10.3, 12.5, 7, 4),
    ):
        style = styles[name]
        style.fontName = "Helvetica-Bold"
        style.fontSize = size
        style.leading = leading
        style.spaceBefore = before
        style.spaceAfter = after
        style.keepWithNext = 1
        style.textColor = colors.HexColor("#" + profile.brand_hex)


def _pdf_block(block, citations, styles, colors, table_cls, table_style_cls, paragraph_cls, spacer_cls, page_break_cls, available_width, profile: ArtifactLayoutProfile):
    if block.kind in {"page_break", "section_break"}:
        return [page_break_cls()]
    if block.kind == "heading":
        return [paragraph_cls(escape(block.text or ""), styles[f"Heading{min(block.level + 1, 3)}"])]
    if block.kind == "callout":
        return _pdf_callout_flowables(block.text or "", block.evidence_ids, citations, styles, colors, table_cls, table_style_cls, paragraph_cls, spacer_cls, available_width, profile)
    if block.kind in {"paragraph", "quotation"}:
        style = styles["QuoteV2"] if block.kind == "quotation" else styles["BodyText"]
        return [
            paragraph_cls(_pdf_with_sources_markup(block.text or "", block.evidence_ids, citations, profile), style),
            spacer_cls(1, 6),
        ]
    if block.kind in {"bullet_list", "numbered_list"}:
        bullet_type = "1" if block.kind == "numbered_list" else "bullet"
        return _pdf_list_flowables(_list_items(block), citations, styles, bullet_type=bullet_type, profile=profile)
    if block.kind == "key_value":
        rows = [["Field", "Value", "Source"]] + [
            [entry.key, entry.value, _source_label(entry.evidence_ids, citations)]
            for entry in block.entries
        ]
        return [_styled_pdf_table(rows, colors, table_cls, table_style_cls, styles, available_width, profile), spacer_cls(1, 9)]
    if block.kind == "table" and block.table is not None:
        rows = [[*block.table.headers, "Source"]] + [
            [*row.values, _source_label(row.evidence_ids, citations)]
            for row in block.table.rows
        ]
        return [_styled_pdf_table(rows, colors, table_cls, table_style_cls, styles, available_width, profile), spacer_cls(1, 9)]
    return []


def _pdf_callout_flowables(text, evidence_ids, citations, styles, colors, table_cls, table_style_cls, paragraph_cls, spacer_cls, available_width, profile: ArtifactLayoutProfile):
    callout = paragraph_cls(f"<b>Important</b><br/>{_pdf_with_sources_markup(text, evidence_ids, citations, profile)}", styles["CalloutV2"])
    table = table_cls([[callout]], colWidths=[available_width], hAlign="LEFT")
    table.setStyle(table_style_cls([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#" + profile.callout_hex)),
        ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#E3D4B3")),
        ("LINEBEFORE", (0, 0), (0, -1), 3, colors.HexColor("#" + profile.accent_hex)),
        ("LEFTPADDING", (0, 0), (-1, -1), 9),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    return [table, spacer_cls(1, 8)]


def _pdf_list_flowables(items: list[tuple[str, list[str]]], citations, styles, *, bullet_type: str, profile: ArtifactLayoutProfile):
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, Spacer

    list_items = [
        ListItem(Paragraph(_pdf_with_sources_markup(item_text, evidence_ids, citations, profile), styles["BodyText"]), leftIndent=0)
        for item_text, evidence_ids in items
        if item_text
    ]
    if not list_items:
        return []
    return [
        ListFlowable(list_items, bulletType=bullet_type, leftIndent=16, bulletIndent=4),
        Spacer(1, 6),
    ]


def _pdf_with_sources_markup(
    text: str,
    evidence_ids: list[str],
    citations: dict[str, EvidenceCitation],
    profile: ArtifactLayoutProfile,
) -> str:
    summary = _source_summary(evidence_ids, citations)
    if not summary:
        return escape(text)
    return f"{escape(text)} <font color=\"#{profile.muted_hex}\" size=\"8\">({escape(summary)})</font>"


def _source_label(evidence_ids: list[str], citations: dict[str, EvidenceCitation]) -> str:
    return _source_summary(evidence_ids, citations).removeprefix("Sources: ")


def _styled_pdf_table(rows, colors, table_cls, table_style_cls, styles, available_width, profile: ArtifactLayoutProfile):
    from reportlab.platypus import Paragraph

    normalized = _normalized_table_rows(rows)
    table_rows = [
        [
            Paragraph(escape(str(cell)), styles["TableHeaderV2" if row_index == 0 else "TableCellV2"])
            for cell in row
        ]
        for row_index, row in enumerate(normalized)
    ]
    table = table_cls(
        table_rows,
        colWidths=_pdf_col_widths(normalized, available_width),
        repeatRows=1,
        hAlign="LEFT",
        splitByRow=1,
    )
    table.setStyle(table_style_cls([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + profile.brand_hex)),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#" + profile.row_hex)]),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("LINEBELOW", (0, 0), (-1, 0), 0.9, colors.HexColor("#" + profile.accent_hex)),
        ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#" + profile.rule_hex)),
        ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#AAB7B8")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    return table


def _normalized_table_rows(rows) -> list[list[str]]:
    width = max((len(row) for row in rows), default=1)
    return [[str(cell) for cell in row] + [""] * (width - len(row)) for row in rows]


def _pdf_col_widths(rows: list[list[str]], available_width: float) -> list[float]:
    columns = max((len(row) for row in rows), default=1)
    if columns <= 1:
        return [available_width]
    has_source_column = rows and rows[0] and rows[0][-1].strip().lower() == "source"
    if has_source_column and columns > 2:
        source_width = min(available_width * 0.22, 95)
        body_width = (available_width - source_width) / (columns - 1)
        return [body_width] * (columns - 1) + [source_width]
    return [available_width / columns] * columns


def _render_pptx(
    bundle: ArtifactContentBundle,
    generated_at: datetime,
    profile: ArtifactLayoutProfile,
    *,
    progress_callback: Callable[[int, int, str], None] | None = None,
) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    citations = {item.evidence_id: item for item in bundle.content.citations}
    total_slides = _pptx_render_slide_count(bundle)
    rendered_slides = 0
    content_slide_index = 0

    def record_slide(label: str) -> None:
        nonlocal rendered_slides
        rendered_slides += 1
        if progress_callback:
            progress_callback(rendered_slides, total_slides, label)

    def add_content_slide(title: str):
        nonlocal content_slide_index
        slide = _add_base_content_slide(presentation, title, content_slide_index, profile)
        content_slide_index += 1
        return slide

    title_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _set_slide_background(title_slide, RGBColor(*profile.pptx_title_background))
    _decorate_title_slide(title_slide, bundle, generated_at, profile)
    record_slide("Title slide")
    for slide_spec in _content_slides(bundle.presentation.slides):
        table_block = next((block for block in slide_spec.blocks if block.kind == "table" and block.table), None)
        if table_block is not None:
            chunks = _table_block_chunks(table_block)
            for chunk_index, chunk in enumerate(chunks, start=1):
                title = slide_spec.title if len(chunks) == 1 else f"{slide_spec.title} ({chunk_index}/{len(chunks)})"
                slide = add_content_slide(title)
                _add_slide_table(slide, chunk, profile)
                _add_slide_sources(slide, _block_evidence_ids(chunk), citations, profile)
                record_slide(title)
            continue

        evidence_ids = list(dict.fromkeys(
            evidence_id
            for block in slide_spec.blocks
            for evidence_id in _block_evidence_ids(block)
        ))
        line_chunks = _slide_line_chunks(_slide_lines(slide_spec.blocks) or ["No supported content was generated."])
        for chunk_index, line_chunk in enumerate(line_chunks, start=1):
            title = slide_spec.title if len(line_chunks) == 1 else f"{slide_spec.title} ({chunk_index}/{len(line_chunks)})"
            slide = add_content_slide(title)
            _add_slide_bullets(slide, line_chunk, profile)
            _add_slide_sources(slide, evidence_ids, citations, profile)
            record_slide(title)
    if bundle.presentation.include_references_slide:
        slide = _add_base_content_slide(presentation, "References", content_slide_index, profile)
        _add_reference_lines(
            slide,
            [_reference_line(index, citation) for index, citation in enumerate(bundle.content.citations[:18], start=1)],
            profile,
        )
        record_slide("References")
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _pptx_render_slide_count(bundle: ArtifactContentBundle) -> int:
    total = 1
    for slide_spec in _content_slides(bundle.presentation.slides):
        table_block = next((block for block in slide_spec.blocks if block.kind == "table" and block.table), None)
        if table_block is not None:
            total += len(_table_block_chunks(table_block))
        else:
            total += len(_slide_line_chunks(_slide_lines(slide_spec.blocks) or ["No supported content was generated."]))
    if bundle.presentation.include_references_slide:
        total += 1
    return max(total, 1)


def _decorate_title_slide(
    slide,
    bundle: ArtifactContentBundle,
    generated_at: datetime,
    profile: ArtifactLayoutProfile,
) -> None:
    title_background = _rgb_tuple_hex(profile.pptx_title_background)
    panel_hex = _scale_hex(title_background, 0.82)
    _add_slide_rect(slide, 0, 0, 0.18, 7.5, profile.accent_hex)
    _add_slide_rect(slide, 9.25, 0, 4.08, 7.5, panel_hex)
    _add_slide_rect(slide, 0.85, 0.85, 1.0, 0.05, profile.accent_hex)
    _add_slide_text(slide, profile.kicker, 0.85, 0.98, 5.2, 0.3, 11, profile.accent_hex, bold=True)
    _add_slide_text(
        slide,
        bundle.presentation.title,
        0.82,
        1.45,
        7.85,
        1.7,
        _title_font_size(bundle.presentation.title),
        "FFFFFF",
        bold=True,
    )
    if bundle.presentation.subtitle:
        _add_slide_text(slide, bundle.presentation.subtitle, 0.88, 3.22, 7.2, 0.82, 15, "D6E4E5")
    _add_slide_text(slide, "Prudentia AI", 9.75, 0.88, 2.7, 0.32, 11, "FFFFFF", bold=True)
    _add_slide_text(slide, profile.label, 9.75, 1.32, 2.7, 0.28, 9, "D6E4E5")
    _add_title_metric(slide, "Generated", generated_at.date().isoformat(), 9.75, 2.25, profile)
    _add_title_metric(slide, "Sources", str(len(bundle.content.citations)), 9.75, 3.22, profile)
    _add_title_metric(slide, "Sections", str(len(bundle.content.sections)), 9.75, 4.19, profile)
    _add_slide_text(slide, "Evidence-backed artifact", 9.75, 6.52, 2.6, 0.25, 8, "AFC7C9")


def _add_title_metric(slide, label: str, value: str, left: float, top: float, profile: ArtifactLayoutProfile) -> None:
    _add_slide_rect(slide, left, top, 2.55, 0.62, _scale_hex(_rgb_tuple_hex(profile.pptx_title_background), 0.92), profile.accent_hex)
    _add_slide_text(slide, label.upper(), left + 0.18, top + 0.12, 1.1, 0.16, 7, "AFC7C9", bold=True)
    _add_slide_text(slide, value, left + 1.42, top + 0.1, 0.85, 0.26, 14, "FFFFFF", bold=True)


def _add_base_content_slide(presentation, title: str, content_slide_index: int, profile: ArtifactLayoutProfile):
    from pptx.dml.color import RGBColor

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _set_slide_background(slide, RGBColor(*_content_background(content_slide_index, profile)))
    _add_slide_rect(slide, 0, 0, 0.15, 7.5, profile.brand_hex)
    _add_slide_rect(slide, 0.15, 0, 0.045, 7.5, profile.accent_hex)
    _add_slide_text(slide, profile.kicker, 0.65, 0.18, 4.0, 0.22, 8, profile.accent_hex, bold=True)
    _add_slide_text(slide, title, 0.65, 0.42, 10.75, 0.62, _heading_font_size(title), profile.brand_hex, bold=True)
    _add_slide_text(slide, f"{content_slide_index + 2:02d}", 12.08, 0.32, 0.48, 0.24, 9, profile.muted_hex, bold=True)
    _add_slide_rect(
        slide,
        PPTX_BODY_PANEL_LEFT_IN,
        PPTX_BODY_PANEL_TOP_IN,
        PPTX_BODY_PANEL_WIDTH_IN,
        PPTX_BODY_PANEL_HEIGHT_IN,
        "FFFFFF",
        profile.rule_hex,
    )
    _add_slide_rect(slide, PPTX_BODY_PANEL_LEFT_IN, PPTX_BODY_PANEL_TOP_IN, 0.08, PPTX_BODY_PANEL_HEIGHT_IN, profile.accent_hex)
    _add_slide_rect(slide, 0.68, 6.71, 11.95, 0.015, profile.rule_hex)
    return slide


def _content_background(index: int, profile: ArtifactLayoutProfile) -> tuple[int, int, int]:
    return profile.pptx_backgrounds[index % len(profile.pptx_backgrounds)]


def _content_slides(slides):
    return [slide for slide in slides if not _is_generated_references_slide(slide)]


def _is_generated_references_slide(slide) -> bool:
    return _normalized_reference_title(slide.title) in _REFERENCE_SLIDE_TITLES


def _normalized_reference_title(value: str) -> str:
    return re.sub(r"[^a-z]+", " ", value.lower()).strip()


def _add_slide_sources(
    slide,
    evidence_ids: list[str],
    citations: dict[str, EvidenceCitation],
    profile: ArtifactLayoutProfile,
) -> None:
    source_text = _source_summary(list(dict.fromkeys(evidence_ids)), citations)
    if source_text:
        _add_slide_text(slide, source_text, 0.78, 6.86, 10.9, 0.27, 7.6, profile.muted_hex)


def _set_slide_background(slide, color) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_slide_rect(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    fill_hex: str,
    line_hex: str | None = None,
):
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill_hex)
    if line_hex:
        shape.line.color.rgb = RGBColor.from_string(line_hex)
    else:
        shape.line.fill.background()
    return shape


def _add_slide_text(slide, text, left, top, width, height, size, color, *, bold=False) -> None:
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor.from_string(color)


def _title_font_size(text: str) -> int:
    length = len(_compact(text, 180))
    if length <= 45:
        return 38
    if length <= 70:
        return 32
    if length <= 95:
        return 27
    return 23


def _heading_font_size(text: str) -> int:
    return 25 if len(text) <= 60 else 21


def _add_slide_bullets(slide, lines: list[str], profile: ArtifactLayoutProfile) -> None:
    display_lines = lines or ["No supported content was generated."]
    max_length = max((len(line) for line in display_lines), default=0)
    row_gap = 0.08 if len(display_lines) <= 6 else 0.05
    available_height = 4.78
    font_size = 14.2 if len(display_lines) <= 4 and max_length <= 220 else 12.8 if len(display_lines) <= 6 else 11.4
    row_heights = _pptx_bullet_row_heights(display_lines, font_size, available_height, row_gap)
    top = 1.42
    left = 0.95
    width = 11.18
    row_top = top
    for index, line in enumerate(display_lines):
        row_height = row_heights[index]
        if len(display_lines) <= 8:
            _add_slide_rect(slide, left, row_top, width, row_height, profile.row_hex if index % 2 else "FFFFFF", profile.rule_hex)
        _add_slide_rect(slide, left + 0.18, row_top + 0.12, 0.055, max(0.22, row_height - 0.24), profile.accent_hex)
        _add_slide_text(slide, f"{index + 1:02d}", left + 0.36, row_top + 0.13, 0.38, 0.18, 7.8, profile.muted_hex, bold=True)
        limit = 360 if row_height >= 0.72 else 285 if row_height >= 0.55 else 220
        _add_slide_text(
            slide,
            _compact(line, limit),
            left + 0.82,
            row_top + 0.1,
            width - 1.05,
            max(0.22, row_height - 0.1),
            font_size,
            profile.text_hex,
        )
        row_top += row_height + row_gap


def _add_reference_lines(slide, lines: list[str], profile: ArtifactLayoutProfile) -> None:
    if not lines:
        _add_slide_text(slide, "No references were available.", 0.95, 1.48, 10.8, 0.35, 12, profile.text_hex)
        return
    columns = (lines[:9], lines[9:18])
    for column_index, column_lines in enumerate(columns):
        left = 0.95 + column_index * 5.72
        for row_index, line in enumerate(column_lines):
            top = 1.46 + row_index * 0.49
            _add_slide_text(slide, _compact(line, 155), left, top, 5.15, 0.35, 8.1, profile.text_hex)


def _pptx_bullet_row_heights(lines: list[str], font_size: float, available_height: float, row_gap: float) -> list[float]:
    if not lines:
        return [available_height]
    usable_height = max(0.44, available_height - row_gap * max(0, len(lines) - 1))
    line_counts = [
        max(1, _pptx_estimated_line_count(line, 9.7, font_size))
        for line in lines
    ]
    desired = [
        min(1.78, max(0.48, 0.34 + line_count * (font_size / 72 * 1.55)))
        for line_count in line_counts
    ]
    total_desired = sum(desired)
    if total_desired <= usable_height:
        return desired
    scale = usable_height / total_desired if total_desired else 1
    scaled = [max(0.44, height * scale) for height in desired]
    overflow = sum(scaled) - usable_height
    if overflow > 0:
        reducible = [
            (index, height - 0.44)
            for index, height in enumerate(scaled)
            if height > 0.44
        ]
        capacity = sum(amount for _index, amount in reducible)
        if capacity > 0:
            for index, amount in reducible:
                scaled[index] -= overflow * (amount / capacity)
    return scaled


def _rgb_tuple_hex(value: tuple[int, int, int]) -> str:
    return "".join(f"{part:02X}" for part in value)


def _scale_hex(value: str, factor: float) -> str:
    normalized = value.strip().lstrip("#")
    if len(normalized) != 6:
        return value
    channels = [
        max(0, min(255, int(int(normalized[index : index + 2], 16) * factor)))
        for index in (0, 2, 4)
    ]
    return "".join(f"{channel:02X}" for channel in channels)


def _add_slide_table(slide, block: ContentBlock, profile: ArtifactLayoutProfile) -> None:
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    assert block.table is not None
    body_rows = block.table.rows
    rows = len(body_rows) + 1
    cols = len(block.table.headers)
    layout = _pptx_table_layout(block, fit_to_frame=True)
    shape = slide.shapes.add_table(
        rows,
        cols,
        Inches(PPTX_TABLE_LEFT_IN),
        Inches(PPTX_TABLE_TOP_IN),
        Inches(PPTX_TABLE_WIDTH_IN),
        Inches(layout.height),
    )
    table = shape.table
    for column, width in enumerate(layout.column_widths):
        table.columns[column].width = Inches(width)
    table.rows[0].height = Inches(layout.header_height)
    for column, header in enumerate(block.table.headers):
        cell = table.cell(0, column)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(profile.brand_hex)
        cell.margin_left = Inches(0.1)
        cell.margin_right = Inches(0.1)
        cell.margin_top = Inches(0.05)
        cell.margin_bottom = Inches(0.05)
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.bold = True
        run.font.size = Pt(10 if cols <= 4 else 9)
    for row_index, row in enumerate(body_rows, start=1):
        table.rows[row_index].height = Inches(layout.row_heights[row_index - 1])
        for column, value in enumerate(row.values):
            cell = table.cell(row_index, column)
            cell.text = _compact(value, _pptx_cell_text_limit(layout.column_widths[column], layout.body_font_size))
            cell.margin_left = Inches(0.1)
            cell.margin_right = Inches(0.1)
            cell.margin_top = Inches(0.05)
            cell.margin_bottom = Inches(0.05)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(profile.row_hex if row_index % 2 == 1 else "FFFFFF")
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = "Aptos"
                paragraph.font.size = Pt(layout.body_font_size)
                paragraph.font.color.rgb = RGBColor.from_string(profile.text_hex)


def _slide_lines(blocks: list[ContentBlock]) -> list[str]:
    lines: list[str] = []
    for block in blocks:
        if block.kind in {"paragraph", "quotation", "callout", "heading"} and block.text:
            lines.append(block.text)
        elif block.kind in {"bullet_list", "numbered_list"}:
            lines.extend(item_text for item_text, _evidence_ids in _list_items(block))
        elif block.kind == "key_value":
            lines.extend(f"{entry.key}: {entry.value}" for entry in block.entries)
    return lines


def _block_evidence_ids(block: ContentBlock) -> list[str]:
    ids = list(block.evidence_ids)
    ids.extend(evidence_id for item in block.list_items for evidence_id in item.evidence_ids)
    ids.extend(evidence_id for entry in block.entries for evidence_id in entry.evidence_ids)
    if block.table is not None:
        ids.extend(evidence_id for row in block.table.rows for evidence_id in row.evidence_ids)
    return ids


def _list_items(block: ContentBlock) -> list[tuple[str, list[str]]]:
    if block.list_items:
        return [(item.text, item.evidence_ids) for item in block.list_items]
    return [(item, block.evidence_ids) for item in block.items]


def _table_block_chunks(block: ContentBlock, *, max_body_rows: int | None = None) -> list[ContentBlock]:
    if block.table is None:
        return [block]
    rows = block.table.rows
    if not rows:
        return [block]
    chunks: list[ContentBlock] = []
    current: list = []
    for row in rows:
        if max_body_rows is not None and current and len(current) >= max_body_rows:
            chunks.append(block.model_copy(update={"table": block.table.model_copy(update={"rows": current})}))
            current = []
        candidate = [*current, row]
        candidate_block = block.model_copy(update={"table": block.table.model_copy(update={"rows": candidate})})
        candidate_height = _pptx_table_layout(candidate_block, fit_to_frame=False).height
        if current and candidate_height > PPTX_TABLE_MAX_HEIGHT_IN:
            chunks.append(block.model_copy(update={"table": block.table.model_copy(update={"rows": current})}))
            current = [row]
        else:
            current = candidate
    if current:
        chunks.append(block.model_copy(update={"table": block.table.model_copy(update={"rows": current})}))
    return chunks or [block]


def _pptx_table_layout(block: ContentBlock, *, fit_to_frame: bool) -> PptxTableLayout:
    assert block.table is not None
    headers = list(block.table.headers)
    rows = block.table.rows
    column_widths = _pptx_table_col_widths(headers, rows)
    body_font_size = _pptx_table_body_font_size(headers, rows)
    header_height = max(
        PPTX_TABLE_HEADER_MIN_HEIGHT_IN,
        0.22 + 0.18 * max(
            _pptx_estimated_line_count(header, column_widths[index], 10)
            for index, header in enumerate(headers)
        ),
    )
    row_heights = [
        _pptx_table_row_height(row.values, column_widths, body_font_size)
        for row in rows
    ]
    layout = PptxTableLayout(
        column_widths=column_widths,
        row_heights=row_heights,
        header_height=min(header_height, 0.72),
        body_font_size=body_font_size,
    )
    return _fit_pptx_table_layout(layout) if fit_to_frame else layout


def _fit_pptx_table_layout(layout: PptxTableLayout) -> PptxTableLayout:
    if layout.height <= PPTX_TABLE_MAX_HEIGHT_IN or not layout.row_heights:
        return layout
    available_body_height = max(PPTX_TABLE_MAX_HEIGHT_IN - layout.header_height, PPTX_TABLE_ROW_MIN_HEIGHT_IN)
    total_body_height = sum(layout.row_heights)
    scale = available_body_height / total_body_height if total_body_height else 1
    fitted_rows = [max(PPTX_TABLE_ROW_MIN_HEIGHT_IN, height * scale) for height in layout.row_heights]
    if layout.header_height + sum(fitted_rows) > PPTX_TABLE_MAX_HEIGHT_IN:
        overflow = layout.header_height + sum(fitted_rows) - PPTX_TABLE_MAX_HEIGHT_IN
        reducible = [
            (index, height - PPTX_TABLE_ROW_MIN_HEIGHT_IN)
            for index, height in enumerate(fitted_rows)
            if height > PPTX_TABLE_ROW_MIN_HEIGHT_IN
        ]
        total_reducible = sum(value for _index, value in reducible)
        if total_reducible > 0:
            for index, value in reducible:
                fitted_rows[index] -= overflow * (value / total_reducible)
    font_size = layout.body_font_size if scale > 0.9 else max(PPTX_TABLE_BODY_FONT_MIN, layout.body_font_size - 0.5)
    return PptxTableLayout(
        column_widths=layout.column_widths,
        row_heights=fitted_rows,
        header_height=layout.header_height,
        body_font_size=font_size,
    )


def _pptx_table_col_widths(headers: list[str], rows) -> list[float]:
    if not headers:
        return [PPTX_TABLE_WIDTH_IN]
    values_by_column = [
        [
            str(row.values[index]) if index < len(row.values) else ""
            for row in rows
        ]
        for index in range(len(headers))
    ]
    weights = [
        _pptx_column_weight(header, values_by_column[index])
        for index, header in enumerate(headers)
    ]
    total_weight = sum(weights) or 1
    widths = [PPTX_TABLE_WIDTH_IN * weight / total_weight for weight in weights]
    minimums = [_pptx_column_min_width(header) for header in headers]
    if sum(minimums) > PPTX_TABLE_WIDTH_IN:
        scale = PPTX_TABLE_WIDTH_IN / sum(minimums)
        minimums = [width * scale for width in minimums]
    maximums = [_pptx_column_max_width(header, len(headers)) for header in headers]
    maximums = [max(maximum, minimums[index]) for index, maximum in enumerate(maximums)]
    return _redistribute_widths(widths, minimums, maximums, PPTX_TABLE_WIDTH_IN)


def _pptx_column_weight(header: str, values: list[str]) -> float:
    normalized = header.strip().lower()
    lengths = [len(_normalized_cell_text(header)), *(len(_normalized_cell_text(value)) for value in values)]
    average = sum(min(length, 160) for length in lengths) / max(len(lengths), 1)
    longest = max(lengths, default=0)
    content_weight = 0.55 + average / 34 + min(longest, 220) / 115
    if normalized in {"id", "ref", "date", "time", "status", "score", "level", "owner"}:
        return max(0.75, content_weight * 0.72)
    if any(token in normalized for token in ("event", "finding", "summary", "description", "evidence", "item", "rationale", "action", "recommendation", "note")):
        return max(1.75, content_weight * 1.25)
    return max(1.0, content_weight)


def _pptx_column_min_width(header: str) -> float:
    normalized = header.strip().lower()
    if normalized in {"id", "ref", "date", "time", "score", "level"}:
        return 0.72
    if normalized in {"status", "owner", "confidence"}:
        return 1.05
    return 1.25


def _pptx_column_max_width(header: str, column_count: int) -> float:
    normalized = header.strip().lower()
    if normalized in {"id", "ref", "date", "time", "score", "level"}:
        return 1.6
    if column_count <= 2:
        return PPTX_TABLE_WIDTH_IN * 0.72
    if any(token in normalized for token in ("event", "description", "summary", "finding", "recommendation", "note")):
        return PPTX_TABLE_WIDTH_IN * 0.58
    return PPTX_TABLE_WIDTH_IN * 0.45


def _redistribute_widths(widths: list[float], minimums: list[float], maximums: list[float], total_width: float) -> list[float]:
    adjusted = [min(max(widths[index], minimums[index]), maximums[index]) for index in range(len(widths))]
    for _iteration in range(8):
        delta = total_width - sum(adjusted)
        if abs(delta) < 0.01:
            break
        if delta > 0:
            candidates = [(index, maximums[index] - adjusted[index]) for index in range(len(adjusted)) if adjusted[index] < maximums[index]]
        else:
            candidates = [(index, adjusted[index] - minimums[index]) for index in range(len(adjusted)) if adjusted[index] > minimums[index]]
        capacity = sum(value for _index, value in candidates)
        if capacity <= 0:
            break
        for index, value in candidates:
            adjusted[index] += delta * (value / capacity)
            adjusted[index] = min(max(adjusted[index], minimums[index]), maximums[index])
    if adjusted:
        adjusted[-1] += total_width - sum(adjusted)
    return adjusted


def _pptx_table_body_font_size(headers: list[str], rows) -> float:
    cell_lengths = [
        len(_normalized_cell_text(value))
        for row in rows
        for value in row.values
    ]
    longest = max(cell_lengths, default=0)
    row_count = len(rows)
    column_count = len(headers)
    font_size = PPTX_TABLE_BODY_FONT_MAX
    if longest > 160 or row_count > 7:
        font_size -= 1.3
    elif longest > 95 or row_count > 5:
        font_size -= 0.6
    if column_count >= 5:
        font_size -= 0.6
    return max(PPTX_TABLE_BODY_FONT_MIN, font_size)


def _pptx_table_row_height(values: list[str], column_widths: list[float], font_size: float) -> float:
    line_count = max(
        (
            _pptx_estimated_line_count(value, column_widths[index], font_size)
            for index, value in enumerate(values)
            if index < len(column_widths)
        ),
        default=1,
    )
    height = 0.20 + line_count * (font_size / 72 * 1.28)
    return min(max(height, PPTX_TABLE_ROW_MIN_HEIGHT_IN), PPTX_TABLE_ROW_MAX_HEIGHT_IN)


def _pptx_estimated_line_count(text: str, column_width: float, font_size: float) -> int:
    normalized = _normalized_cell_text(text)
    if not normalized:
        return 1
    chars_per_line = max(6, int(column_width * 13.0 * (9.0 / max(font_size, 1))))
    lines = 1
    current = 0
    for word in normalized.split(" "):
        length = len(word)
        if length > chars_per_line:
            extra_lines = (length + chars_per_line - 1) // chars_per_line
            if current:
                lines += 1
                current = 0
            lines += max(0, extra_lines - 1)
            current = length % chars_per_line
            continue
        next_length = length if current == 0 else current + 1 + length
        if next_length <= chars_per_line:
            current = next_length
        else:
            lines += 1
            current = length
    return lines


def _pptx_cell_text_limit(column_width: float, font_size: float) -> int:
    lines_that_fit = max(2, int(PPTX_TABLE_ROW_MAX_HEIGHT_IN / (font_size / 72 * 1.28)))
    chars_per_line = max(10, int(column_width * 13.0 * (9.0 / max(font_size, 1))))
    return max(80, min(360, chars_per_line * lines_that_fit))


def _normalized_cell_text(value: str) -> str:
    return re.sub(r"\s+", " ", str(value)).strip()


def _slide_line_chunks(lines: list[str]) -> list[list[str]]:
    chunks: list[list[str]] = []
    current: list[str] = []
    current_units = 0
    for line in lines:
        units = _pptx_text_units(line)
        if current and (
            len(current) >= PPTX_MAX_BULLETS_PER_SLIDE
            or current_units + units > PPTX_MAX_TEXT_UNITS_PER_SLIDE
        ):
            chunks.append(current)
            current = []
            current_units = 0
        current.append(line)
        current_units += units
    if current:
        chunks.append(current)
    return chunks or [[]]


def _pptx_text_units(text: str) -> int:
    return max(1, (len(text.strip()) + PPTX_TEXT_UNIT_CHARS - 1) // PPTX_TEXT_UNIT_CHARS)


def _with_sources(text: str, evidence_ids: list[str], citations: dict[str, EvidenceCitation]) -> str:
    summary = _source_summary(evidence_ids, citations)
    return f"{text} ({summary})" if summary else text


def _source_summary(evidence_ids: list[str], citations: dict[str, EvidenceCitation]) -> str:
    labels = list(dict.fromkeys([
        _short_citation(citations[evidence_id])
        for evidence_id in evidence_ids
        if evidence_id in citations
    ]))
    if len(labels) > PPTX_MAX_SOURCE_LABELS:
        labels = [*labels[:PPTX_MAX_SOURCE_LABELS], f"+{len(labels) - PPTX_MAX_SOURCE_LABELS} more"]
    return "Sources: " + "; ".join(labels) if labels else ""


def _short_citation(citation: EvidenceCitation) -> str:
    return f"{citation.doc_title}, p. {citation.page_start}" if citation.page_start else citation.doc_title


def _reference_line(index: int, citation: EvidenceCitation) -> str:
    page = ""
    if citation.page_start is not None and citation.page_end not in (None, citation.page_start):
        page = f", pages {citation.page_start}-{citation.page_end}"
    elif citation.page_start is not None:
        page = f", page {citation.page_start}"
    return f"[{index}] {citation.doc_title}{page} ({citation.doc_id}:{citation.chunk_id})"


def _pdf_footer(canvas, document, profile: ArtifactLayoutProfile) -> None:
    from reportlab.lib import colors
    from reportlab.lib.units import mm

    canvas.saveState()
    width, height = document.pagesize
    left = document.leftMargin
    right = width - document.rightMargin
    canvas.setStrokeColor(colors.HexColor("#" + profile.rule_hex))
    canvas.setLineWidth(0.6)
    canvas.line(left, height - 16 * mm, right, height - 16 * mm)
    canvas.line(left, 12 * mm, right, 12 * mm)
    canvas.setFont("Helvetica-Bold", 8.5)
    canvas.setFillColor(colors.HexColor("#" + profile.brand_hex))
    canvas.drawString(left, height - 12.4 * mm, "Prudentia AI")
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#" + profile.muted_hex))
    canvas.drawCentredString(width / 2, 8 * mm, profile.label)
    canvas.drawRightString(right, 8 * mm, f"Page {document.page}")
    canvas.restoreState()


def _validate_package(artifact_format: ArtifactFormat, content: bytes) -> None:
    if len(content) < 100:
        raise RuntimeError(f"{artifact_format} renderer returned an empty package")
    if artifact_format == "pdf":
        if not content.startswith(b"%PDF"):
            raise RuntimeError("PDF renderer returned an invalid document")
        return
    required = "word/document.xml" if artifact_format == "docx" else "ppt/presentation.xml"
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            if required not in archive.namelist():
                raise RuntimeError(f"{artifact_format} package is missing {required}")
            archive.testzip()
    except zipfile.BadZipFile as exc:
        raise RuntimeError(f"{artifact_format} renderer returned an invalid OpenXML package") from exc


def _office_smoke_check(
    *,
    artifact_format: ArtifactFormat,
    content: bytes,
    require_libreoffice: bool,
) -> tuple[str, ...]:
    if artifact_format == "pdf":
        return ()
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if executable is None:
        if require_libreoffice:
            raise RuntimeError("LibreOffice is required for generated artifact smoke checks")
        return ("LibreOffice smoke check skipped because LibreOffice is unavailable.",)
    with tempfile.TemporaryDirectory(prefix="artifact-smoke-") as temp_dir:
        source = Path(temp_dir) / f"artifact.{artifact_format}"
        source.write_bytes(content)
        result = subprocess.run(
            [executable, "--headless", "--convert-to", "pdf", "--outdir", temp_dir, str(source)],
            capture_output=True,
            text=True,
            timeout=120,
            check=False,
        )
        converted = Path(temp_dir) / "artifact.pdf"
        if result.returncode != 0 or not converted.exists() or converted.stat().st_size < 500:
            message = f"LibreOffice smoke check failed: {(result.stderr or result.stdout)[-500:]}"
            if require_libreoffice:
                raise RuntimeError(message)
            return (message,)
        pdftotext = shutil.which("pdftotext")
        if pdftotext:
            text_result = subprocess.run(
                [pdftotext, "-layout", str(converted), "-"],
                capture_output=True,
                text=True,
                timeout=60,
                check=False,
            )
            pages = text_result.stdout.split("\f")
            if any(not page.strip() for page in pages[:-1]):
                message = "LibreOffice smoke check found a blank output page"
                if require_libreoffice:
                    raise RuntimeError(message)
                return (message,)
    return ()


def _safe_filename(value: str) -> str:
    candidate = re.sub(r"[^A-Za-z0-9._-]+", "-", value.strip()).strip(".-")
    return _compact(candidate or "evidence-artifact", 80).strip(".-") or "evidence-artifact"


def _compact(value: str, limit: int) -> str:
    normalized = re.sub(r"\s+", " ", value).strip()
    return normalized if len(normalized) <= limit else normalized[: max(0, limit - 3)].rstrip() + "..."
