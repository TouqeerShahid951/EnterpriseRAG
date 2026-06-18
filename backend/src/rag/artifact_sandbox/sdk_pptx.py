"""PPTX layout primitives for sandbox build programs."""

from __future__ import annotations

from datetime import datetime
from typing import Any

from rag.artifact_jobs.contracts import EvidenceCitation

from .sdk_shared import citation_map, compact, item_parts, normalize_evidence_ids, reference_line, row_parts, short_citation, validated_output_path

MAX_TEXT_UNITS_PER_SLIDE = 10
MAX_TABLE_ROWS_PER_SLIDE = 8
TITLE_SLIDE_BACKGROUND = "173B3F"
CONTENT_SLIDE_BACKGROUNDS = ("DCEBE8", "F6E7CB", "E2E8F5", "F1E0DC", "E5EBD5")
CONTENT_SLIDE_BACKGROUND = CONTENT_SLIDE_BACKGROUNDS[0]
REFERENCE_SLIDE_TITLES = {"reference", "references", "source", "sources", "citation", "citations", "bibliography"}


class PptxArtifact:
    def __init__(
        self,
        *,
        path: str,
        title: str,
        subtitle: str | None,
        generated_at: datetime,
        citations: list[EvidenceCitation],
    ) -> None:
        from pptx import Presentation
        from pptx.util import Inches

        self._path = validated_output_path(path)
        self._citations = citation_map(citations)
        self._all_citations = list(citations)
        self._presentation = Presentation()
        self._presentation.slide_width = Inches(13.333)
        self._presentation.slide_height = Inches(7.5)
        self._slide = None
        self._frame = None
        self._base_slide_title = "Details"
        self._text_units = 0
        self._has_table = False
        self._content_slide_index = 0
        self._current_slide_title = ""
        self._slide_source_ids: list[str] = []
        self._source_footer_shape = None
        self.add_title_slide(title, subtitle or "", footer=generated_at.date().isoformat())

    def add_title_slide(self, title: str, subtitle: str = "", *, footer: str = "") -> None:
        slide = self._presentation.slides.add_slide(self._presentation.slide_layouts[6])
        self._set_background(slide, TITLE_SLIDE_BACKGROUND)
        self._add_text(slide, title, 0.8, 1.55, 11.7, 0.8, 30, "FFFFFF", bold=True)
        self._add_text(slide, subtitle, 0.85, 2.65, 10.8, 0.7, 17, "D6E4E5")
        self._add_text(slide, footer, 0.85, 6.6, 3.0, 0.25, 10, "AFC7C9")

    def add_slide(self, title: str | None = None, **_layout_hints: Any) -> None:
        slide_title = str(title or "").strip() or self._base_slide_title or "Details"
        slide = self._presentation.slides.add_slide(self._presentation.slide_layouts[6])
        self._set_background(slide, self._next_content_background())
        self._add_text(slide, slide_title, 0.65, 0.3, 12.0, 0.55, 23, "173B3F", bold=True)
        self._slide = slide
        self._frame = None
        self._base_slide_title = slide_title
        self._current_slide_title = slide_title
        self._text_units = 0
        self._has_table = False
        self._slide_source_ids = []
        self._source_footer_shape = None

    def add_heading(self, text: str, *, level: int = 1) -> None:
        self.add_slide(str(text))

    def add_paragraph(self, text: str, *, evidence_ids: list[str] | None = None) -> None:
        paragraph_text = compact(text, 360)
        units = self._paragraph_units(paragraph_text)
        if self._has_table or self._text_units + units > MAX_TEXT_UNITS_PER_SLIDE:
            self._continue_slide()
        frame = self._text_frame()
        paragraph = frame.paragraphs[0] if self._text_units == 0 else frame.add_paragraph()
        paragraph.text = paragraph_text
        paragraph.font.size = self._pt(14)
        paragraph.space_after = self._pt(7)
        self._text_units += units
        self._update_source_footer(evidence_ids)

    def add_bullet_list(self, items: list[Any], *, evidence_ids: list[str] | None = None) -> None:
        for item in items:
            text, ids = item_parts(item, evidence_ids)
            self.add_paragraph("- " + text, evidence_ids=ids)

    def add_numbered_list(self, items: list[Any], *, evidence_ids: list[str] | None = None) -> None:
        for index, item in enumerate(items, start=1):
            text, ids = item_parts(item, evidence_ids)
            self.add_paragraph(f"{index}. {text}", evidence_ids=ids)

    def add_callout(self, text: str, *, evidence_ids: list[str] | None = None) -> None:
        self.add_paragraph("Key point: " + str(text), evidence_ids=evidence_ids)

    def add_table(self, headers: list[str], rows: list[Any], *, evidence_ids: list[str] | None = None) -> None:
        from pptx.dml.color import RGBColor
        from pptx.util import Inches

        body_rows = list(rows)
        chunks = [
            body_rows[offset : offset + MAX_TABLE_ROWS_PER_SLIDE]
            for offset in range(0, len(body_rows), MAX_TABLE_ROWS_PER_SLIDE)
        ] or [[]]
        base_title = self._base_slide_title or "Details"
        for chunk_index, body in enumerate(chunks, start=1):
            if chunk_index > 1 or self._has_table or self._text_units:
                title = base_title if len(chunks) == 1 else f"{base_title} ({chunk_index}/{len(chunks)})"
                self.add_slide(title)
            else:
                self._ensure_slide()
            shape = self._slide.shapes.add_table(len(body) + 1, len(headers), Inches(0.55), Inches(1.15), Inches(12.2), Inches(5.4))
            table = shape.table
            for column, header in enumerate(headers):
                cell = table.cell(0, column)
                cell.text = str(header)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(23, 59, 63)
                run = cell.text_frame.paragraphs[0].runs[0]
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.bold = True
                run.font.size = self._pt(9 if len(headers) <= 5 else 8)
            chunk_evidence_ids: list[str] = []
            for row_index, row in enumerate(body, start=1):
                values, ids = row_parts(row, evidence_ids)
                chunk_evidence_ids.extend(ids)
                body_values = [*values[: len(headers)], *([""] * max(0, len(headers) - len(values)))]
                for column, value in enumerate(body_values):
                    cell = table.cell(row_index, column)
                    cell.text = compact(value, 120)
                    for paragraph in cell.text_frame.paragraphs:
                        paragraph.font.size = self._pt(8 if len(headers) > 5 or len(body) > 6 else 9)
            self._has_table = True
            self._text_units = MAX_TEXT_UNITS_PER_SLIDE
            self._update_source_footer(chunk_evidence_ids)
        self._base_slide_title = base_title

    def add_page_break(self) -> None:
        self._slide = None
        self._frame = None
        self._text_units = 0
        self._has_table = False
        self._slide_source_ids = []
        self._source_footer_shape = None

    def add_section_break(self) -> None:
        self.add_page_break()

    def add_references_slide(self, *_, **__) -> None:
        if self._is_reusable_references_slide():
            self._clear_references_placeholder_body()
        else:
            self.add_slide("References")
        reference_lines = [
            reference_line(index, citation)
            for index, citation in enumerate(self._all_citations[:18], start=1)
        ]
        self.add_bullet_list(reference_lines)

    def finalize(self) -> None:
        self._presentation.save(str(self._path))

    def _ensure_slide(self) -> None:
        if self._slide is None:
            self.add_slide(self._base_slide_title or "Details")

    def _text_frame(self):
        from pptx.util import Inches

        self._ensure_slide()
        if self._frame is None:
            box = self._slide.shapes.add_textbox(Inches(0.85), Inches(1.15), Inches(11.7), Inches(5.35))
            self._frame = box.text_frame
            self._frame.clear()
            self._frame.word_wrap = True
        return self._frame

    def _continue_slide(self) -> None:
        base_title = self._base_slide_title or "Details"
        self.add_slide(f"{base_title} (cont.)")
        self._base_slide_title = base_title

    def _paragraph_units(self, text: str) -> int:
        return max(1, min(3, (len(text) + 109) // 110))

    def _set_background(self, slide, color: str) -> None:
        from pptx.dml.color import RGBColor

        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(color)

    def _next_content_background(self) -> str:
        color = CONTENT_SLIDE_BACKGROUNDS[self._content_slide_index % len(CONTENT_SLIDE_BACKGROUNDS)]
        self._content_slide_index += 1
        return color

    def _add_text(self, slide, text: str, left: float, top: float, width: float, height: float, size: int, color: str, *, bold: bool = False) -> None:
        from pptx.dml.color import RGBColor
        from pptx.util import Inches

        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = compact(text, 260)
        paragraph.font.name = "Aptos"
        paragraph.font.size = self._pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string(color)

    def _update_source_footer(self, evidence_ids: list[str] | None) -> None:
        from pptx.dml.color import RGBColor
        from pptx.util import Inches

        self._ensure_slide()
        for evidence_id in normalize_evidence_ids(evidence_ids):
            if evidence_id in self._citations and evidence_id not in self._slide_source_ids:
                self._slide_source_ids.append(evidence_id)
        if not self._slide_source_ids:
            return

        labels = [short_citation(self._citations[evidence_id]) for evidence_id in self._slide_source_ids]
        summary = compact("Sources: " + "; ".join(labels), 190)
        if self._source_footer_shape is None:
            self._source_footer_shape = self._slide.shapes.add_textbox(
                Inches(0.85),
                Inches(6.72),
                Inches(11.7),
                Inches(0.35),
            )
            self._source_footer_shape.text_frame.word_wrap = True

        text_frame = self._source_footer_shape.text_frame
        text_frame.clear()
        paragraph = text_frame.paragraphs[0]
        paragraph.text = summary
        paragraph.font.name = "Aptos"
        paragraph.font.size = self._pt(8)
        paragraph.font.color.rgb = RGBColor.from_string("53676A")

    def _is_reusable_references_slide(self) -> bool:
        if self._slide is None or self._has_table or self._source_footer_shape is not None:
            return False
        if _normalized_reference_title(self._current_slide_title) not in REFERENCE_SLIDE_TITLES:
            return False
        body = self._frame_text()
        return not body or _is_reference_placeholder_text(body)

    def _clear_references_placeholder_body(self) -> None:
        if self._frame is not None:
            self._frame.clear()
        self._text_units = 0
        self._slide_source_ids = []

    def _frame_text(self) -> str:
        if self._frame is None:
            return ""
        return compact(" ".join(paragraph.text for paragraph in self._frame.paragraphs), 120)

    def _pt(self, value: int):
        from pptx.util import Pt

        return Pt(value)


def _normalized_reference_title(value: str) -> str:
    import re

    return re.sub(r"[^a-z]+", " ", value.lower()).strip()


def _is_reference_placeholder_text(value: str) -> bool:
    normalized = _normalized_reference_title(value)
    parts = normalized.split()
    return normalized in REFERENCE_SLIDE_TITLES or bool(parts) and all(
        part in REFERENCE_SLIDE_TITLES
        for part in parts
    )
