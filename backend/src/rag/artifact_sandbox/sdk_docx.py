"""DOCX layout primitives for sandbox build programs."""

from __future__ import annotations

from datetime import datetime
from typing import Any

from rag.artifact_jobs.contracts import EvidenceCitation

from .sdk_shared import citation_map, item_parts, reference_line, row_parts, validated_output_path, with_sources


class DocxArtifact:
    def __init__(
        self,
        *,
        path: str,
        title: str,
        subtitle: str | None,
        generated_at: datetime,
        citations: list[EvidenceCitation],
    ) -> None:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt

        self._path = validated_output_path(path)
        self._citations = citation_map(citations)
        self._all_citations = list(citations)
        self._document = Document()
        section = self._document.sections[0]
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        self._document.styles["Normal"].font.name = "Aptos"
        self._document.styles["Normal"].font.size = Pt(10.5)
        heading = self._document.add_heading(title, level=0)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if subtitle:
            para = self._document.add_paragraph(subtitle, style="Subtitle")
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta = self._document.add_paragraph(f"Generated {generated_at.strftime('%Y-%m-%d %H:%M UTC')}")
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def add_heading(self, text: str, *, level: int = 1) -> None:
        self._document.add_heading(str(text), level=max(1, min(level, 4)))

    def add_paragraph(self, text: str, *, evidence_ids: list[str] | None = None) -> None:
        self._document.add_paragraph(with_sources(str(text), evidence_ids, self._citations))

    def add_bullet_list(self, items: list[Any], *, evidence_ids: list[str] | None = None) -> None:
        for item in items:
            text, ids = item_parts(item, evidence_ids)
            self._document.add_paragraph(with_sources(text, ids, self._citations), style="List Bullet")

    def add_numbered_list(self, items: list[Any], *, evidence_ids: list[str] | None = None) -> None:
        for item in items:
            text, ids = item_parts(item, evidence_ids)
            self._document.add_paragraph(with_sources(text, ids, self._citations), style="List Number")

    def add_callout(self, text: str, *, evidence_ids: list[str] | None = None) -> None:
        paragraph = self._document.add_paragraph()
        run = paragraph.add_run(with_sources(str(text), evidence_ids, self._citations))
        run.bold = True

    def add_table(self, headers: list[str], rows: list[Any], *, evidence_ids: list[str] | None = None) -> None:
        if len(headers) > 5:
            self._set_landscape()
        table = self._document.add_table(rows=1, cols=len(headers) + 1)
        table.style = "Table Grid"
        table.autofit = True
        for index, header in enumerate(headers):
            table.rows[0].cells[index].text = str(header)
        table.rows[0].cells[-1].text = "Source"
        for row in rows:
            values, ids = row_parts(row, evidence_ids)
            cells = table.add_row().cells
            for index, value in enumerate(values[: len(headers)]):
                cells[index].text = value
            cells[-1].text = with_sources("", ids, self._citations).strip(" ()")

    def add_page_break(self) -> None:
        self._document.add_page_break()

    def add_references(self, *_, **__) -> None:
        self.add_page_break()
        self.add_heading("References")
        for index, citation in enumerate(self._all_citations, start=1):
            self._document.add_paragraph(reference_line(index, citation))

    def finalize(self) -> None:
        self._document.save(str(self._path))

    def _set_landscape(self) -> None:
        from docx.enum.section import WD_ORIENT

        section = self._document.sections[0]
        if section.orientation == WD_ORIENT.LANDSCAPE:
            return
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
